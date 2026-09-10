"""Generate the final versioned black-only research-complete Word report."""

from __future__ import annotations

import argparse
import csv
import json
import statistics
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.shared import Inches, Pt

import _multiclass_bootstrap
import generate_training_details_black_report as base
from document_io import save_docx_with_fallback


ROOT = _multiclass_bootstrap.REPO_ROOT
MODEL = ROOT / "model"
RESULTS = MODEL / "results"
DOCS = MODEL / "docs"
DEFAULT_OUTPUT = DOCS / "Pneumonia_Model_Training_Details_Black_Research_Complete_20260802.docx"
SCREENSHOT_IDS = {"IMG-06", "IMG-07", "IMG-04", "IMG-03", "IMG-11", "IMG-05", "IMG-VNET", "IMG-02", "IMG-01", "IMG-09"}


def read_json(path: Path):
    return json.loads(path.read_text(encoding="utf-8"))


def read_csv(path: Path):
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def gate() -> list[str]:
    failures = []
    required = [
        RESULTS / "audits/current_state_audit_continuation_v2.json",
        RESULTS / "evidence/research_evidence_matrix.json",
        RESULTS / "preprocessing_ablation_analysis.json",
        RESULTS / "tuning_trials.csv",
        RESULTS / "seed_robustness.csv",
        RESULTS / "final_selection.json",
        RESULTS / "runs/final_locked_test/metrics.json",
        RESULTS / "gradcam/gradcam_audit.json",
        RESULTS / "calibration/final_temperature_scaling.json",
        RESULTS / "efficiency/final_efficiency.json",
    ]
    for path in required:
        if not path.exists():
            failures.append(f"missing required final artifact: {path}")
    if failures:
        return failures
    summary = read_csv(RESULTS / "summary.csv")
    if len([row for row in summary if row["stage"] == "baseline" and row["status"] == "completed"]) != 8:
        failures.append("final report requires eight completed baselines")
    if len([row for row in summary if row["stage"] == "preprocessing_ablation" and row["status"] == "completed"]) != 9:
        failures.append("final report requires nine completed preprocessing ablations")
    robustness = [row for row in summary if row["stage"] == "robustness" and row["status"] == "completed"]
    if len(robustness) != 3 or {row["seed"] for row in robustness} != {"42", "123", "2026"}:
        failures.append("final report requires completed robustness seeds 42, 123, and 2026")
    tuning = read_csv(RESULTS / "tuning_trials.csv")
    if len([row for row in tuning if row["status"] == "completed"]) < 3:
        failures.append("final report requires completed tuning evidence")
    selection = read_json(RESULTS / "final_selection.json")
    if selection.get("status") != "locked_test_evaluated":
        failures.append("final selection is not in locked_test_evaluated state")
    evidence = read_json(RESULTS / "evidence/research_evidence_matrix.json")
    if evidence.get("record_count") != 25 or evidence.get("screenshot_visible_entry_count") != 10:
        failures.append("reference evidence matrix is incomplete")
    phash = read_json(RESULTS / "perceptual_hash_manual_review/perceptual_hash_manual_review.json")
    if not phash.get("visual_review_completed"):
        failures.append("manual pHash review is incomplete")
    return failures


def hyperlink(paragraph, label: str, url: str) -> None:
    relationship_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "000000")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.extend([color, underline])
    text = OxmlElement("w:t")
    text.text = label
    run.extend([properties, text])
    link.append(run)
    paragraph._p.append(link)


def add_link_line(doc, label: str, url: str) -> None:
    paragraph = doc.add_paragraph()
    lead = paragraph.add_run(f"{label}: ")
    base.set_run_font(lead, size=9, bold=True)
    hyperlink(paragraph, url, url)


def page_break(doc) -> None:
    doc.add_page_break()


def add_figure(doc, path: Path, caption: str, *, width: float = 6.35) -> None:
    if not path.exists():
        return
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(path), width=Inches(width))
    caption_p = doc.add_paragraph(style="Caption")
    caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = caption_p.add_run(caption)
    base.set_run_font(run, size=8.5)


def cover(doc, generated_at: str, selection: dict, test: dict) -> None:
    for _ in range(3):
        doc.add_paragraph()
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run("Pneumonia Model Training Details")
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.add_run("Research-complete three-class chest X-ray classification report")
    base.paragraph_bottom_border(subtitle, size="12")
    base.add_body(doc, "Normal | Viral Pneumonia | Bacterial Pneumonia")
    doc.add_paragraph()
    base.add_table(
        doc,
        ("Report field", "Final value"),
        [
            ("Status", "Research pipeline complete; single locked-test evaluation completed"),
            ("Selected run", selection["selected_run_id"]),
            ("Architecture / preprocessing", f"{selection['architecture']} / {selection['preprocessing']}"),
            ("Locked-test macro F1", f"{test['macro_f1']:.6f}"),
            ("Locked-test balanced accuracy", f"{test['balanced_accuracy']:.6f}"),
            ("Locked-test accuracy", f"{test['accuracy']:.6f}"),
            ("Generated", generated_at),
        ],
        (2600, 6760),
        font_size=9.2,
    )
    base.add_keep_note(
        doc,
        "All model-selection decisions use the locked validation set. The test set was evaluated once after architecture, preprocessing, hyperparameters, seed protocol, and checkpoint policy were frozen.",
    )
    page_break(doc)


def executive(doc, audit, selection, test, robustness, calibration, efficiency) -> None:
    base.add_heading(doc, "Executive Summary", 1)
    scores = [float(row["validation_macro_f1"]) for row in robustness]
    base.add_body(
        doc,
        f"Eight baseline classifiers were trained and verified, nine new preprocessing ablations were completed on three representative models, three bounded Optuna studies were executed, and the selected configuration was independently retrained at seeds 42, 123, and 2026. The final selected model is {selection['architecture']} with {selection['preprocessing']} preprocessing. Robustness validation macro F1 was {statistics.mean(scores):.6f} +/- {statistics.stdev(scores):.6f} (sample standard deviation).",
    )
    base.add_body(
        doc,
        f"The single locked-test evaluation produced accuracy {test['accuracy']:.6f}, balanced accuracy {test['balanced_accuracy']:.6f}, macro F1 {test['macro_f1']:.6f}, macro one-vs-rest ROC AUC {test.get('roc_auc_ovr_macro', 0):.6f}, and expected calibration error {test['expected_calibration_error']:.6f} before post-hoc temperature scaling.",
    )
    base.add_body(
        doc,
        f"Temperature {calibration['temperature']:.6f} was fit on validation predictions only. On the already frozen test outputs, ECE changed from {calibration['test_before']['expected_calibration_error']:.6f} to {calibration['test_after']['expected_calibration_error']:.6f}; accuracy and macro F1 are unchanged because temperature scaling does not change the predicted class ordering.",
    )
    base.add_body(
        doc,
        f"End-to-end local batch-1 latency, including file decode, preprocessing, transfer, model, softmax, and argmax, was {efficiency['end_to_end_latency_ms']['mean']:.2f} ms/image on average ({efficiency['end_to_end_latency_ms']['p95']:.2f} ms at the 95th percentile). This is an engineering measurement, not a clinical deployment claim.",
    )
    base.add_bullet(doc, "Source data remained immutable: 17,591 files, 2,483,915,296 bytes, zero inventory mismatches.")
    base.add_bullet(doc, "Locked manifests contain 4,076 train, 874 validation, and 874 test images with zero assigned group or exact-hash leakage.")
    base.add_bullet(doc, "All 16 pHash candidates were manually inspected; all were distinct and no manifest change was made.")
    base.add_bullet(doc, "All eight baseline artifact audits passed; completed baselines were not rerun.")
    base.add_bullet(doc, "YOLOv10 classification remained excluded because no genuine compatible official classification head/checkpoint was verified.")
    base.add_keep_note(doc, "Clinical boundary: this is a controlled benchmark on one public pediatric CXR source. It is not a clinical diagnosis system or external validation study.")


def status_overview(doc, tuning_rows, robustness, selection, gradcam) -> None:
    completed_trials = sum(row["status"] == "completed" for row in tuning_rows)
    base.add_heading(doc, "Research Completion Status", 2)
    base.add_table(
        doc,
        ("Research Stage", "Previous Status", "New Status", "Evidence/File"),
        [
            ("Baseline benchmark", "Complete", "Complete; 8/8 supported models verified", "model/results/audits/completed_baseline_artifact_audit.json"),
            ("Preprocessing ablations", "0 of 3", "Complete; 3/3 methods, 9/9 new runs", "model/results/preprocessing_ablation_analysis.json"),
            ("Focused comparisons", "Not started", "Complete; validation-only aliases, no redundant retraining", "model/results/selection_resolution.json"),
            ("Optuna tuning", "0 trials", f"Complete; {completed_trials}/{len(tuning_rows)} completed trials; every state retained", "model/results/tuning_trials.csv"),
            ("Multi-seed robustness", "Not started", f"Complete; {len(robustness)}/3 seeds", "model/results/seed_robustness.csv"),
            ("Locked-test evaluation", "Not performed", "Complete; one frozen representative evaluation", "model/results/runs/final_locked_test/metrics.json"),
            ("Final selection", "Not established", f"Established and frozen: {selection['selected_run_id']}", "model/results/final_selection.json"),
            ("Explainability", "Not completed", f"Complete; {len(gradcam['records'])} post-freeze Grad-CAM cases", "model/results/gradcam/gradcam_audit.json"),
            ("Documentation", "Partially complete", "Complete; versioned Word plus Markdown/CSV/JSON evidence", "model/docs and model/results"),
        ],
        (1800, 1550, 2800, 3210),
        font_size=7.2,
    )


def repository_and_data(doc, audit, split, source, phash_rows) -> None:
    page_break(doc)
    base.add_heading(doc, "1. Repository, Source, and Locked Dataset Audit", 1)
    base.add_body(doc, "The continuation began by auditing the supplied Word report, repository code, manifests, checkpoints, logs, metrics, and training environment. No baseline was rerun and no source image or locked manifest was modified.")
    base.add_table(
        doc,
        ("Audit item", "Verified value"),
        [
            ("Repository branch", audit["repository"]["branch"]),
            ("Starting commit", audit["repository"]["head"]),
            ("Source root", source["source_root"]),
            ("Source files / bytes", f"{source['source_file_count']:,} / {source['source_bytes']:,}"),
            ("Inventory SHA-256", source["source_content_inventory_hash"]),
            ("Manifest version", split["manifest_version"]),
            ("Manifest SHA-256", split["manifest_version_hash"]),
            ("Class order", "Normal, Viral, Bacterial"),
            ("Assigned leakage", "0 group leakage; 0 exact-hash leakage"),
        ],
        (2600, 6760),
        font_size=8.4,
    )
    base.add_heading(doc, "1.1 Split composition", 2)
    base.add_table(
        doc,
        ("Split", "Normal", "Viral", "Bacterial", "Total", "Groups"),
        [
            (name.title(), values["Normal"], values["Viral"], values["Bacterial"], split["split_total_counts"][name], split["split_group_counts"][name])
            for name, values in (("train", split["split_counts"]["train"]), ("val", split["split_counts"]["val"]), ("test", split["split_counts"]["test"]))
        ],
        (1460, 1300, 1300, 1500, 1500, 2300),
        font_size=8.8,
    )
    base.add_body(doc, "The original folder split was not preserved because 264 derived patient/image groups crossed the original folders. A deterministic group-aware assignment was created without consulting image content or test metrics. Exact-hash groups were unioned before splitting.")
    base.add_heading(doc, "1.2 What 'canonical image' and the duplicate counts mean", 2)
    base.add_body(doc, "A canonical image is the selected primary representative of an intended dataset item after excluding macOS metadata, verified nested copies, unrecognized layout artifacts, and redundant byte-identical content. The source contained 5,856 valid canonical-role physical images before exact deduplication and 5,824 unique canonical images after 32 redundant exact images were excluded.")
    base.add_body(doc, "The value 5,888 in duplicate_report.csv is not the number of training images removed. It is a physical report-row count describing exact-content relationships across canonical files and verified nested copies. SHA-256/byte identity found exact duplicates. The 64-bit DCT pHash was used only to flag visually similar candidates for human review.")
    base.add_table(
        doc,
        ("Duplicate/audit concept", "Count", "Meaning"),
        [
            ("Canonical redundant exact images excluded", "32", "Unique-data deduplication decision"),
            ("Physical exact-duplicate report rows", "5,888", "Relationships including nested copies"),
            ("Original folder group crossings", "264", "Reason original split was rejected"),
            ("Nested-copy mismatches", "4", "Excluded source-layout anomalies"),
            ("pHash candidates", "16", "Manual review queue, not automatic duplicates"),
            ("Confirmed near duplicates", "0", "All 16 pairs were distinct"),
        ],
        (3000, 1000, 5360),
        font_size=8.7,
    )
    base.add_heading(doc, "1.3 Manual review of all 16 pHash candidates", 2)
    base.add_body(doc, "Five candidates crossed assigned splits. Each evidence sheet was inspected for thoracic anatomy, clavicle/rib geometry, cardiomediastinal contour, diaphragm shape, lung volume, positioning, exposure, and markers. Similar centered radiograph geometry explained the low pHash distance; no pair represented the same image.")
    base.add_table(
        doc,
        ("Pair", "Assigned splits", "Decision", "Visual basis"),
        [
            (row["pair_id"], f"{row['left_assigned_split']} / {row['right_assigned_split']}", row["manual_decision"], row["manual_rationale"])
            for row in phash_rows
        ],
        (2250, 1300, 2200, 3610),
        font_size=7.3,
    )


def preprocessing(doc) -> None:
    page_break(doc)
    base.add_heading(doc, "2. Preprocessing and Augmentation", 1)
    base.add_keep_note(doc, "Exact tensor pipeline: locked manifest -> open image -> enhancement/RGB -> training-only flip -> training-only rotation -> aspect-ratio-preserving resize-to-fill and center crop -> float32 [0,1] -> CHW -> ImageNet normalization -> 3 x 224 x 224.")
    steps = [
        ("1", "Open the original", "Pillow decodes the immutable source file named by the locked manifest."),
        ("2", "Convert/enhance", "Baseline converts to RGB. CLAHE/global HE convert to grayscale, enhance, then replicate to RGB. VIRIDIS maps grayscale intensity to artificial RGB."),
        ("3", "Horizontal flip (training only)", "Probability 0.5. Validation/test never flip."),
        ("4", "Rotation (training only)", "Uniform random angle from -5 to +5 degrees with bilinear resampling and black corner fill."),
        ("5", "Resize and crop", "Pillow ImageOps.fit preserves aspect ratio, scales until 224 x 224 is filled, and center-crops excess edges. It does not pad, letterbox, or stretch."),
        ("6", "Scale pixels", "Convert to float32 and divide by 255, yielding values in [0,1]."),
        ("7", "Transpose", "Convert H x W x C into C x H x W for PyTorch."),
        ("8", "Normalize", "Subtract ImageNet means [0.485,0.456,0.406] and divide by [0.229,0.224,0.225]."),
    ]
    base.add_table(doc, ("Step", "Operation", "Exact meaning"), steps, (700, 2400, 6260), font_size=8.3)
    base.add_heading(doc, "2.1 Enhancement conditions", 2)
    base.add_table(
        doc,
        ("Condition", "Included?", "Exact implementation", "Why / boundary"),
        [
            ("No enhancement", "Yes; completed control", "RGB conversion only", "Preserves original intensity appearance and pretrained RGB contract."),
            ("CLAHE", "Yes", "OpenCV CLAHE; clipLimit 2.0; tileGridSize 8 x 8 on grayscale", "Local contrast candidate supported partially by CXR literature; effect tested independently."),
            ("Global HE", "Yes", "Pillow ImageOps.equalize on grayscale", "Direct pediatric CXR evidence motivates a controlled comparison."),
            ("VIRIDIS", "Yes; exploratory", "OpenCV COLORMAP_VIRIDIS on grayscale", "No direct exact-task validation was found; cannot be described as adding clinical information."),
            ("PCA", "No", "No raw-image PCA", "No leakage-safe, train-only feature-space representation/question was defined."),
            ("FILM", "No", "No implementation", "No verified conditioning signal or reproducible definition existed."),
            ("SAS tuning / automated grading", "No", "No implementation", "Undefined in repository evidence; no results fabricated."),
            ("Segmentation/ROI masking", "No", "No lung mask", "Localization literature is Partial evidence only; adding a segmenter would change the task and error chain."),
            ("Vertical flip / large rotation", "No", "Excluded", "Anatomically implausible or excessive for this controlled CXR study."),
        ],
        (1500, 1250, 3000, 3610),
        font_size=7.3,
    )
    base.add_body(doc, "Enhancement is deterministic and applied before augmentation and crop. Augmentation is sampled on-the-fly during every training access; no augmented image files are written into the locked dataset. Validation and test use deterministic transforms only.")
    add_figure(doc, RESULTS / "preprocessing_examples/preprocessing_comparison.png", "Pre-normalization examples for the four supported conditions. VIRIDIS is an exploratory artificial colormap.", width=6.2)


def training(doc, model_review, baseline_rows) -> None:
    page_break(doc)
    base.add_heading(doc, "3. Models and Training Procedure", 1)
    common = model_review["common_control"]
    base.add_table(
        doc,
        ("Setting", "Final controlled definition"),
        [
            ("Task", "Single-label three-class classification: Normal, Viral, Bacterial"),
            ("Weights", "Official/ImageNet classification weights; classifier replaced with three logits"),
            ("Fine-tuning", "Full network trainable"),
            ("Loss", common["loss"]),
            ("Class weights", "Normal 1.229563; Viral 1.307668; Bacterial 0.703244"),
            ("Optimizer", "AdamW; initial learning rate 0.0001; weight decay 0.0001"),
            ("Scheduler", common["scheduler"]),
            ("Epoch ceiling", "12; best checkpoint by validation macro F1; early stopping patience 4"),
            ("Mixed precision", "Enabled on CUDA"),
            ("Seed", "42 for controlled comparison; later robustness seeds 42, 123, 2026"),
            ("Test access", "Prohibited until final selection lock"),
        ],
        (2500, 6860),
        font_size=8.6,
    )
    base.add_heading(doc, "3.1 Meaning of learning rate, weight decay, and 12 epochs", 2)
    base.add_body(doc, "Learning rate 0.0001 is AdamW's initial step size. It is deliberately small for full-network fine-tuning so useful pretrained features are not destroyed by large updates. ReduceLROnPlateau halves the rate after two macro-F1 plateaus.")
    base.add_body(doc, "Weight decay 0.0001 is regularization: it gently discourages excessively large weights. It is not a second learning rate. The value is an engineering control shared by the completed comparisons, not a configuration copied blindly from another modality or paper.")
    base.add_body(doc, "Maximum 12 epochs means at most twelve complete passes through 4,076 training images. It is a ceiling. The best validation checkpoint is retained, and early stopping ends after four epochs without an improvement greater than 0.0001. Baseline best epochs ranged from 5 to 11.")
    base.add_heading(doc, "3.2 How batch size was determined", 2)
    base.add_body(doc, "Architecture-specific batch sizes were selected through actual CUDA forward/backward resource calibration on the 4 GB RTX 2050. They are the largest premeasured safe values used for each architecture. Within every preprocessing comparison, the completed control and its variants use the same batch size.")
    base.add_table(
        doc,
        ("Architecture", "Batch", "Parameters", "Weight source / scientific role"),
        [
            (row["architecture"], row["batch_size"], f"{int(row['parameter_count']):,}", row["literature_transferability"])
            for row in model_review["model_records"]
        ],
        (1700, 800, 1500, 5360),
        font_size=7.4,
    )
    base.add_body(doc, "Different batch sizes are a cross-architecture limitation because batch-induced gradient noise may differ. The preprocessing ablations avoid this confound within architecture. Native Ultralytics RandomResizedCrop, RandAugment, and random erasing were not invoked only for YOLO; all eight models used the same repository transform contract.")
    base.add_heading(doc, "3.3 Training and validation loss distinction", 2)
    base.add_body(doc, "Training uses balanced class-weighted cross entropy. Saved validation/test loss is unweighted cross entropy. Checkpoints are selected by macro F1, not loss. This distinction was audited and is labeled explicitly in continuation histories and TensorBoard. The focal-loss implementation was corrected before tuning so class weights scale the focal term without distorting the target probability p_t.")


def baseline_results(doc, baseline_rows) -> None:
    page_break(doc)
    base.add_heading(doc, "4. Completed Baseline Results (Validation Only)", 1)
    base.add_body(doc, "All eight run directories passed independent checkpoint, hash, prediction-row, metric, confusion-matrix, history, parameter-count, output-shape, and file-size verification. Original TensorBoard files were header-only, but CSV/JSON histories are valid; completed baselines were not rerun merely to populate TensorBoard.")
    ordered = sorted(baseline_rows, key=lambda row: float(row["validation_macro_f1"]), reverse=True)
    base.add_table(
        doc,
        ("Rank", "Run", "Accuracy", "Macro F1", "Recall N/V/B", "Best epoch", "Parameters"),
        [
            (
                index,
                row["run_id"],
                f"{float(row['validation_accuracy']):.4f}",
                f"{float(row['validation_macro_f1']):.4f}",
                f"{float(row['recall_normal']):.3f} / {float(row['recall_viral']):.3f} / {float(row['recall_bacterial']):.3f}",
                row["best_epoch"],
                f"{int(float(row['parameter_count'])):,}",
            )
            for index, row in enumerate(ordered, 1)
        ],
        (700, 2400, 1000, 1000, 1600, 1100, 1560),
        font_size=7.2,
    )
    base.add_body(doc, "DenseNet169 ranked first at validation macro F1 0.840293. DenseNet201 had the strongest Viral recall at 0.699552. Viral recall was the recurring weak point across all baselines, while Normal recall was consistently high. The small differences among DenseNet/ResNet baselines at one seed did not establish universal superiority; multi-seed evidence was required after final tuning.")


def final_configuration_table(doc, model_review, selection) -> None:
    """Add the required all-model configuration table in readable continuation parts."""
    base.add_heading(doc, "7.2 Final configuration table for all supported architectures", 2)
    base.add_body(
        doc,
        "The required wide table is split into three continuation parts so every field remains readable. "
        "Rows not selected for the terminal experiment retain their verified completed-baseline configuration; "
        "the selected row reports the frozen tuned configuration actually used for the three-seed study and test.",
    )
    final_config = read_json(Path(selection["selected_run_dir"]) / "config.json")
    exact_implementation = {
        "densenet121": "torchvision.models.densenet121; classifier Linear(...,3)",
        "densenet169": "torchvision.models.densenet169; classifier Linear(...,3)",
        "densenet201": "torchvision.models.densenet201; classifier Linear(...,3)",
        "resnet50": "torchvision.models.resnet50; fc Linear(...,3)",
        "resnet101": "torchvision.models.resnet101; fc Linear(...,3)",
        "resnet152": "torchvision.models.resnet152; fc Linear(...,3)",
        "yolov5_cls": "Ultralytics YOLOv5s-cls v7.0 + verified logits adapter; Classify Linear(...,3)",
        "yolov8_cls": "Ultralytics YOLOv8n-cls v8.3.0 + verified logits adapter; Classify Linear(...,3)",
    }
    tuned_models = {"densenet169", "densenet201", "resnet50"}
    rows = []
    for record in model_review["model_records"]:
        architecture = record["architecture"]
        selected = architecture == selection["architecture"]
        config = final_config if selected else {
            "input_size": record["input_size"],
            "batch_size": record["batch_size"],
            "optimizer": "adamw",
            "learning_rate": 0.0001,
            "weight_decay": 0.0001,
            "epochs": 12,
            "scheduler": "reduce_on_plateau",
            "freeze_backbone": False,
            "label_smoothing": 0.0,
            "loss": "cross_entropy",
            "class_weighting": "balanced",
            "augmentation": {"horizontal_flip_probability": 0.5, "rotation_degrees": 5.0},
            "preprocessing": "baseline",
        }
        augmentation = config["augmentation"]
        status = (
            "Final selected; tuned; three-seed retrain; locked test"
            if selected
            else "Baseline + ablation + equal tuning; not final"
            if architecture in tuned_models
            else "Completed baseline; outside bounded final tuning"
        )
        rows.append(
            {
                "Model": architecture,
                "Exact Implementation": exact_implementation[architecture],
                "Input Size": config["input_size"],
                "Pretrained Weights": record["pretrained_weight_source"],
                "Batch Size": config["batch_size"],
                "Optimizer": str(config["optimizer"]).upper(),
                "Backbone LR": config["learning_rate"],
                "Head LR": config["learning_rate"],
                "Weight Decay": config["weight_decay"],
                "Epoch Limit": config["epochs"],
                "Scheduler": config["scheduler"],
                "Freeze Strategy": "Full-network fine-tuning" if not config["freeze_backbone"] else "Frozen backbone",
                "Dropout": "No project-level override; architecture default",
                "Label Smoothing": config["label_smoothing"],
                "Loss": "Focal, gamma 2" if config["loss"] == "focal" else "Cross entropy",
                "Balancing": config["class_weighting"],
                "Augmentation": f"flip p={augmentation['horizontal_flip_probability']}; rotation +/-{augmentation['rotation_degrees']} deg",
                "Preprocessing": config["preprocessing"],
                "Selection Metric": "Validation macro F1",
                "Source Basis": record["literature_transferability"],
                "Final Status": status,
            }
        )

    base.add_heading(doc, "Part A - identity, inputs, weights, and status", 3)
    base.add_table(
        doc,
        ("Model", "Exact Implementation", "Input Size", "Pretrained Weights", "Batch Size", "Final Status"),
        [tuple(row[key] for key in ("Model", "Exact Implementation", "Input Size", "Pretrained Weights", "Batch Size", "Final Status")) for row in rows],
        (1100, 2600, 850, 2100, 800, 1910),
        font_size=6.8,
    )
    base.add_heading(doc, "Part B - optimization and fine-tuning", 3)
    base.add_table(
        doc,
        ("Model", "Optimizer", "Backbone LR", "Head LR", "Weight Decay", "Epoch Limit", "Scheduler", "Freeze Strategy", "Dropout"),
        [tuple(row[key] for key in ("Model", "Optimizer", "Backbone LR", "Head LR", "Weight Decay", "Epoch Limit", "Scheduler", "Freeze Strategy", "Dropout")) for row in rows],
        (1050, 850, 850, 850, 900, 700, 1200, 1500, 1460),
        font_size=6.3,
    )
    base.add_heading(doc, "Part C - objective, augmentation, preprocessing, and evidence", 3)
    base.add_table(
        doc,
        ("Model", "Label Smoothing", "Loss", "Balancing", "Augmentation", "Preprocessing", "Selection Metric", "Source Basis"),
        [tuple(row[key] for key in ("Model", "Label Smoothing", "Loss", "Balancing", "Augmentation", "Preprocessing", "Selection Metric", "Source Basis")) for row in rows],
        (950, 900, 1100, 1000, 1600, 1100, 1250, 1460),
        font_size=6.3,
    )


def ablation_results(doc, ablation) -> None:
    page_break(doc)
    base.add_heading(doc, "5. Preprocessing Ablation Results", 1)
    base.add_body(doc, "Three completed controls (DenseNet169, DenseNet201, ResNet50) were reused. Nine new runs evaluated CLAHE, global histogram equalization, and exploratory VIRIDIS. Primary selection used validation macro F1 with balanced accuracy, minimum class recall, log loss, size, latency, and run ID as fixed tie-breakers.")
    base.add_table(
        doc,
        ("Architecture", "Condition", "Macro F1", "Delta vs control", "Paired 95% interval", "Recall N/V/B", "Best epoch"),
        [
            (
                row["architecture"],
                row["preprocessing"],
                f"{row['validation_macro_f1']:.6f}",
                f"{row['delta_macro_f1_vs_same_model_baseline']:+.6f}",
                f"[{row['delta_bootstrap_ci_lower']:+.4f}, {row['delta_bootstrap_ci_upper']:+.4f}]",
                f"{row['recall_normal']:.3f}/{row['recall_viral']:.3f}/{row['recall_bacterial']:.3f}",
                row["best_epoch"],
            )
            for row in ablation["records"]
        ],
        (1450, 1550, 1100, 1300, 1750, 1400, 810),
        font_size=7.0,
    )
    for architecture, ranking in ablation["within_model_rankings"].items():
        base.add_bullet(doc, f"{architecture} ranking: " + " > ".join(ranking))
    base.add_body(doc, "The paired class-stratified bootstrap uses 2,000 validation resamples and is descriptive uncertainty on the same selection set, not an independent confirmatory test. VIRIDIS remains exploratory regardless of rank.")


def tuning_results(doc, tuning_rows) -> None:
    page_break(doc)
    base.add_heading(doc, "6. Bounded Hyperparameter Tuning", 1)
    base.add_body(doc, "Three validation-only Optuna TPE studies were run on the validation-best preprocessing parent for each preregistered representative: DenseNet169, DenseNet201, and ResNet50. Each study used four trials at an equal six-epoch budget. Trial 0 was the exact parent setting as a preregistered control. Input size, batch size, weights, preprocessing, and seed remained fixed.")
    base.add_body(doc, "Search dimensions were optimizer, conditional learning rate, weight decay, scheduler, cross entropy versus corrected focal loss, balanced versus unweighted loss, label smoothing, and conservative versus no augmentation. Dropout, freezing, input size, and batch size were not mixed into the small common search.")
    base.add_table(
        doc,
        ("Study / trial", "Parent", "Status", "Parameters", "Macro F1", "Balanced acc.", "Min recall"),
        [
            (
                row["trial_id"],
                row["parent_run_id"],
                row["status"],
                row["parameters_json"],
                row["validation_macro_f1"] or "-",
                row["validation_balanced_accuracy"] or "-",
                row["minimum_per_class_recall"] or "-",
            )
            for row in tuning_rows
        ],
        (1100, 1700, 900, 3100, 900, 900, 760),
        font_size=6.6,
    )
    base.add_body(doc, "All completed, failed, or interrupted trial states were retained in SQLite plus CSV/JSON exports. The locked test was not accessible to the tuning process.")


def final_results(doc, robustness, selection, test) -> None:
    page_break(doc)
    base.add_heading(doc, "7. Multi-Seed Robustness and Single Locked-Test Result", 1)
    base.add_table(
        doc,
        ("Seed", "Run", "Validation accuracy", "Balanced accuracy", "Macro F1", "Recall N/V/B", "Best epoch"),
        [
            (
                row["seed"], row["run_id"], f"{float(row['validation_accuracy']):.4f}",
                f"{float(row['validation_balanced_accuracy']):.4f}", f"{float(row['validation_macro_f1']):.4f}",
                f"{float(row['recall_normal']):.3f}/{float(row['recall_viral']):.3f}/{float(row['recall_bacterial']):.3f}", row["best_epoch"],
            )
            for row in sorted(robustness, key=lambda row: int(row["seed"]))
        ],
        (700, 2300, 1150, 1200, 1000, 1850, 1160),
        font_size=7.1,
    )
    scores = [float(row["validation_macro_f1"]) for row in robustness]
    base.add_body(doc, f"Validation macro F1 mean = {statistics.mean(scores):.6f}; sample standard deviation = {statistics.stdev(scores):.6f}. Seed 42 was predeclared as the representative checkpoint policy and was not selected using test performance.")
    base.add_heading(doc, "7.1 Locked-test metrics", 2)
    base.add_table(
        doc,
        ("Metric", "Value"),
        [
            ("Accuracy", f"{test['accuracy']:.6f}"),
            ("Balanced accuracy", f"{test['balanced_accuracy']:.6f}"),
            ("Macro precision", f"{test['macro_precision']:.6f}"),
            ("Macro recall", f"{test['macro_recall']:.6f}"),
            ("Macro F1", f"{test['macro_f1']:.6f}"),
            ("Weighted F1", f"{test['weighted_f1']:.6f}"),
            ("ROC AUC OVR macro", f"{test['roc_auc_ovr_macro']:.6f}"),
            ("PR AUC macro", f"{test['pr_auc_macro']:.6f}"),
            ("Log loss", f"{test['log_loss']:.6f}"),
            ("ECE", f"{test['expected_calibration_error']:.6f}"),
            ("Brier score", f"{test['brier_score']:.6f}"),
        ],
        (3000, 6360),
        font_size=8.5,
    )
    base.add_table(
        doc,
        ("Class", "Precision", "Recall", "Specificity", "F1", "ROC AUC", "PR AUC", "Support"),
        [
            (
                name,
                f"{test['per_class'][name]['precision']:.4f}", f"{test['per_class'][name]['recall']:.4f}",
                f"{test['per_class'][name]['specificity']:.4f}", f"{test['per_class'][name]['f1']:.4f}",
                f"{test['per_class'][name]['roc_auc']:.4f}", f"{test['per_class'][name]['pr_auc']:.4f}", test['per_class'][name]['support'],
            )
            for name in ("Normal", "Viral", "Bacterial")
        ],
        (1500, 1100, 1000, 1500, 950, 1100, 1100, 1110),
        font_size=7.3,
    )
    base.add_body(doc, f"Frozen selection file: model/results/final_selection.json. Selected run: {selection['selected_run_id']}. Selection status after evaluation: {selection['status']}. No retraining or configuration change was permitted after observing test metrics.")
    add_figure(doc, RESULTS / "runs/final_locked_test/confusion_matrix.png", "Final locked-test confusion matrix.", width=5.4)


def xai_calibration_efficiency(doc, gradcam, calibration, efficiency) -> None:
    page_break(doc)
    base.add_heading(doc, "8. Explainability, Calibration, and Efficiency", 1)
    base.add_heading(doc, "8.1 Grad-CAM qualitative review", 2)
    base.add_body(doc, "Grad-CAM was generated only after selection and the single locked-test evaluation were complete. For each class, a high-confidence correct case and—where available—a high-confidence error case were selected from the frozen test predictions. The target is the model's predicted class at the final convolutional feature stage.")
    for record in gradcam["records"]:
        add_figure(
            doc,
            Path(record["output_path"]),
            f"Grad-CAM {record['case_type']}: truth {record['true_class']}, prediction {record['predicted_class']}, confidence {record['confidence']:.3f}. Qualitative localization only.",
            width=6.2,
        )
    base.add_keep_note(doc, gradcam["interpretation_boundary"])
    base.add_heading(doc, "8.2 Temperature scaling", 2)
    base.add_table(
        doc,
        ("Split/state", "Accuracy", "Macro F1", "Log loss", "ECE", "Brier"),
        [
            (label, f"{values['accuracy']:.4f}", f"{values['macro_f1']:.4f}", f"{values['log_loss']:.4f}", f"{values['expected_calibration_error']:.4f}", f"{values['brier_score']:.4f}")
            for label, values in (("Validation before", calibration["validation_before"]), ("Validation after", calibration["validation_after"]), ("Test before", calibration["test_before"]), ("Test after", calibration["test_after"]))
        ],
        (1800, 1300, 1300, 1500, 1300, 2160),
        font_size=8.0,
    )
    base.add_body(doc, f"Temperature = {calibration['temperature']:.6f}, fit on validation probabilities only. Temperature scaling does not change class ranking, so classification accuracy and macro F1 remain unchanged.")
    add_figure(doc, RESULTS / "calibration/test_calibration_temperature_scaling.png", "Locked-test reliability curves before and after validation-fit temperature scaling.", width=6.1)
    base.add_heading(doc, "8.3 Efficiency", 2)
    latency = efficiency["end_to_end_latency_ms"]
    base.add_table(
        doc,
        ("Engineering measure", "Value"),
        [
            ("Hardware", efficiency["hardware"]),
            ("Measured images", efficiency["measured_images"]),
            ("Mean end-to-end latency", f"{latency['mean']:.2f} ms/image"),
            ("Median / p95", f"{latency['median']:.2f} / {latency['p95']:.2f} ms"),
            ("Throughput from mean", f"{efficiency['end_to_end_throughput_images_per_second_from_mean']:.2f} images/s"),
            ("Parameters", f"{efficiency['parameter_count']:,}"),
            ("Checkpoint size", f"{efficiency['checkpoint_file_size_bytes'] / (1024**2):.1f} MB"),
            ("Peak training GPU memory", f"{efficiency['peak_training_gpu_memory_bytes'] / (1024**2):.1f} MB"),
        ],
        (3300, 6060),
        font_size=8.5,
    )
    base.add_keep_note(doc, efficiency["interpretation_boundary"])


def evidence_review(doc, evidence) -> None:
    page_break(doc)
    base.add_heading(doc, "9. Verified RRL and Research Evidence Matrix", 1)
    base.add_body(doc, "The supplied screenshot was treated as a supplementary reference index, not a source of truth. It displayed ten study rows and blank Reference Link cells. No original RRL spreadsheet or separate bibliography was present in the repository. Each visible study was located by exact title/authors/year and verified against a publisher, DOI, PubMed, official repository, or primary paper.")
    base.add_keep_note(doc, evidence["required_link_instruction"])
    screenshot = [row for row in evidence["records"] if row["Reference ID"] in SCREENSHOT_IDS]
    additional = [row for row in evidence["records"] if row["Reference ID"] not in SCREENSHOT_IDS]
    base.add_table(
        doc,
        ("ID", "Verified citation", "Category", "Result / merit", "Applicability"),
        [
            (
                row["Screenshot Journal ID"],
                f"{row['Authors']} ({row['Year']}). {row['Article Title']}.",
                row["Research Category"],
                f"{row['Reported Results']} {row['Claimed Merit']}",
                f"{row['Evidence Applicability']}: {row['Evidence Applicability Reason']}",
            )
            for row in screenshot
        ],
        (550, 3000, 1550, 2200, 2060),
        font_size=6.4,
    )
    base.add_heading(doc, "9.1 Detailed verification of screenshot entries", 2)
    for row in screenshot:
        base.add_heading(doc, f"{row['Screenshot Journal ID']} — {row['Article Title']}", 3)
        base.add_table(
            doc,
            ("Field", "Verified detail"),
            [
                ("Authors / year", f"{row['Authors']} | {row['Year']}"),
                ("Venue / DOI", f"{row['Venue']} | {row['DOI']}"),
                ("Objective", row["Research Objective"]),
                ("Methodology", row["Methodology"]),
                ("Task / dataset / model", f"{row['Task and Classes']} | {row['Dataset']} | {row['Architecture or Method']}"),
                ("Preprocessing / augmentation", f"{row['Preprocessing']} | {row['Augmentation']}"),
                ("Training details", f"Weights: {row['Pretrained Weights']}; optimizer: {row['Optimizer']}; LR: {row['Initial Learning Rate']}; WD: {row['Weight Decay']}; batch: {row['Batch Size']}; epochs: {row['Epochs']}; loss: {row['Loss']}"),
                ("Results", row["Reported Results"]),
                ("Merit", row["Claimed Merit"]),
                ("Limitations", row["Limitations"]),
                ("Applicability", f"{row['Evidence Applicability']} — {row['Evidence Applicability Reason']}"),
                ("Source trail", row["Source Obtained From"]),
                ("Corrections", row["Notes or Corrections"]),
            ],
            (2100, 7260),
            font_size=7.2,
        )
        add_link_line(doc, "Verified primary/official link", row["Official Link"])
    base.add_heading(doc, "9.2 Additional architecture, preprocessing, validation, and calibration sources", 2)
    base.add_table(
        doc,
        ("Reference ID", "Authors/year/title", "Use in this project", "Applicability", "DOI / official link"),
        [
            (
                row["Reference ID"],
                f"{row['Authors']} ({row['Year']}). {row['Article Title']}",
                row["Claimed Merit"],
                f"{row['Evidence Applicability']}: {row['Evidence Applicability Reason']}",
                f"{row['DOI']} | {row['Official Link']}",
            )
            for row in additional
        ],
        (1200, 2900, 1900, 1900, 1460),
        font_size=6.3,
    )
    base.add_heading(doc, "9.3 Transferability rules applied", 2)
    base.add_bullet(doc, "Exact three-class Normal/Viral/Bacterial CXR evidence received highest priority.")
    base.add_bullet(doc, "Binary CXR, localization, explainability, CT, multimodal, and general architecture evidence was never copied directly into classifier hyperparameters without an applicability explanation.")
    base.add_bullet(doc, "V-Net was excluded from configuration evidence because it is a 3D prostate MRI segmentation architecture.")
    base.add_bullet(doc, "Jaiswal/Mask R-CNN supports localization discussion but not classifier settings.")
    base.add_bullet(doc, "FaNet supports multimodal context only; CT and symptom fusion do not determine 2D CXR preprocessing.")
    base.add_bullet(doc, "Grad-CAM sources support explanation methods, not training configuration.")
    base.add_bullet(doc, "The verified pseudo-color paper did not use VIRIDIS; the two techniques were not conflated.")


def limitations_reproduction(doc, selection, failures) -> None:
    page_break(doc)
    base.add_heading(doc, "10. Limitations, Reproducibility, and Audit Trail", 1)
    for item in (
        "Single public pediatric dataset/source domain; no external hospital validation.",
        "Labels are dataset annotations, not a new multimodal clinical reference standard.",
        "Validation was used repeatedly for architecture, preprocessing, and tuning selection; final test remains the only held-out estimate.",
        "Group derivation is conservative and filename/hash-based because explicit patient IDs were unavailable.",
        "Center cropping can remove peripheral content from extreme aspect ratios.",
        "Horizontal flip moves laterality markers; rotation creates black corner fill.",
        "ImageNet normalization is retained even for enhanced intensity distributions.",
        "The preprocessing bootstrap interval is descriptive on the selection set.",
        "Grad-CAM is qualitative and can highlight non-causal or non-clinical regions.",
        "Local latency/VRAM results depend on this Windows/RTX 2050 environment.",
        "No YOLOv10 classifier, PCA, FILM, SAS tuning, or automated grading result was fabricated.",
    ):
        base.add_bullet(doc, item)
    base.add_heading(doc, "10.1 Environment", 2)
    environment = read_json(RESULTS / "environment.json")["multiclass_environment"]
    base.add_table(
        doc,
        ("Component", "Version/value"),
        [
            ("Python", environment["python_version"]),
            ("PyTorch", environment["packages"]["torch"]),
            ("Torchvision", environment["packages"]["torchvision"]),
            ("Ultralytics", environment["packages"]["ultralytics"]),
            ("CUDA / cuDNN", f"{environment['gpu']['torch_cuda_runtime']} / {environment['gpu']['cudnn_version']}"),
            ("GPU", environment["gpu"]["devices"][0]["name"]),
            ("VRAM", f"{environment['gpu']['devices'][0]['total_memory_bytes']:,} bytes"),
            ("Manifest hash", selection["manifest_hash"]),
            ("Final config hash", selection["configuration_hash"]),
            ("Repository tests", "12 passed before full continuation runs"),
        ],
        (2800, 6560),
        font_size=8.2,
    )
    base.add_heading(doc, "10.2 Resume and generation commands", 2)
    commands = [
        ".venv\\Scripts\\python.exe model\\scripts\\run_benchmark.py --stage preprocessing_ablation --resume",
        ".venv\\Scripts\\python.exe model\\scripts\\resolve_run_matrix.py --stage focused",
        ".venv\\Scripts\\python.exe model\\scripts\\resolve_run_matrix.py --stage tuning",
        ".venv\\Scripts\\python.exe model\\scripts\\run_tuning.py --matrix-run-id tuning_topN --parent-run-id PARENT --trials 4 --epochs 6 --resume",
        ".venv\\Scripts\\python.exe model\\scripts\\resolve_run_matrix.py --stage robustness",
        ".venv\\Scripts\\python.exe model\\scripts\\run_benchmark.py --stage robustness --config model\\configs\\final_selected.json --resume",
        ".venv\\Scripts\\python.exe model\\scripts\\resolve_run_matrix.py --stage test",
        "Use the single command recorded in model/results/final_selection.json for explicitly unlocked test evaluation.",
        ".venv\\Scripts\\python.exe model\\scripts\\generate_gradcam_review.py",
        ".venv\\Scripts\\python.exe model\\scripts\\calibrate_final_model.py",
        ".venv\\Scripts\\python.exe model\\scripts\\measure_final_efficiency.py",
    ]
    for command in commands:
        base.add_body(doc, command)
    base.add_heading(doc, "10.3 Preserved resolved failures", 2)
    base.add_table(
        doc,
        ("Historical row", "Resolved?", "Evidence"),
        [
            (row["run_id"], row["resolved"], f"{row['issue_id']}: {row['error_type']} — {row['error_message']}")
            for row in failures
        ],
        (2400, 1000, 5960),
        font_size=8.0,
    )
    base.add_body(doc, "The controlled resume interruption later completed epoch 2 from its checkpoint. The original YOLOv8 tuple-output failure was corrected by a classification-logits adapter and verified by completed retries. Original failure rows remain in the ledger with resolution identifiers.")
    base.add_keep_note(doc, "End of research-complete report. Machine-readable evidence remains authoritative for exact values; this document summarizes those persisted artifacts without inventing missing results.")


def build(output: Path) -> dict:
    failures = gate()
    if failures:
        raise RuntimeError("; ".join(failures))
    generated_at = datetime.now().astimezone().strftime("%B %d, %Y at %I:%M %p %Z")
    audit = read_json(RESULTS / "audits/current_state_audit_continuation_v2.json")
    split = read_json(MODEL / "data_manifests/split_summary.json")
    source = read_json(MODEL / "data_manifests/dataset_source_snapshot.json")
    phash_rows = read_csv(RESULTS / "perceptual_hash_manual_review/perceptual_hash_manual_review.csv")
    summary = read_csv(RESULTS / "summary.csv")
    baseline_rows = [row for row in summary if row["stage"] == "baseline" and row["status"] == "completed"]
    robustness = [row for row in summary if row["stage"] == "robustness" and row["status"] == "completed"]
    model_review = read_json(RESULTS / "audits/model_configuration_review_continuation_v2.json")
    ablation = read_json(RESULTS / "preprocessing_ablation_analysis.json")
    tuning = read_csv(RESULTS / "tuning_trials.csv")
    selection = read_json(RESULTS / "final_selection.json")
    test = read_json(RESULTS / "runs/final_locked_test/metrics.json")
    gradcam = read_json(RESULTS / "gradcam/gradcam_audit.json")
    calibration = read_json(RESULTS / "calibration/final_temperature_scaling.json")
    efficiency = read_json(RESULTS / "efficiency/final_efficiency.json")
    evidence = read_json(RESULTS / "evidence/research_evidence_matrix.json")
    failure_rows = read_csv(RESULTS / "failures.csv")

    doc = Document()
    base.configure_styles(doc)
    base.configure_sections(doc)
    doc.core_properties.title = "Pneumonia Model Training Details - Research Complete"
    doc.core_properties.subject = "Data preparation, preprocessing, augmentation, training, validation, locked-test results, explainability, and verified RRL"
    doc.core_properties.author = "Pneumonia Research Pipeline"
    doc.core_properties.keywords = "pneumonia, chest X-ray, DenseNet, ResNet, YOLO classification, CLAHE, Grad-CAM"
    cover(doc, generated_at, selection, test)
    executive(doc, audit, selection, test, robustness, calibration, efficiency)
    status_overview(doc, tuning, robustness, selection, gradcam)
    repository_and_data(doc, audit, split, source, phash_rows)
    preprocessing(doc)
    training(doc, model_review, baseline_rows)
    baseline_results(doc, baseline_rows)
    ablation_results(doc, ablation)
    tuning_results(doc, tuning)
    final_results(doc, robustness, selection, test)
    final_configuration_table(doc, model_review, selection)
    xai_calibration_efficiency(doc, gradcam, calibration, efficiency)
    evidence_review(doc, evidence)
    limitations_reproduction(doc, selection, failure_rows)

    saved = save_docx_with_fallback(doc, output)
    base.force_visible_colors_black(saved)
    structural = base.audit_docx(saved)
    structural.update(
        {
            "requested_output": str(output),
            "saved_output": str(saved),
            "research_complete_gate": "passed",
            "baseline_count": len(baseline_rows),
            "ablation_count": 9,
            "tuning_trial_rows": len(tuning),
            "robustness_count": len(robustness),
            "test_status": selection["status"],
            "reference_count": evidence["record_count"],
            "visual_render_verified": False,
        }
    )
    audit_path = saved.with_suffix(".audit.json")
    audit_path.write_text(json.dumps(structural, indent=2) + "\n", encoding="utf-8")
    if not structural["pass"]:
        raise RuntimeError(f"DOCX structural audit failed: {audit_path}")
    return structural


def parse_args():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--check-only", action="store_true")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    failures = gate()
    if args.check_only:
        print(json.dumps({"ready": not failures, "gate_failures": failures}, indent=2))
        return 1 if failures else 0
    if failures:
        print(json.dumps({"ready": False, "gate_failures": failures}, indent=2))
        return 1
    output = args.output if args.output.is_absolute() else ROOT / args.output
    structural = build(output.resolve())
    print(json.dumps(structural, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
