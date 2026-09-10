"""Generate the progressive Word report from live research sources.

The Markdown and machine-readable artifacts remain the sources of truth. This
DOCX is a regenerated presentation layer and never writes back to those files.

Design system: standard_business_brief preset with an editorial_cover opening.
Named cover override: 30 pt dark-blue title, 15 pt blue-gray subtitle, and a
10 pt gold kicker. Body/list/heading/page tokens match the selected preset.
"""

from __future__ import annotations

import argparse
import re
from datetime import datetime
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from document_io import save_docx_with_fallback


REPO_ROOT = Path(__file__).resolve().parents[2]
MODEL_ROOT = REPO_ROOT / "model"
DOCS_ROOT = MODEL_ROOT / "docs"
DEFAULT_OUTPUT = DOCS_ROOT / "Pneumonia_Research_Progress_Report.docx"

SOURCE_DOCUMENTS = [
    ("Current Status", DOCS_ROOT / "CURRENT_STATUS.md"),
    ("Existing Binary Pipeline Audit", DOCS_ROOT / "00_existing_binary_pipeline_audit.md"),
    ("Dataset Audit", DOCS_ROOT / "01_dataset_audit.md"),
    ("Experiment Plan", DOCS_ROOT / "02_experiment_plan.md"),
    ("Implementation Notes", DOCS_ROOT / "03_implementation_notes.md"),
    ("Testing and Validation", DOCS_ROOT / "04_testing_and_validation.md"),
    ("Results Interpretation", DOCS_ROOT / "05_results_interpretation.md"),
    ("Decision Log", DOCS_ROOT / "DECISION_LOG.md"),
    ("Issues and Fixes", DOCS_ROOT / "ISSUES_AND_FIXES.md"),
    ("Development Log", DOCS_ROOT / "DEVELOPMENT_LOG.md"),
    ("Training Journal", DOCS_ROOT / "TRAINING_JOURNAL.md"),
    ("Reproduction Guide", DOCS_ROOT / "06_reproduction_guide.md"),
]

LEGACY_FIGURES = [
    ("Legacy DenseNet121 validation confusion matrix", MODEL_ROOT / "outputs" / "densenet121" / "confusion_matrix.png"),
    ("Legacy DenseNet169 validation confusion matrix", MODEL_ROOT / "outputs" / "densenet169" / "confusion_matrix.png"),
    ("Legacy DenseNet201 validation confusion matrix", MODEL_ROOT / "outputs" / "densenet201" / "confusion_matrix.png"),
    ("Legacy DenseNet121 training history", MODEL_ROOT / "outputs" / "densenet121" / "training_history.png"),
    ("Legacy DenseNet169 training history", MODEL_ROOT / "outputs" / "densenet169" / "training_history.png"),
    ("Legacy DenseNet201 training history", MODEL_ROOT / "outputs" / "densenet201" / "training_history.png"),
]

DATASET_FIGURES = [
    (
        "Retained three-class image counts by locked split",
        MODEL_ROOT / "results" / "dataset_audit" / "class_counts.png",
    ),
    (
        "Retained image width, height, and aspect-ratio distributions",
        MODEL_ROOT / "results" / "dataset_audit" / "image_size_distributions.png",
    ),
]

SMOKE_FIGURES = [
    (
        "DenseNet121 end-to-end smoke training and validation curves",
        MODEL_ROOT / "results" / "runs" / "smoke_densenet121_e2e" / "training_curves.png",
    ),
    (
        "Controlled interruption/resume smoke training and validation curves",
        MODEL_ROOT / "results" / "runs" / "smoke_resume_contract" / "training_curves.png",
    ),
    (
        "Corrected YOLOv8 end-to-end smoke validation confusion matrix",
        MODEL_ROOT / "results" / "runs" / "smoke_yolov8_cls_e2e_retry2" / "confusion_matrix.png",
    ),
]

BASELINE_RUN_IDS = [
    "baseline_densenet121_s42",
    "baseline_densenet169_s42",
    "baseline_densenet201_s42",
    "baseline_resnet50_s42",
    "baseline_resnet101_s42",
    "baseline_resnet152_s42",
    "baseline_yolov5_cls_s42",
    "baseline_yolov8_cls_s42",
]

PREPROCESSING_FIGURES = [
    (
        "Validation-only visual comparison of implemented preprocessing paths",
        MODEL_ROOT / "results" / "preprocessing_examples" / "preprocessing_comparison.png",
    ),
]

BLUE = RGBColor(0x2E, 0x74, 0xB5)
DARK_BLUE = RGBColor(0x1F, 0x4D, 0x78)
COVER_BLUE = RGBColor(0x20, 0x37, 0x48)
SUBTITLE_BLUE = RGBColor(0x2B, 0x51, 0x63)
GOLD = RGBColor(0x7A, 0x5A, 0x00)
MUTED = RGBColor(0x66, 0x66, 0x66)
BLACK = RGBColor(0x00, 0x00, 0x00)


def set_run_font(
    run,
    *,
    name: str = "Calibri",
    size: float | None = None,
    color: RGBColor | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
) -> None:
    run.font.name = name
    if run._element.rPr is None:
        run._element.get_or_add_rPr()
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_repeatable_style_font(style, name: str, size: float, color: RGBColor, bold: bool = False) -> None:
    style.font.name = name
    style.font.size = Pt(size)
    style.font.color.rgb = color
    style.font.bold = bold
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)


def set_style_spacing(style, before: float, after: float, line: float) -> None:
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.line_spacing = line


def configure_styles(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    set_repeatable_style_font(normal, "Calibri", 11, BLACK)
    set_style_spacing(normal, 0, 6, 1.10)

    title = styles["Title"]
    set_repeatable_style_font(title, "Calibri", 30, COVER_BLUE, True)
    set_style_spacing(title, 0, 8, 1.0)

    subtitle = styles["Subtitle"]
    set_repeatable_style_font(subtitle, "Calibri", 15, SUBTITLE_BLUE, False)
    set_style_spacing(subtitle, 0, 8, 1.0)

    heading_tokens = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for style_name, (size, color, before, after) in heading_tokens.items():
        style = styles[style_name]
        set_repeatable_style_font(style, "Calibri", size, color, True)
        set_style_spacing(style, before, after, 1.0)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        set_repeatable_style_font(style, "Calibri", 11, BLACK)
        set_style_spacing(style, 0, 8, 1.167)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.tab_stops.add_tab_stop(Inches(0.5))

    if "Code Block" not in styles:
        code_style = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code_style = styles["Code Block"]
    set_repeatable_style_font(code_style, "Consolas", 8.5, RGBColor(0x22, 0x22, 0x22))
    set_style_spacing(code_style, 3, 6, 1.0)
    code_style.paragraph_format.left_indent = Inches(0.18)
    code_style.paragraph_format.right_indent = Inches(0.18)

    if "Caption" in styles:
        caption = styles["Caption"]
        set_repeatable_style_font(caption, "Calibri", 9, MUTED, False)
        set_style_spacing(caption, 3, 8, 1.0)
        caption.font.italic = True


def shade_paragraph(paragraph, fill: str) -> None:
    properties = paragraph._p.get_or_add_pPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def add_page_number(paragraph) -> None:
    paragraph.add_run("Page ")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    display = OxmlElement("w:t")
    display.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()
    run._r.extend([begin, instruction, separate, display, end])


def configure_page(document: Document) -> None:
    for section in document.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)
        section.different_first_page_header_footer = True

        header = section.header
        header_p = header.paragraphs[0]
        header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        header_p.paragraph_format.space_after = Pt(0)
        header_run = header_p.add_run("Pneumonia Three-Class Research | Progress Report")
        set_run_font(header_run, size=8.5, color=MUTED)

        footer = section.footer
        footer_p = footer.paragraphs[0]
        footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        footer_p.paragraph_format.space_before = Pt(0)
        add_page_number(footer_p)
        for run in footer_p.runs:
            set_run_font(run, size=8.5, color=MUTED)


def clean_text(text: str) -> str:
    replacements = {
        "\u2013": "-",
        "\u2014": "-",
        "\u2018": "'",
        "\u2019": "'",
        "\u201c": '"',
        "\u201d": '"',
        "\u2192": "->",
        "\u00d7": "x",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    return text


def add_inline_markdown(paragraph, text: str) -> None:
    text = clean_text(text)
    token_pattern = re.compile(r"(\*\*.*?\*\*|`.*?`|\*.*?\*)")
    position = 0
    for match in token_pattern.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position : match.start()])
            set_run_font(run)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, bold=True)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, name="Consolas", size=9.5, color=DARK_BLUE)
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, italic=True)
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        set_run_font(run)


def emit_paragraph(document: Document, lines: list[str]) -> None:
    if not lines:
        return
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.widow_control = True
    add_inline_markdown(paragraph, " ".join(line.strip() for line in lines))


def emit_markdown_table(document: Document, lines: list[str]) -> None:
    def cells(line: str) -> list[str]:
        return [item.strip() for item in line.strip().strip("|").split("|")]

    separator = re.compile(r"^:?-{3,}:?$")
    if len(lines) < 2 or not all(separator.match(item) for item in cells(lines[1])):
        emit_paragraph(document, lines)
        return
    rows = [cells(lines[0]), *[cells(line) for line in lines[2:]]]
    column_count = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=column_count)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for row_index, row in enumerate(rows):
        for column_index in range(column_count):
            cell = table.cell(row_index, column_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            value = row[column_index] if column_index < len(row) else ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            add_inline_markdown(paragraph, value)
            for run in paragraph.runs:
                set_run_font(run, size=8.0 if column_count > 6 else 8.5, bold=row_index == 0)
            if row_index == 0:
                properties = cell._tc.get_or_add_tcPr()
                shading = properties.find(qn("w:shd"))
                if shading is None:
                    shading = OxmlElement("w:shd")
                    properties.append(shading)
                shading.set(qn("w:fill"), "D9EAF7")
    after = document.add_paragraph()
    after.paragraph_format.space_after = Pt(2)


def render_markdown(document: Document, markdown: str, *, skip_first_h1: bool = True) -> None:
    lines = markdown.splitlines()
    paragraph_lines: list[str] = []
    in_code = False
    code_lines: list[str] = []
    table_lines: list[str] = []
    skipped_h1 = False

    def flush_paragraph() -> None:
        emit_paragraph(document, paragraph_lines)
        paragraph_lines.clear()

    def flush_table() -> None:
        if table_lines:
            emit_markdown_table(document, table_lines)
            table_lines.clear()

    for raw_line in lines:
        line = raw_line.rstrip()
        is_table_line = line.strip().startswith("|") and line.strip().endswith("|")
        if is_table_line and not in_code:
            flush_paragraph()
            table_lines.append(line)
            continue
        flush_table()
        if line.strip().startswith("```"):
            flush_paragraph()
            if in_code:
                code = document.add_paragraph(style="Code Block")
                code.add_run("\n".join(code_lines))
                shade_paragraph(code, "F2F4F7")
                code_lines.clear()
                in_code = False
            else:
                in_code = True
            continue
        if in_code:
            code_lines.append(clean_text(line))
            continue
        if not line.strip():
            flush_paragraph()
            continue

        heading = re.match(r"^(#{1,3})\s+(.*)$", line)
        if heading:
            flush_paragraph()
            level = len(heading.group(1))
            if level == 1 and skip_first_h1 and not skipped_h1:
                skipped_h1 = True
                continue
            mapped_level = min(level, 3)
            paragraph = document.add_paragraph(style=f"Heading {mapped_level}")
            add_inline_markdown(paragraph, heading.group(2))
            continue

        bullet = re.match(r"^\s*-\s+(.*)$", line)
        if bullet:
            flush_paragraph()
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline_markdown(paragraph, bullet.group(1))
            continue

        numbered = re.match(r"^\s*\d+\.\s+(.*)$", line)
        if numbered:
            flush_paragraph()
            paragraph = document.add_paragraph(style="List Number")
            add_inline_markdown(paragraph, numbered.group(1))
            continue

        if line.strip() == "---":
            flush_paragraph()
            continue
        paragraph_lines.append(line)

    flush_paragraph()
    flush_table()
    if code_lines:
        code = document.add_paragraph(style="Code Block")
        code.add_run("\n".join(code_lines))
        shade_paragraph(code, "F2F4F7")


def add_cover(document: Document, generated_at: str, phase_label: str) -> None:
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_before = Pt(112)
    spacer.paragraph_format.space_after = Pt(0)

    kicker = document.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(18)
    run = kicker.add_run("RESEARCH PROGRESS REPORT")
    set_run_font(run, size=10, color=GOLD, bold=True)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Pneumonia Three-Class\nImage Classification")

    subtitle = document.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Normal | Viral Pneumonia | Bacterial Pneumonia")

    phase = document.add_paragraph()
    phase.alignment = WD_ALIGN_PARAGRAPH.CENTER
    phase.paragraph_format.space_before = Pt(22)
    phase.paragraph_format.space_after = Pt(4)
    run = phase.add_run(phase_label)
    set_run_font(run, size=11, color=COVER_BLUE, bold=True)

    generated = document.add_paragraph()
    generated.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = generated.add_run(f"Generated {generated_at}")
    set_run_font(run, size=9.5, color=MUTED, italic=True)

    disclaimer = document.add_paragraph()
    disclaimer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    disclaimer.paragraph_format.space_before = Pt(70)
    run = disclaimer.add_run(
        "Research use only. Not validated or approved for clinical diagnosis or deployment."
    )
    set_run_font(run, size=9, color=RGBColor(0x9B, 0x1C, 0x1C), bold=True)
    document.add_page_break()


def add_contents(document: Document) -> None:
    document.add_paragraph("Contents", style="Heading 1")
    intro = document.add_paragraph()
    intro.add_run(
        "This regenerated report presents the live audit, status, decisions, issues, testing, "
        "training journal, and reproduction sources in the following order."
    )
    for title, _ in SOURCE_DOCUMENTS:
        paragraph = document.add_paragraph(style="List Number")
        paragraph.add_run(title)
    paragraph = document.add_paragraph(style="List Number")
    paragraph.add_run("Dataset Audit Figures")
    paragraph = document.add_paragraph(style="List Number")
    paragraph.add_run("Preprocessing Evidence Figures")
    paragraph = document.add_paragraph(style="List Number")
    paragraph.add_run("Implementation Smoke Figures")
    paragraph = document.add_paragraph(style="List Number")
    paragraph.add_run("Controlled Baseline Figures")
    paragraph = document.add_paragraph(style="List Number")
    paragraph.add_run("Legacy Result Figures (historical diagnostics only)")


def add_source_section(document: Document, title: str, path: Path) -> None:
    heading = document.add_paragraph(title, style="Heading 1")
    heading.paragraph_format.page_break_before = True
    source = document.add_paragraph()
    source.paragraph_format.space_after = Pt(8)
    run = source.add_run(f"Source artifact: {path.relative_to(REPO_ROOT).as_posix()}")
    set_run_font(run, size=8.5, color=MUTED, italic=True)
    if path.exists():
        render_markdown(document, path.read_text(encoding="utf-8", errors="replace"))
    else:
        paragraph = document.add_paragraph()
        run = paragraph.add_run("Source file is not yet available.")
        set_run_font(run, color=RGBColor(0x9B, 0x1C, 0x1C), bold=True)


def add_dataset_figures(document: Document) -> None:
    heading = document.add_paragraph("Dataset Audit Figures", style="Heading 1")
    heading.paragraph_format.page_break_before = True
    note = document.add_paragraph()
    note.add_run(
        "These figures are generated from the retained canonical image audit and locked "
        "manifest version three_class_seed42_3b815cf769c8."
    )

    for index, (caption_text, path) in enumerate(DATASET_FIGURES, start=1):
        if not path.exists():
            continue
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.keep_with_next = True
        paragraph.add_run().add_picture(str(path), width=Inches(6.25))
        caption = document.add_paragraph(style="Caption")
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.add_run(
            f"Figure {index}. {caption_text}. Source: {path.relative_to(REPO_ROOT).as_posix()}"
        )


def add_legacy_figures(document: Document) -> None:
    heading = document.add_paragraph("Legacy Result Figures", style="Heading 1")
    heading.paragraph_format.page_break_before = True
    note = document.add_paragraph()
    run = note.add_run(
        "These figures are retained historical diagnostics from the incomplete generated binary "
        "dataset. They are not three-class benchmark or locked-test evidence."
    )
    set_run_font(run, color=RGBColor(0x9B, 0x1C, 0x1C), bold=True)

    for index, (caption_text, path) in enumerate(
        LEGACY_FIGURES,
        start=(
            len(DATASET_FIGURES)
            + len(PREPROCESSING_FIGURES)
            + len(SMOKE_FIGURES)
            + len(completed_baseline_figures())
            + 1
        ),
    ):
        if not path.exists():
            continue
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.keep_with_next = True
        width = Inches(5.8 if "confusion" in caption_text.lower() else 6.25)
        paragraph.add_run().add_picture(str(path), width=width)
        caption = document.add_paragraph(style="Caption")
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.add_run(f"Figure {index}. {caption_text}. Source: {path.relative_to(REPO_ROOT).as_posix()}")


def add_smoke_figures(document: Document) -> None:
    heading = document.add_paragraph("Implementation Smoke Figures", style="Heading 1")
    heading.paragraph_format.page_break_before = True
    note = document.add_paragraph()
    run = note.add_run(
        "These figures validate code paths on tiny balanced subsets with random initialization. "
        "They are not architecture-comparison or model-selection evidence."
    )
    set_run_font(run, color=RGBColor(0x9B, 0x1C, 0x1C), bold=True)
    for index, (caption_text, path) in enumerate(
        SMOKE_FIGURES, start=len(DATASET_FIGURES) + len(PREPROCESSING_FIGURES) + 1
    ):
        if not path.exists():
            continue
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.keep_with_next = True
        width = Inches(6.25 if "curves" in caption_text.lower() else 5.8)
        paragraph.add_run().add_picture(str(path), width=width)
        caption = document.add_paragraph(style="Caption")
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.add_run(
            f"Figure {index}. {caption_text}. Source: {path.relative_to(REPO_ROOT).as_posix()}"
        )


def completed_baseline_figures() -> list[tuple[str, Path]]:
    figures: list[tuple[str, Path]] = []
    for run_id in BASELINE_RUN_IDS:
        run_dir = MODEL_ROOT / "results" / "runs" / run_id
        status_path = run_dir / "status.json"
        if not status_path.exists() or '"status": "completed"' not in status_path.read_text(
            encoding="utf-8", errors="replace"
        ):
            continue
        display_name = run_id.removeprefix("baseline_").removesuffix("_s42")
        figures.extend(
            [
                (
                    f"Controlled baseline {display_name} training and validation curves",
                    run_dir / "training_curves.png",
                ),
                (
                    f"Controlled baseline {display_name} validation confusion matrix",
                    run_dir / "confusion_matrix.png",
                ),
            ]
        )
    return figures


def add_baseline_figures(document: Document) -> None:
    figures = completed_baseline_figures()
    heading = document.add_paragraph("Controlled Baseline Figures", style="Heading 1")
    heading.paragraph_format.page_break_before = True
    note = document.add_paragraph()
    note.add_run(
        "These figures come from completed Stage B runs on the locked training and validation "
        "manifests. They are model-selection evidence, but they are not locked-test results."
    )
    if not figures:
        paragraph = document.add_paragraph()
        run = paragraph.add_run("No controlled baseline had completed when this report was generated.")
        set_run_font(run, color=MUTED, italic=True)
        return
    start = len(DATASET_FIGURES) + len(PREPROCESSING_FIGURES) + len(SMOKE_FIGURES) + 1
    for index, (caption_text, path) in enumerate(figures, start=start):
        if not path.exists():
            continue
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.keep_with_next = True
        width = Inches(6.25 if "curves" in caption_text.lower() else 5.8)
        paragraph.add_run().add_picture(str(path), width=width)
        caption = document.add_paragraph(style="Caption")
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.add_run(
            f"Figure {index}. {caption_text}. Source: {path.relative_to(REPO_ROOT).as_posix()}"
        )


def add_preprocessing_figures(document: Document) -> None:
    heading = document.add_paragraph("Preprocessing Evidence Figures", style="Heading 1")
    heading.paragraph_format.page_break_before = True
    note = document.add_paragraph()
    note.add_run(
        "Examples are selected only from the assigned validation manifest for visual QA. "
        "The implemented methods are fixed transforms and no parameters are fitted on validation data."
    )
    for index, (caption_text, path) in enumerate(
        PREPROCESSING_FIGURES, start=len(DATASET_FIGURES) + 1
    ):
        if not path.exists():
            continue
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.keep_with_next = True
        paragraph.add_run().add_picture(str(path), width=Inches(6.25))
        caption = document.add_paragraph(style="Caption")
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.add_run(
            f"Figure {index}. {caption_text}. Source: {path.relative_to(REPO_ROOT).as_posix()}"
        )


def build_report(output_path: Path, phase_label: str) -> None:
    generated_at = datetime.now().astimezone().isoformat(timespec="seconds")
    document = Document()
    document.core_properties.title = "Pneumonia Three-Class Classification Research Progress Report"
    document.core_properties.subject = phase_label
    document.core_properties.author = "Pneumonia Research Pipeline"
    document.core_properties.comments = "Generated from live Markdown, JSON, CSV, and legacy result artifacts."

    configure_styles(document)
    configure_page(document)
    add_cover(document, generated_at, phase_label)
    add_contents(document)
    for title, path in SOURCE_DOCUMENTS:
        add_source_section(document, title, path)
    add_dataset_figures(document)
    add_preprocessing_figures(document)
    add_smoke_figures(document)
    add_baseline_figures(document)
    add_legacy_figures(document)

    saved_path = save_docx_with_fallback(document, output_path)
    print(f"Wrote progress report: {saved_path}")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument(
        "--phase-label",
        default="Phase 5 complete - preprocessing evidence and controlled run matrix",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    output = args.output if args.output.is_absolute() else REPO_ROOT / args.output
    build_report(output.resolve(), args.phase_label)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
