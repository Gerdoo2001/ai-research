"""Create a monochrome presenter guide for the pneumonia research documents."""

from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


REPO = Path(__file__).resolve().parents[2]
OUTPUT = (
    REPO
    / "model"
    / "docs"
    / "Pneumonia_Classification_Discussion_Notes_and_Overview_20260813.docx"
)

BLACK = RGBColor(0, 0, 0)
DARK_GRAY = RGBColor(68, 68, 68)
MUTED = RGBColor(95, 95, 95)
LIGHT_GRAY = "E7E7E7"
VERY_LIGHT_GRAY = "F4F4F4"
WHITE = "FFFFFF"
TABLE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


def load_json(path: Path) -> dict[str, Any]:
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return {}


def pct(value: Any, digits: int = 3) -> str:
    if value is None:
        return "Pending"
    return f"{float(value) * 100:.{digits}f}%"


def set_run_font(
    run,
    *,
    name: str = "Arial",
    size: float | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
    color: RGBColor = BLACK,
) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = color


def set_style_font(style, size: float, *, bold: bool = False, color: RGBColor = BLACK) -> None:
    style.font.name = "Arial"
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Arial")
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = color


def configure_styles(doc: Document) -> None:
    # Compact reference guide preset with one named override:
    # monochrome_academic = all headings black; table fills grayscale only.
    normal = doc.styles["Normal"]
    set_style_font(normal, 11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    title = doc.styles["Title"]
    set_style_font(title, 26, bold=True)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(8)
    title_ppr = title._element.get_or_add_pPr()
    for border in list(title_ppr.findall(qn("w:pBdr"))):
        title_ppr.remove(border)

    subtitle = doc.styles["Subtitle"]
    set_style_font(subtitle, 13, color=DARK_GRAY)
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(18)

    heading_specs = {
        "Heading 1": (16, 18, 10),
        "Heading 2": (13, 14, 7),
        "Heading 3": (12, 10, 5),
    }
    for name, (size, before, after) in heading_specs.items():
        style = doc.styles[name]
        set_style_font(style, size, bold=True, color=BLACK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        set_style_font(style, 11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    caption = doc.styles["Caption"]
    set_style_font(caption, 9.5, bold=True, color=DARK_GRAY)
    caption.paragraph_format.space_before = Pt(5)
    caption.paragraph_format.space_after = Pt(4)


def add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_run_font(run, size=9, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Arial")
    fonts.set(qn("w:hAnsi"), "Arial")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "18")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "666666")
    rpr.extend([fonts, size, color])
    r.append(rpr)
    text = OxmlElement("w:t")
    text.text = "1"
    r.append(text)
    fld.append(r)
    paragraph._p.append(fld)


def configure_sections(doc: Document) -> None:
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)

        hp = section.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        hp.paragraph_format.space_after = Pt(0)
        hr = hp.add_run("PNEUMONIA CLASSIFICATION | DISCUSSION GUIDE")
        set_run_font(hr, size=8.5, bold=True, color=MUTED)
        ppr = hp._p.get_or_add_pPr()
        pbdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "4")
        bottom.set(qn("w:space"), "4")
        bottom.set(qn("w:color"), "B5B5B5")
        pbdr.append(bottom)
        ppr.append(pbdr)

        add_page_field(section.footer.paragraphs[0])


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_row_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    tr_pr.append(cant)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]) -> None:
    if sum(widths) != TABLE_WIDTH_DXA:
        raise ValueError(f"table widths must sum to {TABLE_WIDTH_DXA}, got {sum(widths)}")
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(TABLE_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        set_row_cant_split(row)
        for idx, (cell, width) in enumerate(zip(row.cells, widths)):
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def set_cell_text(cell, text: str, *, bold: bool = False, size: float = 9.2) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold)


def add_table(
    doc: Document,
    headers: list[str],
    rows: Iterable[Iterable[str]],
    widths: list[int],
    *,
    font_size: float = 9.2,
) -> Any:
    data = [list(row) for row in rows]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, bold=True, size=font_size)
        shade_cell(table.rows[0].cells[i], LIGHT_GRAY)
    set_repeat_table_header(table.rows[0])
    for row_idx, row_data in enumerate(data, start=1):
        cells = table.add_row().cells
        for i, value in enumerate(row_data):
            set_cell_text(cells[i], str(value), size=font_size)
            shade_cell(cells[i], WHITE if row_idx % 2 else VERY_LIGHT_GRAY)
    set_table_geometry(table, widths)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    return table


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True


def add_para(doc: Document, text: str, *, bold_lead: str | None = None, italic: bool = False) -> None:
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, size=11, bold=True)
        r2 = p.add_run(text[len(bold_lead) :])
        set_run_font(r2, size=11, italic=italic)
    else:
        r = p.add_run(text)
        set_run_font(r, size=11, italic=italic)


def add_bullets(doc: Document, items: Iterable[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(item)
        set_run_font(r, size=11)


def new_decimal_numbering(doc: Document) -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [
        int(node.get(qn("w:abstractNumId")))
        for node in numbering.findall(qn("w:abstractNum"))
        if node.get(qn("w:abstractNumId")) is not None
    ]
    num_ids = [
        int(node.get(qn("w:numId")))
        for node in numbering.findall(qn("w:num"))
        if node.get(qn("w:numId")) is not None
    ]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal")
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), "%1.")
    level_jc = OxmlElement("w:lvlJc")
    level_jc.set(qn("w:val"), "left")
    ppr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "270")
    tabs.append(tab)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "540")
    indent.set(qn("w:hanging"), "270")
    ppr.extend([tabs, indent])
    rpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Arial")
    fonts.set(qn("w:hAnsi"), "Arial")
    rpr.append(fonts)
    level.extend([start, num_fmt, level_text, level_jc, ppr, rpr])
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def add_numbers(doc: Document, items: Iterable[str]) -> None:
    num_id = new_decimal_numbering(doc)
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.25
        ppr = p._p.get_or_add_pPr()
        num_pr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), "0")
        num_id_node = OxmlElement("w:numId")
        num_id_node.set(qn("w:val"), str(num_id))
        num_pr.extend([ilvl, num_id_node])
        ppr.append(num_pr)
        r = p.add_run(item)
        set_run_font(r, size=11)


def add_callout(doc: Document, label: str, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_geometry(table, [TABLE_WIDTH_DXA])
    cell = table.cell(0, 0)
    shade_cell(cell, VERY_LIGHT_GRAY)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r1 = p.add_run(label + ": ")
    set_run_font(r1, size=10.5, bold=True)
    r2 = p.add_run(text)
    set_run_font(r2, size=10.5)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def page_break(doc: Document) -> None:
    doc.add_page_break()


def build_document() -> Path:
    ft010_progress = load_json(
        REPO / "model" / "results" / "ft010_label_smoothing_multiseed" / "progress.json"
    )
    ft010_summary = load_json(
        REPO / "model" / "results" / "ft010_label_smoothing_multiseed" / "summary.json"
    )
    final_decision = load_json(
        REPO / "model" / "results" / "final_ft004_ensemble" / "decision.json"
    )
    paired_bootstrap = load_json(
        REPO / "model" / "results" / "final_ft004_ensemble" / "paired_bootstrap.json"
    )
    ft009b = load_json(REPO / "model" / "results" / "ft009b_hierarchical_s42" / "metrics.json")
    ft009b_stage1 = load_json(
        REPO
        / "model"
        / "results"
        / "runs"
        / "fine_tuning_ft009b_stage1_normal_pneumonia_s42"
        / "metrics.json"
    )
    ft009b_stage2 = load_json(
        REPO
        / "model"
        / "results"
        / "runs"
        / "fine_tuning_ft009b_stage2_viral_bacterial_s42"
        / "metrics.json"
    )

    now = datetime.now().astimezone()
    status_line = "FT-010 status unavailable"
    if ft010_progress.get("status") == "completed" and ft010_summary:
        status_line = (
            f"FT-010 completed 3/3 seeds; mean macro F1 "
            f"{pct(ft010_summary.get('mean_macro_f1'))} and sample standard deviation "
            f"{pct(ft010_summary.get('sample_std_macro_f1'))}."
        )

    doc = Document()
    configure_styles(doc)

    # Page 1 - cover and executive overview.
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(42)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("RESEARCH DISCUSSION NOTES")
    set_run_font(r, size=10, bold=True, color=MUTED)
    title = doc.add_paragraph("Pneumonia Classification", style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    subtitle = doc.add_paragraph(
        "Presenter's overview of the full journal and post-selection fine-tuning experiments",
        style="Subtitle",
    )
    subtitle.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_table(
        doc,
        ["Discussion field", "Current value"],
        [
            ["Task", "Classify pediatric chest X-rays as Normal, Viral Pneumonia, or Bacterial Pneumonia"],
            ["Development-selected pipeline", "DenseNet-201 + VIRIDIS -> GE -> CLAHE"],
            ["Chosen final research model", "FT-004 three-member DenseNet-201 ensemble; validation macro F1 85.305% across three seeds"],
            ["Best simpler alternative", "FT-010 label smoothing 0.05; mean validation macro F1 85.029% across three seeds"],
            ["Finalization status", status_line],
            ["Clinical boundary", "Research benchmark only; not a medical device or clinical diagnosis system"],
        ],
        [2300, 7060],
        font_size=9.5,
    )
    add_callout(
        doc,
        "Main message",
        "The study is strongest as a controlled, leakage-aware comparison. The central unresolved problem is not detecting pneumonia versus normal; it is separating Viral from Bacterial pneumonia reliably and externally.",
    )
    add_heading(doc, "One-minute overview", 1)
    add_para(
        doc,
        "We created a leakage-controlled pediatric chest-X-ray classifier for Normal, Viral, and Bacterial cases. After auditing duplicates and locking group-aware splits, we compared eight pretrained architectures and every ordered VIRIDIS, GE, and CLAHE combination for the leading models. Fine-tuning tested resolution, augmentation, ensembles, alternative losses, lung segmentation, and hierarchical classification. FT-010 completed all three label-smoothing seeds, but its 85.029% mean macro F1 remained below FT-004's 85.305%. FT-004 is therefore frozen as the final validation-selected research model: three DenseNet-201 members apply complementary preprocessing and their three-class probabilities are averaged equally.",
    )

    # Page 2 - document map and speaking structure.
    page_break(doc)
    add_heading(doc, "1. How the three source documents fit together", 1)
    add_table(
        doc,
        ["Document", "Role", "Important cutoff"],
        [
            [
                "Pneumonia_Three_Class_Classification_Full_Journal_20260810.docx",
                "Main journal: data preparation, eight-model baselines, preprocessing ablations, tuning, three-seed selection, earlier locked-test result, explainability, and references.",
                "Updated 10 Aug 2026; it predates FT-006 through FT-010 outcomes.",
            ],
            [
                "DenseNet201_VIRIDIS_GE_CLAHE_Fine_Tuning_Report.docx",
                "Reader-ready post-selection report covering FT-001 through FT-009B design and progress.",
                "Updated 13 Aug 2026; FT-009B, FT-010, and final model selection are now complete.",
            ],
            [
                "DenseNet201_VIRIDIS_GE_CLAHE_Fine_Tuning_Journal.md",
                "Detailed audit journal with configurations, rationale, references, gates, artifact paths, and chronological evidence.",
                "Updated through FT-010 and the frozen FT-004 selection; preserves the chronological audit trail.",
            ],
            [
                "This discussion guide",
                "Bridges the source documents and incorporates FT-009B, completed FT-010, paired bootstrap findings, and the frozen FT-004 decision.",
                f"Snapshot created {now.strftime('%d %B %Y, %H:%M %Z')}.",
            ],
        ],
        [2650, 4350, 2360],
        font_size=8.6,
    )
    add_heading(doc, "Recommended order when discussing", 2)
    add_bullets(
        doc,
        [
            "Start with the research problem and the three classes.",
            "Explain the data audit and leakage controls before discussing scores.",
            "Explain the shared training configuration and why macro F1 was primary.",
            "Show how the ordered preprocessing winner was selected across multiple seeds.",
            "Walk through fine-tuning as controlled questions, not as random trial-and-error.",
            "End with the selection boundary: validation evidence is strong, but independent external validation is still required.",
        ],
    )
    add_callout(
        doc,
        "Do not say",
        "'The model is clinically proven' or 'the model achieved the old locked-test score.' The old locked-test result belongs only to the earlier DenseNet-201 + GE model.",
    )

    # Page 3 - data and pipeline.
    page_break(doc)
    add_heading(doc, "2. Study design and data preparation", 1)
    add_para(
        doc,
        "The task uses the public Kermany pediatric chest-X-ray collection. The cohort consists of children approximately one to five years old. Individual filenames do not provide a reliable exact age, so age must be described at the cohort level rather than inferred from image appearance.",
    )
    add_table(
        doc,
        ["Preparation stage", "Verified result", "Why it matters"],
        [
            ["Source audit", "17,591 physical source files inspected", "The source contained nested copies, metadata, and redundant material that could not be treated as independent samples."],
            ["Canonical set", "5,824 unique canonical images", "One canonical record represents the unique image identity after exact-hash and source-layout checks."],
            ["Redundant exact images", "32 excluded", "Prevents identical image content from inflating sample size or crossing splits."],
            ["Training split", "4,076 images", "Used for weight updates and training-only augmentation."],
            ["Validation split", "874 images", "Used for model selection, checkpoint selection, and fine-tuning comparisons."],
            ["Sealed test split", "874 images", "Originally reserved for one final evaluation; it was opened once for the earlier DenseNet-201 + GE model."],
            ["Leakage checks", "Zero assigned exact-hash or patient/group leakage", "Reduces optimistic results caused by the same image or related patient group appearing in multiple assigned splits."],
        ],
        [2200, 2300, 4860],
        font_size=8.8,
    )
    add_heading(doc, "Data flow to the model", 2)
    add_para(
        doc,
        "Locked manifest -> open original image -> apply the selected enhancement order -> training-only flip/rotation -> aspect-ratio-preserving resize and center crop to 224 x 224 -> float32 scaling -> channel-first tensor -> ImageNet normalization -> DenseNet-201.",
    )
    add_heading(doc, "Augmentation boundary", 2)
    add_bullets(
        doc,
        [
            "Training only: random horizontal flip with probability 0.5.",
            "Training only: random rotation within +/-5 degrees.",
            "Validation and test: deterministic processing only; no random augmentation.",
            "Excluded: vertical flips, large rotations, and aggressive crops because they can create anatomically implausible chest images or remove disease evidence.",
        ],
    )

    # Page 4 - preprocessing and configuration.
    page_break(doc)
    add_heading(doc, "3. What VIRIDIS, GE, and CLAHE do", 1)
    add_table(
        doc,
        ["Step", "Plain-language meaning", "Reason for inclusion", "Important limitation"],
        [
            ["VIRIDIS", "Maps grayscale intensity values to an artificial RGB color scale.", "Tests whether a different intensity representation interacts favorably with ImageNet-pretrained filters.", "It adds no new clinical information and must not be described as revealing disease that was absent in the X-ray."],
            ["GE", "Global Histogram Equalization redistributes brightness values across the whole image.", "Can make broad low-contrast structures easier for the network to distinguish.", "Global enhancement may overemphasize noise or flatten local differences."],
            ["CLAHE", "Enhances contrast within small local tiles while limiting extreme amplification.", "Can emphasize subtle regional lung patterns without allowing unlimited contrast growth.", "Its benefit is empirical and architecture-dependent."],
        ],
        [1200, 2450, 2850, 2860],
        font_size=8.6,
    )
    add_callout(
        doc,
        "Why order matters",
        "These operations are nonlinear. A -> B does not generally equal B -> A. Each step consumes pixels already changed by the previous step, so every order can produce a different input distribution.",
    )
    add_heading(doc, "Shared DenseNet-201 training configuration", 2)
    add_table(
        doc,
        ["Setting", "Controlled value", "Meaning"],
        [
            ["Weights", "ImageNet pretrained", "Start from learned visual features instead of random weights."],
            ["Input", "224 x 224, three channels", "Matches the standard pretrained input contract and the resource-tested baseline."],
            ["Batch size", "8", "Largest stable controlled choice for DenseNet-201 on the RTX 2050 under the tested setup."],
            ["Optimizer", "AdamW", "Updates parameters using adaptive gradients while separating weight decay."],
            ["Learning rate", "0.0001", "Small step size appropriate for fine-tuning pretrained weights."],
            ["Weight decay", "0.0001", "Regularization discouraging excessively large weights."],
            ["Epoch limit", "12", "Maximum passes through training data; early stopping can finish sooner."],
            ["Early stopping", "Patience 4", "Stop when validation macro F1 fails to improve sufficiently for four epochs."],
            ["Class weighting", "Balanced", "Gives minority classes more influence during loss calculation."],
            ["Primary metric", "Validation macro F1", "Gives Normal, Viral, and Bacterial classes equal weight."],
        ],
        [1900, 2100, 5360],
        font_size=8.7,
    )

    # Page 5 - pre-fine-tuning selection.
    page_break(doc)
    add_heading(doc, "4. What was established before fine-tuning", 1)
    add_bullets(
        doc,
        [
            "Eight supported ImageNet-initialized classifiers completed a controlled baseline comparison.",
            "DenseNet-201 and ResNet-50 each completed all 16 possible ordered subsets of CLAHE, GE, and VIRIDIS, including the no-enhancement baseline.",
            "The leading pipeline was transferred to the other supported architectures rather than assumed to work universally.",
            "Finalists were repeated with seeds 42, 123, and 2026 to measure repeatability.",
        ],
    )
    add_table(
        doc,
        ["Candidate", "Mean macro F1", "Mean balanced accuracy", "Mean accuracy", "Interpretation"],
        [
            ["DenseNet-201 + VIRIDIS -> GE -> CLAHE", "84.168%", "83.932%", "85.202%", "Selected validation pipeline under the declared primary metric."],
            ["DenseNet-169 + VIRIDIS -> GE -> CLAHE", "84.142%", "83.804%", "85.088%", "Practically tied, more stable, and slightly better mean Viral recall."],
            ["ResNet-50 + CLAHE -> VIRIDIS -> GE", "83.472%", "83.207%", "84.477%", "Best architecture-specific ResNet-50 order."],
        ],
        [2800, 1350, 1550, 1350, 2310],
        font_size=8.7,
    )
    add_para(
        doc,
        "The DenseNet-201 advantage over DenseNet-169 was only 0.026 percentage points in mean macro F1. Therefore, the correct wording is 'selected validation candidate,' not 'proven substantially superior.'",
        bold_lead="The DenseNet-201 advantage",
    )
    add_heading(doc, "Why macro F1 is primary", 2)
    add_para(
        doc,
        "Accuracy counts every correct image equally and is influenced by the largest class. Macro F1 calculates an F1 score separately for Normal, Viral, and Bacterial and then averages them equally. It prevents strong Normal or Bacterial performance from hiding weak Viral performance. Balanced accuracy complements it by averaging the three recalls.",
    )
    add_heading(doc, "What a seed means", 2)
    add_para(
        doc,
        "A seed is a repeatable starting number for pseudo-random operations such as classifier initialization, training-batch order, flips, and rotations. Seeds 42, 123, and 2026 are simply three predetermined random conditions. A seed does not represent images, epochs, a model version, or a percentage.",
    )

    # Page 6 - fine-tuning FT-001 to FT-005.
    page_break(doc)
    add_heading(doc, "5. Fine-tuning experiments FT-001 through FT-005", 1)
    add_table(
        doc,
        ["Experiment", "Question", "Key evidence", "Decision and why"],
        [
            ["FT-001: 320 x 320", "Would more pixels preserve subtle detail?", "Macro F1 84.215%; accuracy 85.469%.", "Stopped. It used more computation but remained below the 224 x 224 reference macro F1 of 85.294%."],
            ["FT-002/2B/2C: PSPNet ROI audit", "Could a pretrained lung mask remove irrelevant borders and anatomy?", "Initial manual acceptance 69/90; recalibrated independent audit 82/90, with class-dependent failures.", "Did not authorize classification. The adult/general CXR segmenter was not reliable enough across the pediatric disease classes."],
            ["FT-003: mild affine augmentation", "Would small translation/scale variation reduce overfitting?", "Macro F1 83.605%; Viral recall 65.471%.", "Stopped. Viral recall improved, but overall macro F1 declined by 1.688 points."],
            ["FT-004: equal-probability ensemble", "Would complementary preprocessing models correct different errors?", "Three-seed mean macro F1 85.305%; accuracy 86.270%; Viral recall 66.517%.", "Advanced as the strongest confirmed average-performing system, with the cost of running three models per image."],
            ["FT-005: full PSPNet ROI feasibility", "Could the frozen ROI policy process all 4,950 development images safely?", "116 unusable crops, all in Viral or Bacterial cases.", "Stopped before classifier training. A class-dependent preprocessing failure would create an unfair or incomplete comparison."],
        ],
        [1500, 2450, 2300, 3110],
        font_size=8.3,
    )
    add_callout(
        doc,
        "Segmentation lesson",
        "An automatic geometric pass is not the same as an anatomically correct mask. ROI methods were required to pass manual review before they could influence classifier training.",
    )
    add_heading(doc, "How to describe a negative experiment", 2)
    add_para(
        doc,
        "A stopped experiment is still a completed scientific result when the stopping rule was defined in advance. FT-005 did not 'fail to produce an accuracy'; it correctly stopped because the preprocessing itself was not feasible for every class.",
    )

    # Page 7 - fine-tuning FT-006 to FT-010.
    page_break(doc)
    add_heading(doc, "6. Fine-tuning experiments FT-006 through FT-010", 1)
    add_table(
        doc,
        ["Experiment", "Controlled change", "Verified result", "Interpretation"],
        [
            ["FT-006: focal loss", "Cross-entropy -> balanced focal loss, gamma 2.0", "Macro F1 83.917%; Viral recall 67.713%; accuracy 84.439%.", "It focused more on difficult Viral cases but shifted too many Bacterial cases toward Viral; overall balance declined."],
            ["FT-007: label smoothing", "Label smoothing 0.00 -> 0.05; all else fixed", "Macro F1 85.631%; balanced accuracy 85.075%; accuracy 86.384%; Viral recall 66.816%.", "Best observed single seed-42 model. The original improvement gate was missed by 0.163 points, but FT-010 later received explicit authorization for confirmation."],
            ["FT-008A2: two U-Nets", "Two frozen U-Net artifacts for lung ROI", "Classic accepted 10/90; Legacy accepted 2/90; neither accepted any Viral or Bacterial image.", "Stopped permanently. The local U-Net artifacts did not generalize to the pediatric disease images."],
            ["FT-009A: error analysis", "No training; compare saved per-image predictions", "87.18%-93.28% of remaining errors were Viral/Bacterial.", "Confirmed that etiology separation, not Normal/Pneumonia detection, was the dominant error source."],
            ["FT-009B: two-stage hierarchy", "Model 1: Normal/Pneumonia; Model 2: Viral/Bacterial", f"Combined macro F1 {pct(ft009b.get('macro_f1'))}; accuracy {pct(ft009b.get('accuracy'))}; balanced accuracy {pct(ft009b.get('balanced_accuracy'))}; Viral recall {pct(ft009b.get('per_class', {}).get('Viral', {}).get('recall'))}.", "Viral recall improved substantially, but combined macro F1 missed the frozen advancement gate; stopped at seed 42."],
            ["FT-010A: label-smoothing confirmation", "Repeat FT-007 at seeds 123 and 2026", "Completed 3/3 seeds: mean macro F1 85.029%; accuracy 85.850%; balanced accuracy 84.501%; Viral recall 65.919%.", "Consistent and simpler than FT-004, but its primary-metric mean remained 0.276 points lower."],
        ],
        [1450, 2150, 2850, 2910],
        font_size=8.15,
    )
    add_heading(doc, "Why FT-009B was scientifically useful even though it stopped", 2)
    add_para(
        doc,
        f"Stage 1 achieved {pct(ft009b_stage1.get('macro_f1'))} macro F1 for Normal versus Pneumonia, while Stage 2 achieved {pct(ft009b_stage2.get('macro_f1'))} macro F1 for Viral versus Bacterial. The hierarchy verified that pneumonia detection is comparatively easy and etiology separation is the bottleneck. Its combined Viral recall rose to {pct(ft009b.get('per_class', {}).get('Viral', {}).get('recall'))}, but this came with lower Bacterial recall and insufficient overall macro F1.",
    )

    # Page 8 - model decision.
    page_break(doc)
    add_heading(doc, "7. Which model is best right now?", 1)
    add_table(
        doc,
        ["Claim", "Configuration", "Evidence", "Correct wording"],
        [
            ["Best original validation-selected pipeline", "DenseNet-201 + VIRIDIS -> GE -> CLAHE", "Three-seed mean macro F1 84.168%.", "Final validation-selected pipeline before post-selection fine-tuning."],
            ["Chosen final research model", "FT-004 equal-probability ensemble", "Three-seed mean macro F1 85.305% and accuracy 86.270%.", "Highest fully three-seed-confirmed mean macro F1; uses three DenseNet-201 members per image."],
            ["Best simpler single-network alternative", "FT-010 label smoothing 0.05", "Three-seed mean macro F1 85.029%; sample standard deviation 0.568 points.", "Statistically competitive and cheaper, but ranked second under the fixed macro-F1 rule."],
            ["Highest Viral recall among current focused candidates", "FT-009B hierarchy", f"Viral recall {pct(ft009b.get('per_class', {}).get('Viral', {}).get('recall'))}; macro F1 {pct(ft009b.get('macro_f1'))}.", "Useful sensitivity tradeoff, not the best overall model."],
            ["Only locked-test-evaluated model", "Earlier DenseNet-201 + GE", "Locked-test macro F1 78.877%; accuracy 79.863%.", "Official one-time locked-test result for that earlier frozen configuration only."],
        ],
        [1900, 2450, 2150, 2860],
        font_size=8.2,
    )
    add_callout(
        doc,
        "Final decision",
        "Freeze FT-004 as the final validation-selected research model. For deployment, use exactly three original seed-42 checkpoints: VIRIDIS -> GE -> CLAHE, VIRIDIS -> GE, and GE. Average their Normal, Viral, and Bacterial probabilities equally. Seeds 42, 123, and 2026 confirm repeatability; they do not mean nine networks are loaded.",
    )
    add_heading(doc, "Why we should not chase 90% accuracy by training longer", 2)
    add_para(
        doc,
        "More epochs do not guarantee better generalization. Several runs reached their best validation score before the final epoch and then worsened. The training system therefore saves the best validation checkpoint and uses early stopping. A higher training score without higher validation macro F1 would indicate memorization rather than a better model.",
    )

    # Page 9 - likely questions.
    page_break(doc)
    add_heading(doc, "8. Likely questions and short answers", 1)
    qa_rows = [
        ["Why use 224 x 224?", "It matches the standard ImageNet-pretrained contract, fit the available GPU, and beat the controlled 320 x 320 experiment on the primary metric."],
        ["Why not select by accuracy?", "The classes are imbalanced. Macro F1 gives each class equal weight, so poor Viral performance cannot be hidden by the larger Bacterial class."],
        ["Why three seeds?", "To see whether performance repeats under different controlled random initializations, batch orders, and augmentations."],
        ["Why does the chosen system use three models?", "The preprocessing views make complementary errors. Each member outputs three probabilities, and their equal mean produced the highest confirmed average macro F1."],
        ["Are there nine deployed models because there are three seeds?", "No. The deployable package uses three seed-42 members. The other seeds are repeatability experiments for the same three-member design."],
        ["Does VIRIDIS add medical information?", "No. It remaps existing intensity values. Any benefit is an empirical interaction with the pretrained network and later histogram operations."],
        ["Why did ROI segmentation stop?", "The masks and crops failed disproportionately on Viral and Bacterial images. Using them would remove information or create class-dependent preprocessing."],
        ["Were the segmentation images pediatric?", "Yes. The complete Kermany cohort is pediatric, approximately ages one to five; exact individual ages are not encoded in the filenames."],
        ["Why did the hierarchical model help Viral recall?", "It gave Viral/Bacterial separation a dedicated model instead of making one network solve both pneumonia detection and etiology simultaneously."],
        ["Why was FT-009B not selected?", "Its Viral recall improved, but its combined macro F1 did not beat the frozen advancement requirement and Bacterial recall declined."],
        ["Can the old locked-test score be used for the new model?", "No. Metrics belong to the exact architecture, preprocessing, weights, seed policy, and checkpoint that generated them."],
        ["Is the system ready for hospitals?", "No. It needs independent external validation, clinical review, calibration checks, and prospective governance."],
    ]
    add_table(doc, ["Question", "Suggested answer"], qa_rows, [2850, 6510], font_size=8.7)

    # Page 10 - discussion script and next steps.
    page_break(doc)
    add_heading(doc, "9. Five-minute discussion script", 1)
    script_parts = [
        ("Opening", "Our project evaluates three-class pediatric pneumonia classification: Normal, Viral, and Bacterial. The emphasis is fair measurement and leakage prevention, not accuracy alone."),
        ("Data and controls", "We audited 17,591 files, retained 5,824 unique canonical images, and created group-aware training, validation, and test assignments. Eight pretrained classifiers were compared under shared controls."),
        ("Preprocessing selection", "DenseNet-201 and ResNet-50 received all 16 ordered VIRIDIS, GE, and CLAHE conditions. DenseNet-201 with VIRIDIS -> GE -> CLAHE ranked first on three-seed mean validation macro F1."),
        ("Fine-tuning findings", "Higher resolution and added affine augmentation did not improve macro F1. Focal loss raised Viral recall but reduced overall balance. Lung segmentation failed class-dependent gates. The hierarchy showed that Viral-versus-Bacterial separation is the main bottleneck. Label smoothing produced the best single seed-42 model."),
        ("Conclusion", "FT-010 completed at 85.029% mean macro F1. FT-004 remained higher at 85.305% and is frozen as the final validation-selected three-member system. Its improvement over the original pipeline was supported by paired bootstrap analysis, while external validation remains required."),
    ]
    for label, text in script_parts:
        add_para(doc, f"{label}: {text}", bold_lead=f"{label}:")

    add_heading(doc, "10. Next steps", 1)
    add_bullets(
        doc,
        [
            "Use the frozen three-checkpoint FT-004 inference package for demonstrations and research comparison.",
            "Keep FT-010 documented as the lower-compute alternative and FT-009B as the higher-Viral-recall alternative.",
            "Use the six ensemble Grad-CAM panels for qualitative error discussion, not as segmentation proof.",
            "Integrate the new PyTorch three-class predictor into the application only through a separate, tested endpoint; do not overwrite the legacy TensorFlow binary model silently.",
            "Use an independent external pediatric dataset before any generalization or clinical-readiness claim.",
        ],
    )
    add_callout(
        doc,
        "Final sentence to use",
        "The final validation-selected system is the frozen FT-004 three-member DenseNet-201 ensemble, chosen by the predeclared macro-F1 rule and supported by paired validation analysis; independent external confirmation is still required, especially for Viral-versus-Bacterial discrimination.",
    )

    # Appendix - source paths and snapshot note.
    page_break(doc)
    add_heading(doc, "Appendix: source files used for this guide", 1)
    source_rows = [
        ["Full journal", "model/docs/Pneumonia_Three_Class_Classification_Full_Journal_20260810.docx"],
        ["Fine-tuning Word report", "model/docs/fine_tuning/DenseNet201_VIRIDIS_GE_CLAHE_Fine_Tuning_Report.docx"],
        ["Fine-tuning audit journal", "model/docs/fine_tuning/DenseNet201_VIRIDIS_GE_CLAHE_Fine_Tuning_Journal.md"],
        ["FT-009B final result", "model/results/ft009b_hierarchical_s42/ft009b_result.json"],
        ["FT-010 protocol", "model/configs/fine_tuning/ft010_label_smoothing_multiseed_protocol.json"],
        ["FT-010 completed summary", "model/results/ft010_label_smoothing_multiseed/summary.json"],
        ["Frozen FT-004 selection", "model/results/final_ft004_ensemble/frozen_selection.json"],
        ["Paired bootstrap result", "model/results/final_ft004_ensemble/paired_bootstrap.json"],
        ["Ensemble Grad-CAM audit", "model/results/final_ft004_ensemble/gradcam/gradcam_audit.json"],
    ]
    add_table(doc, ["Evidence", "Repository path"], source_rows, [2700, 6660], font_size=8.7)
    add_para(
        doc,
        "This guide records the completed FT-010 confirmation and frozen FT-004 validation selection as of 13 August 2026. The old locked-test score is not assigned to FT-004; independent external evaluation remains pending.",
        italic=True,
    )

    configure_sections(doc)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_document())
