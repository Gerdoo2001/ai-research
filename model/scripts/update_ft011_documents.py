"""Publish actual FT-011 evidence and update black-text Word research reports.

Run with the Codex bundled document Python, not the training environment.
"""
from __future__ import annotations

import csv
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

import create_fine_tuning_report as word

ROOT = Path(__file__).resolve().parents[2]
DOCS = ROOT / "model/docs"
FINE = DOCS / "fine_tuning"
RESULTS = ROOT / "model/results/ft011"
WORK = DOCS / "_ft011_word_qa"


def read(path):
    return json.loads(path.read_text(encoding="utf-8"))


def make_figures(history, metrics, protocol):
    from PIL import Image, ImageDraw, ImageFont
    import math
    figures = FINE / "figures"
    figures.mkdir(exist_ok=True)
    epochs = [int(row["epoch"]) for row in history]
    font_path=Path("C:/Windows/Fonts/arial.ttf")
    font=ImageFont.truetype(str(font_path),30) if font_path.exists() else ImageFont.load_default(size=30)
    small=ImageFont.truetype(str(font_path),26) if font_path.exists() else ImageFont.load_default(size=26)
    def dashed(draw,a,b,period=18,width=4):
        length=math.dist(a,b)
        if not length: return
        for distance in range(0,int(length),period*2):
            end=min(length,distance+period)
            draw.line((a[0]+(b[0]-a[0])*distance/length,a[1]+(b[1]-a[1])*distance/length,
                       a[0]+(b[0]-a[0])*end/length,a[1]+(b[1]-a[1])*end/length),fill="black",width=width)
    def line_chart(filename,title,series,minimum,maximum,ticks,threshold=None):
        canvas=Image.new("RGB",(1800,850),"white"); draw=ImageDraw.Draw(canvas)
        left,right,top,bottom=150,1740,100,610
        point=lambda x,y:(left+(x-1)/(max(epochs)-1)*(right-left),bottom-(y-minimum)/(maximum-minimum)*(bottom-top))
        draw.text((left,22),title,font=font,fill="black")
        for value in ticks:
            y=point(1,value)[1]; draw.line((left,y,right,y),fill="#D9D9D9",width=2)
            draw.text((35,y-18),f"{value:g}",font=small,fill="black")
        draw.line((left,top,left,bottom,right,bottom),fill="black",width=3)
        for epoch in epochs:
            x=point(epoch,minimum)[0]; draw.text((x-12,bottom+18),str(epoch),font=small,fill="black")
        draw.text(((left+right)//2-35,bottom+58),"Epoch",font=font,fill="black")
        for index,(label,values) in enumerate(series):
            points=[point(e,v) for e,v in zip(epochs,values)]
            for a,b in zip(points,points[1:]):
                if index: dashed(draw,a,b,period=10 if index==1 else 3)
                else: draw.line((a,b),fill="black",width=4)
            for x,y in points:
                if index==0: draw.ellipse((x-7,y-7,x+7,y+7),outline="black",fill="white",width=3)
                elif index==1: draw.rectangle((x-7,y-7,x+7,y+7),outline="black",fill="white",width=3)
                else: draw.polygon([(x,y-8),(x-8,y+8),(x+8,y+8)],outline="black",fill="white")
            y=720+index*38
            if index: dashed(draw,(left,y+12),(left+80,y+12),period=10 if index==1 else 3)
            else: draw.line((left,y+12,left+80,y+12),fill="black",width=4)
            draw.text((left+100,y-5),label,font=small,fill="black")
        if threshold is not None:
            y=point(1,threshold)[1]; dashed(draw,(left,y),(right,y),period=25,width=2)
            draw.text((right-520,y-36),f"Required macro F1 {threshold:.3f}%",font=small,fill="black")
        canvas.save(figures/filename)
    line_chart("ft011_loss.png","FT 011 actual loss history",[
        ("Weighted smoothed training loss",[float(r["train_loss"]) for r in history]),
        ("Unweighted validation loss",[float(r["val_loss"]) for r in history])],0,.8,[0,.2,.4,.6,.8])
    line_chart("ft011_validation.png","FT 011 actual validation history percent",[
        ("Macro F1",[float(r["val_macro_f1"])*100 for r in history]),
        ("Balanced accuracy",[float(r["val_balanced_accuracy"])*100 for r in history]),
        ("Accuracy",[float(r["val_accuracy"])*100 for r in history])],70,90,[70,75,80,85,90],
        threshold=protocol["controls"]["screening_gates"][0]["threshold"]*100)
    names = ["FT 010 seed 42", "FT 004 seed 42", "FT 011 seed 42"]
    scores = [protocol["controls"]["systems"][s]["per_seed"]["42"]["macro_f1"]*100 for s in ("FT010", "FT004")]
    scores.append(metrics["macro_f1"]*100)
    canvas=Image.new("RGB",(1800,850),"white"); draw=ImageDraw.Draw(canvas)
    draw.text((150,25),"Matched seed 42 validation macro F1 percent",font=font,fill="black")
    for value in (0,20,40,60,80,100):
        y=650-value*5; draw.line((150,y,1730,y),fill="#D9D9D9",width=2); draw.text((45,y-15),str(value),font=small,fill="black")
    for index,(name,score) in enumerate(zip(names,scores)):
        x=320+index*500; y=650-score*5
        draw.rectangle((x,y,x+230,650),fill=["#E0E0E0","#BEBEBE","#8A8A8A"][index],outline="black",width=3)
        draw.text((x+45,y-40),f"{score:.3f}%",font=small,fill="black")
        draw.text((x,680),name,font=font,fill="black")
    canvas.save(figures/"ft011_matched_comparison.png")
    matrix=metrics["confusion_matrix"]
    canvas=Image.new("RGB",(1200,850),"white"); draw=ImageDraw.Draw(canvas)
    names = ["Normal", "Viral", "Bacterial"]
    draw.text((150,25),"FT 011 selected epoch 5 confusion matrix",font=font,fill="black")
    draw.text((550,90),"Predicted class",font=font,fill="black")
    for i in range(3):
        draw.text((30,235+i*175),names[i],font=font,fill="black")
        draw.text((245+i*260,145),names[i],font=font,fill="black")
        for j in range(3):
            shade=255-int(matrix[i][j]/max(max(r) for r in matrix)*150)
            draw.rectangle((220+j*260,200+i*175,480+j*260,375+i*175),fill=(shade,shade,shade),outline="black",width=2)
            draw.text((330+j*260,260+i*175),str(matrix[i][j]),font=font,fill="black")
    draw.text((220,780),"Rows are true Normal Viral Bacterial classes",font=small,fill="black")
    canvas.save(figures/"ft011_confusion_matrix.png")


def markdown(decision, protocol, metrics, history, telemetry):
    p = lambda value: f"{value*100:.3f}%"
    rows = ["# FT 011 Final Controlled Experiment and Selected Model", "",
        "Updated 13 October 2026 Asia Manila. Decision: **retain FT-004**. FT-011 completed seed 42, failed four of five screening gates, and closed this tuning phase. Seeds 123 and 2026 were not run. We retain the existing frozen three-member DenseNet-201 ensemble rather than replace it with the new single-network candidate.", "",
        "## Purpose and comparison", "",
        "We tested whether a smaller backbone learning rate would improve the established VIRIDIS -> Global Histogram Equalization (GE) -> Contrast Limited Adaptive Histogram Equalization (CLAHE) classifier. The backbone is the feature-extracting part; the classifier converts those features into three class scores. FT-010 used 0.0001 throughout. FT-011 used 0.00001 for the backbone and 0.0001 for the classifier, with every layer trainable. This was a fixed comparison, not a new search.", "",
        "Transferred features can vary in task specificity, which motivates testing how aggressively they should adapt [1]. Domain-adaptation research provides related context [2]. Neither publication establishes that these exact rates are optimal for this dataset. The rates and numerical decision gates are project choices. Prespecifying the gates reduces post-result selection flexibility; repeated validation tuning can still introduce selection bias [3].", "",
        "## Data preparation and preprocessing", "",
        "The experiment reused 5,824 canonical unique images and their locked group-controlled assignments: 4,076 training, 874 validation and 874 sealed-test cases. The training classes contain 1,105 Normal, 1,039 Viral and 1,932 Bacterial images. Validation contains 237 Normal, 223 Viral and 414 Bacterial images. Assigned split comes from the manifest, not the legacy directory prefix. FT-011 opened only training and validation images.", "",
        "No new split, relabelling, duplicate removal or ROI was introduced. Git had converted the four manifest CSVs to CRLF; restoring LF recovered the original exact hashes without changing a row or label. Derived filename and exact-hash groups control known overlap but do not establish verified patient independence.", "",
        "The pipeline opens the original image and converts it to RGB. VIRIDIS applies the OpenCV colour map to grayscale intensity. GE uses PIL global grayscale equalization. CLAHE then uses OpenCV grayscale enhancement with clip limit 2.0 and an 8 by 8 tile grid. Each grayscale enhancement returns replicated RGB channels. Therefore VIRIDIS affects the intermediate intensity transformation, but the final input is not a retained colourful palette.", "",
        "Training only then uses a random horizontal flip with probability 0.5 and a uniformly sampled rotation from -5 to +5 degrees, using bilinear interpolation and black fill. Both are retained from the control to keep the comparison matched. Validation has neither random operation. The image is aspect-ratio-preservingly fitted and center cropped to 224 by 224, converted to float32 pixels divided by 255, transposed from height-width-channel to channel-height-width, and normalized with ImageNet mean [0.485, 0.456, 0.406] and standard deviation [0.229, 0.224, 0.225]. Output is 3 by 224 by 224.", "",
        "## Training configuration and reasons", "",
        "Every candidate starts from fresh ImageNet DenseNet-201 weights with a seed-matched new three-output classifier. Fresh means no continued FT-010 or smoke-trained checkpoint; it does not mean a random backbone. Seed 42 fixes the random-number sequence used for initialization, order and augmentation. The number 42 is an arbitrary reproducibility label, not 42 models or 42 training attempts.", "",
        "AdamW uses weight decay 0.0001. The backbone initial rate 0.00001 is ten times smaller than the classifier rate 0.0001. Learning rate controls update size; weight decay discourages large weights. The lower backbone rate was intended to preserve transferable features while the new head adapted. Its weaker result does not prove a mechanism such as underfitting; this experiment only establishes that this fixed policy did not meet the gates.", "",
        "Balanced weighted cross-entropy uses the training-count formula N divided by K times each class count. The Normal, Viral and Bacterial weights are 1.2295626401901245, 1.30766761302948 and 0.703243613243103. Label smoothing is 0.05, reducing the hard certainty of training targets. Validation loss is ordinary unweighted cross-entropy, so its magnitude is not directly interchangeable with the weighted smoothed training loss.", "",
        "Batch size remains 8, which defines examples per optimizer update. The maximum is 12 epochs, where an epoch is one pass over the assigned training cases. Early stopping waits four epochs without macro-F1 improvement exceeding 0.0001. A new best checkpoint requires strictly greater macro F1 than the prior best plus 0.0001. These unchanged rules selected epoch 5 and stopped after epoch 9.", "",
        "ReduceLROnPlateau monitors validation macro F1 in max mode, with factor 0.5 and patience 2. Its relative threshold is 0.0001, cooldown 0, minimum rates [0,0] and epsilon 1e-8. Epoch 8 triggered the rate reduction for the next epoch: backbone 0.000005 and classifier 0.00005. CUDA mixed precision, all BatchNorm layers, zero loader workers, original seed handling and data ordering were retained. The parameter groups are disjoint and cover all 18,098,691 trainable parameters.", "",
        "The runtime matched Python 3.12.6, PyTorch 2.5.1+cu121, torchvision 0.20.1+cu121, CUDA 12.1 and cuDNN 90100. ImageNet weights resolve to IMAGENET1K_V1 with SHA256 c1103571aad4f45d0b5b15556d92944db94c33bee68cd58babba3d76225a6401. Historical runs did not save initial-state hashes or per-run source commits; FT-011 records them for the current implementation and fresh initialization.", "",
        "## Exact screening gates", "",
        "Every gate had to pass before the two confirmation seeds were authorized. Full-precision unrounded metrics determine the decision; displayed percentages are for reading only. The macro-F1 margin is 0.005 on the zero-to-one scale, which is 0.500 percentage points, not a 0.5 percent relative increase.", "",
        "| Metric | Exact requirement | Exact observed | Result |", "| --- | --- | --- | --- |"]
    labels = {"macro_f1":"Macro F1", "balanced_accuracy":"Balanced accuracy", "viral_recall":"Viral recall", "normal_recall":"Normal recall", "bacterial_recall":"Bacterial recall"}
    for gate in decision["screening_gates"]:
        rows.append(f"| {labels[gate['metric']]} | {gate['operator']} {gate['threshold']:.16f} | {gate['observed']:.16f} | {'PASS' if gate['passed'] else 'FAIL'} |")
    rows.extend(["", "Only Normal recall passed. The protocol therefore requires retaining FT-004, stopping FT-011 and not running seeds 123 or 2026. Three-seed replacement gates are not evaluated, and a one-seed result has no three-seed sample standard deviation.", "", "## Actual results and training history", "",
        f"The selected epoch-5 checkpoint achieved accuracy {p(metrics['accuracy'])}, balanced accuracy {p(metrics['balanced_accuracy'])} and macro F1 {p(metrics['macro_f1'])}. Macro F1 gives each class equal importance; accuracy counts the overall fraction of correct images. Balanced accuracy is the mean of class recalls. Recall asks how many actual examples of a class were found, precision asks how many predicted examples were correct, and F1 combines precision and recall.", "",
        "| Class | Precision | Recall | F1 | Images |", "| --- | --- | --- | --- | --- |"])
    for name in ("Normal","Viral","Bacterial"):
        c = metrics["per_class"][name]
        rows.append(f"| {name} | {p(c['precision'])} | {p(c['recall'])} | {p(c['f1'])} | {c['support']} |")
    rows.extend(["", "| Epoch | Train loss | Val loss | Macro F1 | Duration seconds |", "| --- | --- | --- | --- | --- |"])
    for h in history:
        rows.append(f"| {h['epoch']} | {float(h['train_loss']):.6f} | {float(h['val_loss']):.6f} | {p(float(h['val_macro_f1']))} | {float(h['elapsed_seconds']):.3f} |")
    rows.extend(["", "![Actual loss history](figures/ft011_loss.png)", "Figure 1 Actual weighted training and unweighted validation loss histories", "",
        "![Actual validation history](figures/ft011_validation.png)", "Figure 2 Actual validation scores with the fixed macro F1 screening target", "",
        "![Selected confusion matrix](figures/ft011_confusion_matrix.png)", "Figure 3 True classes are rows and predicted classes are columns", "",
        "## Matched comparisons and uncertainty", "",
        "| System at seed 42 | Accuracy | Balanced accuracy | Macro F1 |", "| --- | --- | --- | --- |"])
    for name in ("FT010", "FT004", "FT011"):
        m = metrics if name == "FT011" else protocol["controls"]["systems"][name]["per_seed"]["42"]
        rows.append(f"| {name} | {p(m['accuracy'])} | {p(m['balanced_accuracy'])} | {p(m['macro_f1'])} |")
    rows.extend(["", "![Matched seed comparison](figures/ft011_matched_comparison.png)", "Figure 4 Matched seed 42 comparisons do not mix a single seed with a three seed mean", "",
        "The paired bootstrap used 10,000 resamples with identical sampled cases across systems. It sampled intact derived groups within class strata, calculated each seed's metric separately and then averaged metrics when multiple seeds existed. Only seed 42 exists for FT-011. This uses 228 proxy groups, including 181 repeated-image groups. They are derived filename and hash clusters rather than verified patient identifiers.", "",
        "| Comparison | Macro F1 difference points | 95 percent interval points |", "| --- | --- | --- |"])
    for comp in decision["bootstrap"]["comparisons"]:
        v=comp["metrics"]["macro_f1"]; lo,hi=v["ci_95_percentile"]
        rows.append(f"| FT011 minus {comp['comparator']} | {v['point_difference']*100:.6f} | {lo*100:.6f} to {hi*100:.6f} |")
    rows.extend(["", "Both observed macro-F1 differences and their validation bootstrap intervals are negative. This supports retaining the incumbent for this matched comparison. It does not establish population performance across arbitrary random seeds, independent external accuracy or verified patient independence [3,4,6].", "",
        "## Laptop protection and technical recovery", "",
        "The protected smoke used only 16 training and 16 validation cases; its scores are execution checks, not experimental results. Its simulated unsafe reading exercised complete atomic checkpointing, state restoration and at least 120 continuous healthy seconds before resuming. Full seed-42 training started fresh afterward.", "",
        "The first smoke attempt paused for RAM pressure, then disappeared without a recorded termination cause. Its checkpoint and logs were preserved under interrupted_checks/20261013_1349 before the identical smoke was repeated. No prior full FT-011 run is claimed from the protocol's earlier ongoing status label.", "",
        "During full epoch 4 at 14:16:39 Manila, a Windows access-denied error prevented atomic telemetry-file replacement. Required monitoring was treated as unavailable, GPU work paused at a safe batch boundary, and the full model, optimizer, scheduler, scaler, random states, early-stopping state, pending transformed batch and sampler position were checkpointed and restored. The run resumed at 14:18:45 after 122.547 continuous healthy seconds. This interruption must be disclosed even though the aggregate monitoring-gap counter is zero.", "",
        f"Recorded active runtime was {metrics['training_duration_seconds']:.3f} seconds and wall runtime {metrics['wall_duration_seconds']:.3f} seconds. The GPU peaked at {telemetry['peaks']['gpu_temperature_c']:.1f} C. Device VRAM peaked at {telemetry['peaks']['gpu_vram_used_bytes']:,} bytes; peak PyTorch allocated memory was {metrics['peak_gpu_memory_bytes']:,} bytes. System RAM used peaked at {telemetry['peaks']['ram_used_bytes']:,} bytes. Actual overheating and thermal-throttle samples were zero. Shared monitoring history contains one simulated smoke pause and two cooling pauses in total. CPU temperature was unavailable, not zero.", "",
        "The RTX 2050 reports target 87 C, slowdown 97 C and shutdown 100 C. Project margins warn at 80 C, pause at 84 C, stop critically at 92 C, and resume below 75 C. RAM pauses below 1 GiB available and resumes at least 2 GiB; VRAM pauses at 98 percent and resumes below 92 percent. Required heartbeat age is at most six seconds. Checkpoint disk reserve is 1.5 GiB. These are hardware-anchored project choices rather than universal safe-temperature claims.", "",
        f"Selected-checkpoint inference measured {metrics['inference']['latency_ms_per_image']:.3f} milliseconds per image and {metrics['inference']['throughput_images_per_second']:.3f} images per second, using batch one, five warm-ups and twenty timed full-float passes. Preprocessing is excluded; this is not total application latency or the three-network FT-004 cost.", "",
        "## Retained best model and publication", "",
        "The selected system remains FT-004: three seed-42 DenseNet-201 checkpoints with member pipelines VIRIDIS -> GE -> CLAHE, VIRIDIS -> GE, and GE. Each produces Normal, Viral and Bacterial probabilities. We take their equal arithmetic mean and choose the largest probability. The three confirmation seeds describe repeatability of this design; deployment uses three networks, not nine. Ensemble research motivates this aggregation but does not prove its benefit for every population [5].", "",
        "The previously completed FT-004 three-seed means are macro F1 85.30516857750923 percent, accuracy 86.27002288329519 percent, balanced accuracy 84.92846583463107 percent and Viral recall 66.51718983557549 percent. These means are different summaries from the matched seed-42 table above. FT-004 remains a validation-selected research prototype, not a clinically validated device.", "",
        "Its original frozen selection SHA256 remains bbb65ee1202b2c34490cdb1047339795e5751f2c71edf6626d1b1f235acea57b. The portable repository publication copies its three original checkpoints without changing their bytes, records each hash and original configuration provenance, and resolves member files relative to the published selection. Git LFS carries the large weights. No FT-011 weights replace the selected model and the live website remains unchanged.", "",
        "Canonical evidence is model/results/ft011/protocol.json, final_decision.json, frozen_selection.json, the seed-42 run directory, and the timestamped analysis directory. The documentation index identifies the current Word reports and model package. Historical records and the interrupted smoke remain preserved.", "",
        "## Conclusion and next evaluation", "",
        "FT-011 is a completed negative result for this fixed learning-rate policy. It did not improve the matched control enough to satisfy the predeclared gates, and the tuning phase is closed. The next step is genuinely independent evaluation of the frozen FT-004 system on a separately collected, appropriately labelled cohort. Do not transfer the earlier DenseNet-201 plus GE locked-test score to FT-004 or describe validation reuse as external performance [4,6].", "",
        "## Numbered references", ""])
    references=[
        ("Yosinski et al 2014 How transferable are features in deep neural networks", "https://papers.nips.cc/paper_files/paper/2014/hash/532a2f85b6977104bc93f8580abbb330-Abstract.html"),
        ("Long et al 2015 Learning Transferable Features with Deep Adaptation Networks", "https://proceedings.mlr.press/v37/long15.html"),
        ("Cawley and Talbot 2010 On Over fitting in Model Selection and Subsequent Selection Bias", "https://www.jmlr.org/papers/v11/cawley10a.html"),
        ("Collins et al 2024 TRIPOD AI statement", "https://pmc.ncbi.nlm.nih.gov/articles/PMC11019967/"),
        ("Lakshminarayanan et al 2017 Simple and Scalable Predictive Uncertainty Estimation using Deep Ensembles", "https://papers.nips.cc/paper_files/paper/2017/hash/9ef2ed4b7fd2c810847ffa5fa85bce38-Abstract.html"),
        ("Riley et al 2024 How to undertake an external validation study", "https://pmc.ncbi.nlm.nih.gov/articles/PMC10788734/")]
    rows += [f"[{i}] [{name}]({url})" for i,(name,url) in enumerate(references,1)]
    return "\n".join(rows)+"\n"


def add_table(document, rows):
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    total=int(document.sections[-1].page_width-document.sections[-1].left_margin-document.sections[-1].right_margin)//635
    widths=word.allocate_widths(rows,total)
    word.set_table_geometry(table,widths,total)
    for i,row in enumerate(rows):
        for j,text in enumerate(row):
            cell=table.cell(i,j); cell.text=""
            p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(3); p.paragraph_format.line_spacing=1.05
            p.paragraph_format.keep_with_next=(i==0 or (len(rows)<=4 and i<len(rows)-1))
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
            word.add_inline_markdown(p,text,size=9.5)
            for run in p.runs: run.bold=i==0
            if i==0:
                shading=OxmlElement("w:shd"); shading.set(qn("w:fill"),"E7E6E6"); cell._tc.get_or_add_tcPr().append(shading)
            word.set_cell_margins(cell,top=100,bottom=100,start=100,end=100)
    borders=OxmlElement("w:tblBorders")
    for edge in ("top","left","bottom","right","insideH","insideV"):
        item=OxmlElement("w:"+edge); item.set(qn("w:val"),"single"); item.set(qn("w:sz"),"4"); item.set(qn("w:color"),"D9D9D9"); borders.append(item)
    table._tbl.tblPr.append(borders); word.mark_repeat_header(table.rows[0]); document.add_paragraph()


def append_content(document, content, *, title=False):
    lines=content.splitlines(); i=0
    while i<len(lines):
        line=lines[i].strip()
        if not line: i+=1; continue
        if line.startswith("|"):
            table=[]
            while i<len(lines) and lines[i].strip().startswith("|"):
                cells=[c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(set(c)<=set(" :-") for c in cells): table.append(cells)
                i+=1
            add_table(document,table); continue
        image=re.match(r"!\[(.*?)\]\((.*?)\)",line)
        if image:
            p=document.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next=True
            inline=p.add_run().add_picture(str(FINE/image.group(2)),width=Inches(6.1))
            inline._inline.docPr.set("descr",image.group(1)); i+=1; continue
        if line.startswith("# "):
            p=document.add_paragraph(line[2:],style="Title" if title else "Heading 1")
        elif line.startswith("## "):
            p=document.add_paragraph(line[3:],style="Heading 1")
        elif line.startswith("Figure "):
            p=document.add_paragraph(line,style="Caption")
        else:
            p=document.add_paragraph(); word.add_inline_markdown(p,line)
        i+=1


def save(document,path):
    for style in document.styles:
        if style.name in ("Title","Subtitle"):
            for border in list(style.element.findall(".//"+qn("w:pBdr"))):
                border.getparent().remove(border)
    for p in document.paragraphs[:12]:
        for border in list(p._element.findall(".//"+qn("w:pBdr"))):
            border.getparent().remove(border)
    for table in document.tables:
        if not table.rows: continue
        word.mark_repeat_header(table.rows[0])
        for i,row in enumerate(table.rows):
            trpr=row._tr.get_or_add_trPr()
            if trpr.find(qn("w:cantSplit")) is None: trpr.append(OxmlElement("w:cantSplit"))
            if i==0:
                for cell in row.cells:
                    for p in cell.paragraphs: p.paragraph_format.keep_with_next=True
    word.force_all_text_colors_black(document)
    settings=document.settings.element
    update=OxmlElement("w:updateFields"); update.set(qn("w:val"),"true"); settings.append(update)
    document.save(path); word.scrub_saved_ooxml_text_colors(path)


def build():
    WORK.mkdir(exist_ok=True)
    decision,protocol=read(RESULTS/"final_decision.json"),read(RESULTS/"protocol.json")
    run_dir=RESULTS/"runs/ft011_densenet201_differential_lr_s42"
    metrics,telemetry=read(run_dir/"metrics.json"),read(RESULTS/"monitoring/final_telemetry_summary.json")
    with (run_dir/"history.csv").open(encoding="utf-8",newline="") as handle: history=list(csv.DictReader(handle))
    make_figures(history,metrics,protocol)
    content=markdown(decision,protocol,metrics,history,telemetry)
    (FINE/"FT011_Final_Report_20261013.md").write_text(content,encoding="utf-8")
    standalone=Document(); word.configure_styles(standalone); word.configure_section(standalone.sections[0],landscape=False)
    append_content(standalone,content,title=True)
    save(standalone,FINE/"FT011_Final_Report_20261013.docx")
    fine=Document(FINE/"DenseNet201_VIRIDIS_GE_CLAHE_Fine_Tuning_Report_Final_20260813.docx")
    fine.paragraphs[1].text="DenseNet 201 Pneumonia Classification"; fine.paragraphs[1].style="Title"
    for table in fine.tables:
        for row in table.rows:
            if len(row.cells)==2 and row.cells[0].text=="Document date": row.cells[1].text="13 October 2026"
            if len(row.cells)==2 and row.cells[0].text=="Scope": row.cells[1].text="FT-001 through FT-011 and retained FT-004 selection; complete chronological fine-tuning and freeze record"
    for p in fine.paragraphs:
        if p.text.startswith("Current status:"):
            p.text="Current status: FT-011 completed seed 42 and failed four screening gates. FT-004 remains the selected frozen three-member system. FT-011 seeds 123 and 2026 were not run; the tuning phase is closed."
        elif p.text.startswith("Created: 2026-08-10;"):
            p.text="Created: 2026-08-10; historical confirmation finalized: 2026-08-13; FT-011 outcome updated: 2026-10-13"
        elif p.text.startswith("Final artifact:"):
            p.text="Current artifact: DenseNet201_VIRIDIS_GE_CLAHE_Fine_Tuning_Report.docx; historical dated versions remain preserved."
    reporting_rules=False
    for p in fine.paragraphs:
        if p.text=="Reporting rules": reporting_rules=True
        elif reporting_rules and p.style.name.startswith("Heading"): reporting_rules=False
        elif reporting_rules and p.style.name=="List Number":
            p.paragraph_format.space_before=Pt(0)
            p.paragraph_format.space_after=Pt(2)
            p.paragraph_format.line_spacing=1.05
    fine.add_page_break(); append_content(fine,content)
    save(fine,FINE/"DenseNet201_VIRIDIS_GE_CLAHE_Fine_Tuning_Report.docx")
    journal=Document(DOCS/"Pneumonia_Three_Class_Classification_Full_Journal_20260810.docx")
    for p in journal.paragraphs:
        if p.style.name=="Title": p.text="Three Class Pneumonia Classification Research Journal"
        elif p.text.startswith("Updated 10 August 2026"):
            p.text="Updated 13 October 2026 | Earlier baseline and preprocessing findings retained as historical evidence"
        elif p.text.startswith("Reporting boundary:"):
            p.text="Reporting boundary: the current validation-selected system is the FT-004 three-member DenseNet-201 ensemble. Earlier sections preserve the preceding baseline and preprocessing-only phase. FT-011 used training and validation only; the earlier DenseNet-201 plus GE test result does not transfer to FT-004. Independent external evaluation remains required."
        elif p.text.startswith("In one sentence:"):
            p.text="In one sentence: the project compared pediatric chest X-ray classifiers and preprocessing policies, then fine-tuning selected the FT-004 three-member DenseNet-201 ensemble; the final FT-011 learning-rate experiment did not meet its screening gates, so FT-004 remains frozen."
        elif p.text.startswith("DenseNet-201 + VIRIDIS -> GE -> CLAHE is therefore the final"):
            p.text="The earlier preprocessing-only phase selected DenseNet-201 + VIRIDIS -> GE -> CLAHE. Subsequent fine-tuning confirmed FT-004 as the highest-mean fully three-seed-confirmed finalist. The October FT-011 learning-rate comparison failed its screening gates, so FT-004 remains the current frozen validation-selected research system."
        elif p.text.startswith("DenseNet-201 + VIRIDIS -> GE -> CLAHE is the selected final"):
            p.text="The historical preprocessing-only winner is DenseNet-201 + VIRIDIS -> GE -> CLAHE; the current retained research system after fine-tuning is FT-004."
        elif p.text.startswith("Final validation-selected candidate."):
            p.text="Final validation-selected candidate. The system chosen using validation comparisons. The current retained system is FT-004, a three-member DenseNet-201 ensemble; the earlier preprocessing-only winner was the single VIRIDIS -> GE -> CLAHE model."
        elif p.text.startswith("SAS."):
            p.text="SAS. Spatial Attention Superposition, the project's classification attention adaptation, not lung segmentation. It was not part of the selected standard DenseNet-201 model."
        elif p.text.startswith("MFF."):
            p.text="MFF. Multilayer Feature Fusion: the project's channel-ranking, global-pooling and 1D-convolution adaptation, not fusion of separately extracted network depths. It did not exceed the selected standard model."
    for row in journal.tables[0].rows:
        label=row.cells[0].text
        replacements={"Status":"FT-011 completed; four screening gates failed; retain FT-004 and close tuning",
            "Final validation-selected candidate":"FT-004 three-member DenseNet-201 probability ensemble",
            "Seeds":"FT-004 confirmed at 42, 123 and 2026; deployment uses three original seed-42 checkpoints",
            "Three-seed mean validation macro F1":"0.8530516857750923 (85.305%)",
            "Three-seed mean balanced accuracy":"0.8492846583463107 (84.928%)",
            "Three-seed mean accuracy":"0.8627002288329519 (86.270%)",
            "Locked-test boundary":"FT-004 remains validation-selected; earlier GE test metrics do not transfer",
            "Generated":"13 October 2026 Asia Manila"}
        if label in replacements: row.cells[1].text=replacements[label]
    for row in journal.tables[1].rows:
        if row.cells[0].text=="Final validation-selected candidate": row.cells[0].text="Historical preprocessing-only selection"
    for table in journal.tables:
        for row in table.rows:
            if row.cells[0].text.strip()=="Final validation-selected candidate":
                for cell in row.cells[1:]:
                    if "Here it is" in cell.text or "Hereitis" in cell.text.replace(" ",""):
                        cell.text="The current retained research system is FT-004, a three-member DenseNet-201 ensemble. The earlier preprocessing-only winner was the single VIRIDIS -> GE -> CLAHE model."
            if row.cells[0].text.strip()=="SAS":
                row.cells[1].text="Spatial Attention Superposition, the project's classification attention adaptation. It is not lung segmentation."
            if row.cells[0].text.strip()=="MFF":
                row.cells[1].text="Multilayer Feature Fusion, implemented here as a channel-ranking, global-average-pooling and one-dimensional-convolution classification adaptation. It does not combine separately extracted feature maps from multiple network depths."
    for p in list(journal.paragraphs):
        if p.style.name.startswith("Heading") and ("12.11" in p.text or "14. Conclusion" in p.text or "Final selection:" in p.text):
            note=journal.add_paragraph("Historical conclusion for the earlier preprocessing-only phase. The current retained system is FT-004; see the final October update.")
            p._element.addnext(note._element)
    first=next(p for p in journal.paragraphs if p.style.name=="Heading 1")
    notice=first.insert_paragraph_before("Current decision on 13 October 2026: retain FT-004, the three-member DenseNet-201 ensemble. Earlier sections describe the preceding baseline and preprocessing phases. The final update below records FT-011; the separate fine-tuning report contains its full method, recovery details and gate evidence.")
    notice.paragraph_format.space_after=Pt(8)
    journal.add_page_break()
    summary="\n".join(content.splitlines()[:2])+"\n\n"+content.splitlines()[2]+"\n\n"
    summary+="## Current selected system\n\nFT-004 uses three original seed-42 DenseNet-201 checkpoints with VIRIDIS -> GE -> CLAHE, VIRIDIS -> GE, and GE. The equal mean of their Normal, Viral and Bacterial probability vectors defines the prediction. Historical three-seed means are macro F1 85.305%, accuracy 86.270%, balanced accuracy 84.928% and Viral recall 66.517%. This is validation-selected research evidence, not independent external performance.\n\n"
    summary+="## Completed FT 011 result\n\nThe fresh ImageNet candidate changed only the backbone initial rate to 0.00001 while its classifier remained 0.0001. It completed nine epochs and selected epoch five. Accuracy was 83.410%, balanced accuracy 81.731%, macro F1 82.028%, Normal recall 98.734%, Viral recall 58.296% and Bacterial recall 88.164%. Normal recall passed its screening gate; the other four failed. Seeds 123 and 2026 were therefore not run.\n\n"
    summary+="![Matched seed comparison](figures/ft011_matched_comparison.png)\nFigure 1 Matched seed 42 comparison for the final controlled experiment\n\n"
    summary+="## Protection and next step\n\nActive training took 69 minutes 45 seconds; wall time was 71 minutes 51 seconds. Peak GPU temperature was 51 C, with no actual overheating or thermal throttling. A telemetry-file write failure caused a complete-state protective pause; recovery followed two healthy minutes. The first low-RAM smoke was archived; the repeated protected smoke passed before fresh training.\n\nThe fine-tuning report contains full methods and evidence. The portable package is model/artifacts/ft004; the website is unchanged. Independently evaluate frozen FT-004. Another configuration's earlier test scores do not apply.\n"
    append_content(journal,summary)
    save(journal,DOCS/"Pneumonia_Three_Class_Classification_Full_Journal_20261013.docx")
    print("Three Word reports and the full FT-011 Markdown report updated")


if __name__=="__main__":
    build()
