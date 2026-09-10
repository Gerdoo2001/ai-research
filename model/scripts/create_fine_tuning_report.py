"""Build the black-only fine-tuning Word report from the canonical Markdown journal."""

from __future__ import annotations

import argparse
import os
import re
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor


BLACK = RGBColor(0, 0, 0)
USABLE_PORTRAIT_DXA = 9360
USABLE_LANDSCAPE_DXA = 12960
TABLE_INDENT_DXA = 120


def set_run_font(run, name: str = "Calibri", size: float | None = None, *, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.color.rgb = BLACK
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int], total_width: int):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total_width))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(widths[index] / 1440)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)


def mark_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_run_font(run, size=9)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, separate, text, end))


def configure_section(section, *, landscape: bool = False):
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Inches(11)
        section.page_height = Inches(8.5)
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    header = section.header.paragraphs[0]
    header.text = "DenseNet-201 Fine-Tuning and ROI Experiment Record"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.paragraph_format.space_after = Pt(0)
    for run in header.runs:
        set_run_font(run, size=8.5)
    footer = section.footer.paragraphs[0]
    footer.clear()
    add_page_field(footer)


def configure_styles(document: Document):
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    tokens = {
        "Title": (26, 0, 6),
        "Subtitle": (14, 0, 12),
        "Heading 1": (16, 18, 10),
        "Heading 2": (13, 14, 7),
        "Heading 3": (12, 10, 5),
    }
    for name, (size, before, after) in tokens.items():
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = BLACK
        style.font.bold = name != "Subtitle"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.font.color.rgb = BLACK
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25


def force_all_text_colors_black(document: Document):
    """Remove inherited Word theme colors so every explicit text color is black."""

    roots = [document.element, document.styles.element]
    for section in document.sections:
        roots.extend((section.header._element, section.footer._element))
    for root in roots:
        for color in root.iter(qn("w:color")):
            color.set(qn("w:val"), "000000")
            for attribute in ("themeColor", "themeTint", "themeShade"):
                color.attrib.pop(qn(f"w:{attribute}"), None)


def scrub_saved_ooxml_text_colors(path: Path):
    """Blacken text colors in auxiliary style parts not exposed by python-docx."""

    temporary = path.with_name(f"{path.name}.color-scrub.tmp")
    color_tag = re.compile(rb"<w:color\b[^>]*/?>")

    def blacken(match: re.Match[bytes]) -> bytes:
        tag = match.group(0)
        tag = re.sub(rb'w:val="[^"]*"', b'w:val="000000"', tag)
        tag = re.sub(rb'\s+w:theme(?:Color|Tint|Shade)="[^"]*"', b"", tag)
        return tag

    with ZipFile(path, "r") as source, ZipFile(temporary, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename.endswith(".xml"):
                content = color_tag.sub(blacken, content)
            target.writestr(item, content)
    os.replace(temporary, path)


def add_hyperlink(paragraph, text: str, url: str, *, size: float | None = None):
    relationship_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = paragraph.add_run(text)
    set_run_font(run, size=size)
    run.font.underline = True
    paragraph._p.remove(run._r)
    hyperlink.append(run._r)
    paragraph._p.append(hyperlink)


def add_inline_markdown(paragraph, text: str, *, size: float | None = None):
    pattern = re.compile(r"(\[[^\]]+\]\(https?://[^)]+\)|`[^`]+`|\*\*[^*]+\*\*)")
    cursor = 0
    for match in pattern.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor : match.start()])
            set_run_font(run, size=size)
        token = match.group(0)
        if token.startswith("["):
            link_match = re.match(r"\[([^\]]+)\]\((https?://[^)]+)\)", token)
            if link_match is None:
                raise ValueError(f"Invalid Markdown hyperlink: {token}")
            add_hyperlink(paragraph, link_match.group(1), link_match.group(2), size=size)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, name="Courier New", size=(size or 11) - 0.5)
        else:
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, bold=True)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run, size=size)


def allocate_widths(rows: list[list[str]], total_width: int) -> list[int]:
    columns = len(rows[0])
    if columns == 2:
        first_lengths = [len(row[0]) for row in rows]
        if max(first_lengths, default=0) < 38:
            return [2700, total_width - 2700]
    weights = []
    for index in range(columns):
        maximum = max(len(row[index]) for row in rows)
        weights.append(max(8.0, min(34.0, maximum ** 0.65)))
    widths = [max(650, int(total_width * weight / sum(weights))) for weight in weights]
    difference = total_width - sum(widths)
    widths[-1] += difference
    return widths


def add_table(document: Document, rows: list[list[str]], *, landscape: bool):
    total_width = USABLE_LANDSCAPE_DXA if landscape else USABLE_PORTRAIT_DXA
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    widths = allocate_widths(rows, total_width)
    for row_index, values in enumerate(rows):
        for column_index, value in enumerate(values):
            cell = table.cell(row_index, column_index)
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
                if column_index > 0 and len(value) < 28
                else WD_ALIGN_PARAGRAPH.LEFT
            )
            add_inline_markdown(paragraph, value, size=8.3 if landscape else 9)
            for run in paragraph.runs:
                run.bold = row_index == 0
    mark_repeat_header(table.rows[0])
    set_table_geometry(table, widths, total_width)
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(2)


def parse_markdown(document: Document, lines: list[str]):
    index = 0
    first_title_skipped = False
    while index < len(lines):
        stripped = lines[index].strip()
        if not stripped:
            index += 1
            continue
        if stripped.startswith("|") and index + 1 < len(lines) and set(lines[index + 1].replace("|", "").replace("-", "").replace(":", "").strip()) == set():
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            parsed = [[cell.strip() for cell in line.strip("|").split("|")] for line in table_lines]
            rows = [parsed[0]] + parsed[2:]
            landscape = len(rows[0]) > 6
            if landscape:
                # Keep an immediately preceding heading with its landscape
                # table instead of orphaning the heading on a portrait page.
                carried_heading = None
                if document.paragraphs:
                    candidate = document.paragraphs[-1]
                    if candidate.style and candidate.style.name.startswith("Heading"):
                        carried_heading = (candidate.text, candidate.style.name)
                        candidate._element.getparent().remove(candidate._element)
                section = document.add_section(WD_SECTION.NEW_PAGE)
                configure_section(section, landscape=True)
                if carried_heading is not None:
                    heading_text, heading_style = carried_heading
                    paragraph = document.add_paragraph(style=heading_style)
                    add_inline_markdown(paragraph, heading_text)
            add_table(document, rows, landscape=landscape)
            if landscape:
                section = document.add_section(WD_SECTION.NEW_PAGE)
                configure_section(section, landscape=False)
            continue
        if stripped.startswith("# "):
            if first_title_skipped:
                paragraph = document.add_paragraph(style="Heading 1")
                add_inline_markdown(paragraph, stripped[2:])
            first_title_skipped = True
        elif stripped.startswith("## "):
            paragraph = document.add_paragraph(style="Heading 1")
            add_inline_markdown(paragraph, stripped[3:])
        elif stripped.startswith("### "):
            paragraph = document.add_paragraph(style="Heading 2")
            add_inline_markdown(paragraph, stripped[4:])
        elif re.match(r"^\d+\.\s+", stripped):
            paragraph = document.add_paragraph(style="List Number")
            add_inline_markdown(paragraph, re.sub(r"^\d+\.\s+", "", stripped))
        elif stripped.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline_markdown(paragraph, stripped[2:])
        else:
            paragraph = document.add_paragraph()
            add_inline_markdown(paragraph, stripped)
        index += 1


def add_cover(document: Document):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(36)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run("FINE-TUNING AND ROI EXPERIMENT RECORD")
    set_run_font(run, size=11, bold=True)

    title = document.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(8)
    add_inline_markdown(title, "DenseNet-201 Pneumonia Classification", size=26)
    for run in title.runs:
        run.bold = True
    subtitle = document.add_paragraph(style="Subtitle")
    add_inline_markdown(subtitle, "VIRIDIS -> Global Histogram Equalization -> CLAHE", size=14)

    metadata = (
        ("Scope", "FT-001 through FT-010 and final FT-004 selection; chronological fine-tuning, ROI, hierarchy, confirmation, and model-freeze record"),
        ("Primary metric", "Validation macro F1"),
        ("Development data", "Locked training and validation assignments"),
        ("Locked test", "Not used for tuning, model selection, or reported metrics"),
        ("Document date", "13 August 2026"),
        ("Visual style", "Black-only technical report"),
    )
    table = document.add_table(rows=len(metadata), cols=2)
    table.style = "Table Grid"
    for row_index, (label, value) in enumerate(metadata):
        for column_index, text in enumerate((label, value)):
            cell = table.cell(row_index, column_index)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_inline_markdown(p, text, size=10)
            if column_index == 0:
                p.runs[0].bold = True
    set_table_geometry(table, [2400, USABLE_PORTRAIT_DXA - 2400], USABLE_PORTRAIT_DXA)
    document.add_page_break()


def build(source: Path, output: Path):
    document = Document()
    configure_styles(document)
    configure_section(document.sections[0], landscape=False)
    add_cover(document)
    lines = source.read_text(encoding="utf-8").splitlines()
    parse_markdown(document, lines)
    properties = document.core_properties
    properties.title = "DenseNet-201 Fine-Tuning and ROI Experiment Record"
    properties.subject = "Chronological FT-001 through FT-010 documentation and final FT-004 three-member ensemble selection"
    properties.author = "Pneumonia Classification Research Project"
    properties.keywords = "DenseNet-201, fine-tuning, PSPNet, U-Net, ROI, validation error analysis, hierarchical classification, VIRIDIS, histogram equalization, CLAHE"
    force_all_text_colors_black(document)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    scrub_saved_ooxml_text_colors(output)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--source",
        type=Path,
        default=Path("model/docs/fine_tuning/DenseNet201_VIRIDIS_GE_CLAHE_Fine_Tuning_Journal.md"),
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("model/docs/fine_tuning/DenseNet201_VIRIDIS_GE_CLAHE_Fine_Tuning_Report.docx"),
    )
    args = parser.parse_args()
    build(args.source.resolve(), args.output.resolve())
    print(args.output.resolve())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
