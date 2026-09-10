"""Generate the final Markdown and Word reports from completed real artifacts.

The generator fails closed until baselines, preprocessing, all three tuning
studies, three robustness seeds, the frozen locked-test evaluation, and the
consolidation consistency audit are complete. It never manufactures missing
values and writes a machine-readable document-package audit beside the report.
"""

from __future__ import annotations

import argparse
import csv
import json
import math
import statistics
import zipfile
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

import _multiclass_bootstrap
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

import generate_progress_report as progress
from document_io import save_docx_with_fallback


REPO_ROOT = _multiclass_bootstrap.REPO_ROOT
MODEL_ROOT = REPO_ROOT / "model"
RESULTS = MODEL_ROOT / "results"
DOCS = MODEL_ROOT / "docs"
MARKDOWN_OUTPUT = DOCS / "Pneumonia_Three_Class_Classification_Final_Report.md"
DOCX_OUTPUT = DOCS / "Pneumonia_Three_Class_Classification_Final_Report.docx"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--check-only", action="store_true", help="Report gate status without writing final reports")
    return parser.parse_args()


def read_csv(path: Path) -> list[dict[str, str]]:
    if not path.exists() or not path.stat().st_size:
        return []
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def final_gate() -> list[str]:
    failures: list[str] = []
    audit_path = RESULTS / "consolidation_audit.json"
    audit = json.loads(audit_path.read_text(encoding="utf-8")) if audit_path.exists() else {}
    if audit.get("consistency_error_count") != 0:
        failures.append("consolidation consistency audit is missing or nonzero")
    if audit.get("completed_baselines") != 8:
        failures.append(f"completed supported baselines must be 8, found {audit.get('completed_baselines')}")
    summary = read_csv(RESULTS / "summary.csv")
    completed_ablation = [
        row for row in summary if row.get("stage") == "preprocessing_ablation" and row.get("status") == "completed"
    ]
    if len(completed_ablation) != 9:
        failures.append(f"completed supported preprocessing ablations must be 9, found {len(completed_ablation)}")
    for index in (1, 2, 3):
        path = RESULTS / "tuning" / f"tuning_top{index}" / "study_summary.json"
        if not path.exists():
            failures.append(f"tuning_top{index} study summary is missing")
    robustness = [row for row in summary if row.get("stage") == "robustness" and row.get("status") == "completed"]
    if len(robustness) != 3 or {row.get("seed") for row in robustness} != {"42", "123", "2026"}:
        failures.append("robustness runs at seeds 42, 123, and 2026 are not all completed")
    locked = [row for row in summary if row.get("run_id") == "final_locked_test" and row.get("status") == "completed"]
    if len(locked) != 1:
        failures.append("final_locked_test is not completed")
    selection_path = RESULTS / "final_selection.json"
    selection = json.loads(selection_path.read_text(encoding="utf-8")) if selection_path.exists() else {}
    if selection.get("status") != "locked_test_evaluated" or not selection.get("test_metrics"):
        failures.append("final selection lock does not contain the completed single test evaluation")
    return failures


def f(value: Any, digits: int = 4) -> str:
    try:
        number = float(value)
    except (TypeError, ValueError):
        return "not available"
    if not math.isfinite(number):
        return "not available"
    return f"{number:.{digits}f}"


def md_table(columns: list[tuple[str, str]], rows: list[dict[str, Any]], *, limit: int | None = None) -> str:
    visible = rows[:limit] if limit is not None else rows
    lines = [
        "| " + " | ".join(label for label, _ in columns) + " |",
        "| " + " | ".join("---" for _ in columns) + " |",
    ]
    for row in visible:
        values = []
        for _, key in columns:
            value = row.get(key, "")
            rendered = str(value).replace("|", "/").replace("\n", " ")
            values.append(rendered if rendered else "-")
        lines.append("| " + " | ".join(values) + " |")
    return "\n".join(lines)


def summarize_errors(prediction_path: Path) -> dict[str, Any]:
    rows = read_csv(prediction_path)
    incorrect = [row for row in rows if row["true_class"] != row["predicted_class"]]
    directions = Counter(f"{row['true_class']} -> {row['predicted_class']}" for row in incorrect)
    low_confidence = sorted(rows, key=lambda row: float(row["confidence"]))[:10]
    high_confidence_errors = sorted(incorrect, key=lambda row: -float(row["confidence"]))[:10]
    return {
        "total": len(rows),
        "incorrect": len(incorrect),
        "directions": directions,
        "low_confidence": low_confidence,
        "high_confidence_errors": high_confidence_errors,
    }


def build_markdown() -> tuple[str, dict[str, Any]]:
    generated = datetime.now().astimezone().isoformat(timespec="seconds")
    split = json.loads((MODEL_ROOT / "data_manifests/split_summary.json").read_text(encoding="utf-8"))
    environment = json.loads((RESULTS / "environment.json").read_text(encoding="utf-8"))
    multiclass_environment = environment.get("multiclass_environment", {})
    summary = read_csv(RESULTS / "summary.csv")
    matrix = read_csv(RESULTS / "run_matrix.csv")
    ranking = read_csv(RESULTS / "model_ranking.csv")
    ablation = read_csv(RESULTS / "preprocessing_ablation.csv")
    trials = read_csv(RESULTS / "tuning_trials.csv")
    robustness = read_csv(RESULTS / "seed_robustness.csv")
    test_rows = read_csv(RESULTS / "test_results.csv")
    failures = read_csv(RESULTS / "failures.csv")
    selection = json.loads((RESULTS / "final_selection.json").read_text(encoding="utf-8"))
    test_metrics = selection["test_metrics"]
    selected_run_dir = Path(selection["selected_run_dir"])
    selected_config = json.loads((selected_run_dir / "config.json").read_text(encoding="utf-8"))
    errors = summarize_errors(RESULTS / "runs/final_locked_test/predictions.csv")
    completed_trials = [row for row in trials if row.get("status") == "completed"]
    unsupported = [row for row in matrix if row.get("status") in {"unsupported", "skipped"}]
    active_failures = [row for row in failures if row.get("resolved", "").lower() != "true"]
    robustness_scores = [float(row["validation_macro_f1"]) for row in robustness]

    baseline_columns = [
        ("Rank", "rank"), ("Run", "run_id"), ("Architecture", "architecture"),
        ("Macro F1", "validation_macro_f1"), ("Balanced acc.", "validation_balanced_accuracy"),
        ("Normal recall", "recall_normal"), ("Viral recall", "recall_viral"),
        ("Bacterial recall", "recall_bacterial"),
    ]
    matrix_columns = [
        ("Run", "run_id"), ("Stage", "stage"), ("Architecture", "architecture"),
        ("Preprocessing", "preprocessing"), ("Seed", "seed"), ("Status", "status"),
        ("Reason", "reason"),
    ]
    trial_columns = [
        ("Study", "study_id"), ("Trial", "trial_number"), ("Run", "run_id"),
        ("Status", "status"), ("Macro F1", "validation_macro_f1"),
        ("Balanced acc.", "validation_balanced_accuracy"), ("Parameters", "parameters_json"),
    ]
    resource_columns = [
        ("Run", "run_id"), ("Params", "parameter_count"), ("File bytes", "model_file_size_bytes"),
        ("Peak GPU bytes", "peak_gpu_memory_bytes"), ("Latency ms/image", "latency_ms_per_image"),
        ("Throughput img/s", "throughput_images_per_second"),
    ]
    error_rows = [
        {"direction": direction, "count": count}
        for direction, count in errors["directions"].most_common()
    ]

    lines = [
        "# Pneumonia Three-Class Classification Final Report",
        "",
        f"**Document version:** 1.0  ",
        f"**Generated:** {generated}  ",
        "**Research-use status:** Final report from the completed locked-manifest experiment workflow.",
        "",
        "## Executive Summary",
        "",
        f"This work preserved the legacy binary classifier and built a parallel reproducible classifier for Normal, Viral Pneumonia, and Bacterial Pneumonia. The final validation-selected configuration is `{selection['selected_run_id']}` using {selection['architecture']} with {selection['preprocessing']} preprocessing. Across robustness seeds 42, 123, and 2026, validation macro F1 was {f(statistics.mean(robustness_scores))} +/- {f(statistics.stdev(robustness_scores))} (sample standard deviation). The single frozen locked-test evaluation produced macro F1 {f(test_metrics['macro_f1'])}, balanced accuracy {f(test_metrics['balanced_accuracy'])}, and accuracy {f(test_metrics['accuracy'])}.",
        "",
        "The system remains a research classifier. It has not undergone external clinical validation, prospective evaluation, bias assessment across institutions or populations, calibration review for clinical decisions, regulatory review, or medical-device quality assurance.",
        "",
        "## Research Background and Objective",
        "",
        "The repository originally implemented binary Normal-versus-Pneumonia transfer learning. The objective was to audit that behavior first, derive defensible three-class labels from the actual source, prevent duplicate/group leakage, compare requested supported architectures fairly, evaluate conservative preprocessing choices, tune selected configurations without test feedback, measure seed variability, and evaluate one frozen model once on the locked test split.",
        "",
        "## Class Definition and Diagram Interpretation",
        "",
        "The locked output order is Normal = 0, Viral = 1, Bacterial = 2. Normal is derived from canonical NORMAL folders; Viral and Bacterial are derived only from explicit `_virus_` and `_bacteria_` filename evidence inside canonical PNEUMONIA folders.",
        "",
        "The supplied research concepts were interpreted against repository evidence. CLAHE, global histogram equalization, and pseudo-coloring were implementable as independent transforms. Raw high-resolution PCA was deferred because no justified train-only feature-space design was supplied. FILM, SAS Tuning, and automated chest-X-ray grading remained unresolved and were not invented. Optuna TPE is reported only as general tuning.",
        "",
        "## Existing Binary-System Audit",
        "",
        "The legacy TensorFlow/Keras workflow uses DenseNet121/169/201 with frozen ImageNet backbones and a one-unit sigmoid head. It trains from a generated binary folder tree, lacks locked manifests and resumable checkpoints, and is vulnerable to split/duplicate issues. Its source, HDF5 models, output reports, and FastAPI binary contract were preserved. The complete evidence is in `model/docs/00_existing_binary_pipeline_audit.md`.",
        "",
        "## Repository and Implementation Overview",
        "",
        "The new `model/src/multiclass` package supplies configuration validation, locked-manifest loading, transforms, model adapters, training, metrics, artifact generation, tracking, evaluation, and inference. Every long run writes atomic status/config data, immediate epoch history, best/last checkpoints, a readable log, validation predictions, plots, and a journal entry. Windows data workers are configurable and default to zero. CUDA is required unless CPU fallback is explicitly enabled.",
        "",
        "## Development Process Timeline",
        "",
        "The append-only development log records the Phase 0 environment snapshot, legacy audit, immutable dataset audit, implementation, smoke/recovery validation, preprocessing evidence, resource estimates, controlled matrix, each benchmark stage, tuning, robustness, final test, consolidation, and reporting. Source: `model/docs/DEVELOPMENT_LOG.md`. Epoch-level events and checkpoint changes are in `model/docs/TRAINING_JOURNAL.md`.",
        "",
        "## Major Engineering and Research Decisions",
        "",
        "Stable decisions in `model/docs/DECISION_LOG.md` cover label evidence, conservative grouping, exact duplicate exclusion, environment isolation, model support, class order, primary metric, balanced loss, measured batches, checkpoint retention, preprocessing scope, bounded Optuna search, three independent robustness seeds, and predeclared seed-42 test representation.",
        "",
        "## Problems Encountered and Verified Fixes",
        "",
        f"The issue log preserves dependency, integration, rendering, and third-party model-adapter problems. The failure ledger contains {len(failures)} rows, of which {len(active_failures)} remain marked unresolved. Superseded smoke failures were retained instead of deleted. Source: `model/docs/ISSUES_AND_FIXES.md` and `model/results/failures.csv`.",
        "",
        "## Dataset Source, Counts, and Label Derivation",
        "",
        f"The immutable source inventory contains 17,591 physical files. The audit identified 5,856 canonical image candidates, excluded 32 redundant exact canonical copies, and retained 5,824 unique cryptographic images. The locked split version is `{split['manifest_version']}` with hash `{split['manifest_version_hash']}`.",
        "",
        "| Split | Normal | Viral | Bacterial | Total |",
        "| --- | --- | --- | --- | --- |",
        "| Train | 1105 | 1039 | 1932 | 4076 |",
        "| Validation | 237 | 223 | 414 | 874 |",
        "| Test | 237 | 223 | 414 | 874 |",
        "",
        "## Duplicate, Leakage, and Split Analysis",
        "",
        "The physical source contains a complete nested copy tree and 32 redundant exact files inside the canonical tree. Exact hashes and conservative filename-derived groups were kept within one assigned split. Assigned exact-hash leakage and assigned group leakage are both zero. Sixteen pHash candidate pairs were retained for review rather than automatically treated as identity. No authoritative patient metadata was available, so the grouping is conservative and is not claimed as proven patient identity.",
        "",
        "## Architecture Coverage",
        "",
        "DenseNet121, DenseNet169, and DenseNet201 use dense feature reuse; ResNet50, ResNet101, and ResNet152 use residual blocks. Torchvision ImageNet weights initialize all six. Official pinned YOLOv5s-cls and Ultralytics YOLOv8n-cls use genuine classification heads adapted to the common three-logit interface. YOLOv10 classification is unsupported because Ultralytics 8.3.168 provides YOLOv10 detection heads but no genuine YOLOv10 classification YAML/checkpoint; no detector was mislabeled as a classifier.",
        "",
        "## Preprocessing Methods",
        "",
        "Baseline uses safe RGB conversion, centered square fit, tensor conversion, and ImageNet normalization. CLAHE operates on grayscale luminance with clip limit 2.0 and 8x8 tiles. Global histogram equalization is distinct from CLAHE and uses grayscale equalization. Pseudo-coloring applies OpenCV VIRIDIS to grayscale and does not add clinical information. Training-only augmentation is conservative horizontal flip and rotation. No learned preprocessing is fitted on validation or test data.",
        "",
        "## Training Configuration",
        "",
        f"The final configuration uses `{selected_config['architecture']}`, `{selected_config['preprocessing']}` preprocessing, {selected_config['input_size']}x{selected_config['input_size']} input, batch {selected_config['batch_size']}, `{selected_config['optimizer']}`, learning rate {selected_config['learning_rate']}, weight decay {selected_config['weight_decay']}, `{selected_config['scheduler']}` scheduling, `{selected_config['loss']}` loss, class weighting `{selected_config['class_weighting']}`, AMP `{selected_config['amp']}`, and at most {selected_config['epochs']} epochs with early-stopping patience {selected_config['early_stopping_patience']}.",
        "",
        "## Hyperparameter-Tuning Methodology",
        "",
        f"Three validation-selected parent configurations received separate Optuna 4.9.0 TPE studies. Four equal-budget six-epoch trials per study searched optimizer-specific learning rate, weight decay, optimizer, scheduler, cross-entropy versus focal loss, class weighting, label smoothing, and conservative versus disabled augmentation. {len(completed_trials)} trials completed; failed/pruned states remain in the table. The small search budget is a hardware-driven limitation.",
        "",
        "## Hardware and Software Environment",
        "",
        f"Training used the audited NVIDIA GeForce RTX 2050 with 4,096 MiB VRAM. The isolated environment uses Python {multiclass_environment.get('python_version', 'not available')}, PyTorch {multiclass_environment.get('packages', {}).get('torch', 'not available')}, torchvision {multiclass_environment.get('packages', {}).get('torchvision', 'not available')}, cuDNN {multiclass_environment.get('gpu', {}).get('cudnn_version', 'not available')}, Ultralytics {multiclass_environment.get('packages', {}).get('ultralytics', 'not available')}, and Optuna {multiclass_environment.get('packages', {}).get('optuna', 'not available')}. Exact machine-readable details are in `model/results/environment.json` and the pinned requirements file.",
        "",
        "## Complete Experiment Matrix",
        "",
        md_table(matrix_columns, matrix),
        "",
        "## Training-Progress Summary",
        "",
        f"The final matrix contains {len(matrix)} rows spanning diagnostics, eight supported baselines, one unsupported YOLOv10 row, preprocessing ablation, focused comparison, three tuning studies, three independent robustness seeds, and one locked-test evaluation. All epoch histories and checkpoint changes remain available under `model/results/runs` and in the training journal.",
        "",
        "## Main Validation Results and Ranked Model Comparison",
        "",
        md_table(baseline_columns, ranking),
        "",
        "Macro F1 is primary. Balanced accuracy, minimum per-class recall, validation loss/calibration, and then efficiency break close ties; accuracy alone was never the selection rule.",
        "",
        "## Preprocessing Ablation",
        "",
        md_table(baseline_columns[1:], ablation),
        "",
        "## Tuning Results",
        "",
        md_table(trial_columns, trials),
        "",
        "## Multi-Seed Robustness",
        "",
        md_table(baseline_columns[1:], robustness),
        "",
        f"Validation macro F1 mean is {f(statistics.mean(robustness_scores))}; sample standard deviation is {f(statistics.stdev(robustness_scores))}.",
        "",
        "## Locked Test Results",
        "",
        f"The single frozen test evaluation used `{selection['selected_run_id']}` by the predeclared seed-42 policy. Macro F1 is {f(test_metrics['macro_f1'])}, balanced accuracy is {f(test_metrics['balanced_accuracy'])}, accuracy is {f(test_metrics['accuracy'])}, macro ROC-AUC is {f(test_metrics['roc_auc_ovr_macro'])}, macro PR-AUC is {f(test_metrics['pr_auc_macro'])}, log loss is {f(test_metrics['log_loss'])}, ECE is {f(test_metrics['expected_calibration_error'])}, and multiclass Brier score is {f(test_metrics['brier_score'])}.",
        "",
        "| Class | Precision | Recall/sensitivity | Specificity | F1 | Support | ROC-AUC | PR-AUC |",
        "| --- | --- | --- | --- | --- | --- | --- | --- |",
    ]
    for name in ("Normal", "Viral", "Bacterial"):
        row = test_metrics["per_class"][name]
        lines.append(
            f"| {name} | {f(row['precision'])} | {f(row['recall'])} | {f(row['specificity'])} | "
            f"{f(row['f1'])} | {int(row['support'])} | {f(row['roc_auc'])} | {f(row['pr_auc'])} |"
        )
    lines.extend(
        [
            "",
            "## Efficiency and Model-Size Trade-offs",
            "",
            md_table(resource_columns, ranking),
            "",
            "Latency values use batch one on CUDA after warmup and exclude preprocessing and host-to-device transfer. They are controlled diagnostics, not end-to-end clinical deployment claims.",
            "",
            "## Error Analysis",
            "",
            f"The final test set contains {errors['total']} images and {errors['incorrect']} incorrect predictions. Confusion directions are:",
            "",
            md_table([("Direction", "direction"), ("Count", "count")], error_rows),
            "",
            "The highest-confidence incorrect predictions and lowest-confidence predictions remain in `model/results/runs/final_locked_test/sample_errors.csv` and `predictions.csv` with de-identified sample IDs and relative source paths. Viral-versus-Bacterial errors are interpreted cautiously because filename-derived etiology and single-source label quality may be imperfect.",
            "",
            "## Final Selected Model and Justification",
            "",
            f"The final configuration is {selection['architecture']} with {selection['preprocessing']} preprocessing. It was selected before test access from validation macro F1 with declared tie-breakers, then independently retrained across three seeds. The seed-42 checkpoint was predeclared as representative rather than selected after viewing test outcomes. Final checkpoint: `{selection['selected_checkpoint']}`.",
            "",
            "## Limitations",
            "",
            "The dataset is a single public source with filename-derived Viral/Bacterial labels and no authoritative patient metadata. Conservative grouping can merge unrelated reused identifiers, while unrecognized patient relations may remain. Class balance is imperfect, source/device diversity is limited, label noise is plausible, the tuning budget is small, and no external dataset is evaluated. Confidence calibration requires domain-specific review. The versioned research-complete Word deliverable is subject to a separate rendered-page visual QA record.",
            "",
            "## Clinical-Safety Disclaimer",
            "",
            "This classifier is for research only and is not validated, approved, or suitable for clinical diagnosis, triage, treatment, or deployment. Clinical use would require external and prospective validation, broader representative datasets, subgroup bias analysis, calibration and threshold review, human-factors testing, medical oversight, security/privacy review, and applicable regulatory and quality-system processes.",
            "",
            "## Recommended Future Work",
            "",
            "Obtain authoritative patient identifiers and radiologist-reviewed etiology labels; evaluate external institutions and acquisition devices; study label noise; expand tuning under a preregistered compute budget; examine feature-space PCA only with train-only fitting; define FILM or grading concepts before implementation; evaluate ensembling prospectively; conduct subgroup and calibration analyses; and integrate the API only after its three-class contract and clinical-use restrictions are reviewed.",
            "",
            "## Reproduction Commands",
            "",
            "Exact Windows commands for environment setup, audits, manifests, tests, training, resume, matrix execution, tuning, frozen test evaluation, consolidation, and report generation are in `model/docs/06_reproduction_guide.md`. The selected seed-42 training command is:",
            "",
            "```text",
            ".venv\\Scripts\\python.exe model\\scripts\\train_multiclass.py --config model\\configs\\final_selected.json --seed 42 --run-id final_selected_s42",
            "```",
            "",
            "## Output-File Inventory",
            "",
            "Primary outputs are the locked manifests, environment/support audits, run matrix and ledgers, per-run directories, consolidated CSV/JSON files, continuous Markdown documentation, progressive Word report, this final Markdown report, and the final Word report. Paths are rooted under `model/data_manifests`, `model/results`, and `model/docs`.",
            "",
            "## Appendix A - Final Configuration",
            "",
            "```json",
            json.dumps(selected_config, indent=2),
            "```",
            "",
            "## Appendix B - Metric Definitions",
            "",
            "Macro F1 is the unweighted mean of class F1 values. Balanced accuracy is mean class recall. Specificity is one-vs-rest true-negative rate. ROC-AUC and PR-AUC are one-vs-rest. ECE uses confidence bins; the multiclass Brier score sums squared probability error across classes. Reported standard deviation across seeds is the sample standard deviation.",
            "",
            "## Appendix C - Failed, Unsupported, Skipped, and Unresolved Work",
            "",
            md_table(matrix_columns, unsupported),
            "",
            f"The failure ledger contains {len(failures)} preserved entries. FILM, SAS Tuning, and automated grading remain unresolved. Raw-image PCA and YOLOv10 classification remain unsupported for the documented technical reasons; no result is fabricated for them.",
            "",
            "## Appendix D - Source Artifacts",
            "",
            "All numeric claims derive from `model/results/summary.csv`, `model/results/model_ranking.csv`, `model/results/preprocessing_ablation.csv`, `model/results/tuning_trials.csv`, `model/results/seed_robustness.csv`, `model/results/test_results.csv`, the final per-run metrics/predictions, and the locked manifest summary. Consolidation verified those values against original per-run metrics with zero consistency errors.",
            "",
        ]
    )
    context = {
        "generated_at": generated,
        "selected_run_dir": str(selected_run_dir),
        "selected_run_id": selection["selected_run_id"],
        "test_metrics": test_metrics,
        "baseline_ranking": ranking,
    }
    return "\n".join(lines), context


def add_contents(document: Document) -> None:
    document.add_paragraph("Contents", style="Heading 1")
    items = [
        "Executive summary and objective",
        "Audit, dataset, split, and implementation",
        "Architectures, preprocessing, training, tuning, and environment",
        "Experiment matrix and validation results",
        "Robustness and locked-test evaluation",
        "Efficiency, calibration, and error analysis",
        "Final selection, limitations, safety, and future work",
        "Reproduction, inventory, and appendices",
        "Evidence figures",
    ]
    for item in items:
        document.add_paragraph(item, style="List Number")


def add_figure(document: Document, path: Path, caption: str, number: int, *, width: float = 6.2) -> int:
    if not path.exists():
        raise FileNotFoundError(f"final report figure is missing: {path}")
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(path), width=Inches(width))
    caption_p = document.add_paragraph(style="Caption")
    caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_p.add_run(f"Figure {number}. {caption}. Source: {path.relative_to(REPO_ROOT).as_posix()}")
    return number + 1


def build_docx(markdown: str, context: dict[str, Any]) -> dict[str, Any]:
    document = Document()
    document.core_properties.title = "Pneumonia Three-Class Classification Final Report"
    document.core_properties.subject = "Completed locked-manifest research workflow"
    document.core_properties.author = "Pneumonia Research Pipeline"
    progress.configure_styles(document)
    progress.configure_page(document)
    progress.add_cover(document, context["generated_at"], "Final report - completed research workflow")
    add_contents(document)
    document.add_page_break()
    progress.render_markdown(document, markdown)

    heading = document.add_paragraph("Evidence Figures", style="Heading 1")
    heading.paragraph_format.page_break_before = True
    note = document.add_paragraph()
    run = note.add_run(
        "Validation-selection figures and the final locked-test figures are labeled separately. "
        "All paths point to persisted source artifacts."
    )
    progress.set_run_font(run, color=RGBColor(0x66, 0x66, 0x66))
    number = 1
    number = add_figure(
        document,
        MODEL_ROOT / "results/dataset_audit/class_counts.png",
        "Locked-manifest class counts by split",
        number,
    )
    number = add_figure(
        document,
        MODEL_ROOT / "results/preprocessing_examples/preprocessing_comparison.png",
        "Validation-only preprocessing comparison",
        number,
    )
    for row in context["baseline_ranking"][:3]:
        run_dir = Path(row["run_dir"])
        number = add_figure(document, run_dir / "training_curves.png", f"Validation rank {row['rank']} training curves ({row['run_id']})", number)
        number = add_figure(document, run_dir / "confusion_matrix.png", f"Validation rank {row['rank']} confusion matrix ({row['run_id']})", number, width=5.8)
    selected_dir = Path(context["selected_run_dir"])
    number = add_figure(document, selected_dir / "training_curves.png", "Final representative seed-42 training curves", number)
    final_dir = RESULTS / "runs/final_locked_test"
    for filename, caption, width in (
        ("confusion_matrix.png", "Final locked-test confusion matrix", 5.8),
        ("roc_curves.png", "Final locked-test one-vs-rest ROC curves", 6.2),
        ("pr_curves.png", "Final locked-test precision-recall curves", 6.2),
        ("calibration_curve.png", "Final locked-test confidence calibration", 5.8),
    ):
        number = add_figure(document, final_dir / filename, caption, number, width=width)
    saved_output = save_docx_with_fallback(document, DOCX_OUTPUT)

    with zipfile.ZipFile(saved_output) as archive:
        bad = archive.testzip()
        names = archive.namelist()
        xml = ElementTree.fromstring(archive.read("word/document.xml"))
    namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    style_attribute = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val"
    heading_1 = []
    section_headings = []
    for paragraph in xml.findall(".//w:p", namespace):
        style = paragraph.find("w:pPr/w:pStyle", namespace)
        if style is None:
            continue
        style_name = style.get(style_attribute)
        if style_name == "Heading1":
            heading_1.append(paragraph)
        if style_name in {"Heading1", "Heading2"}:
            section_headings.append(paragraph)
    audit = {
        "generated_at": context["generated_at"],
        "requested_docx_path": str(DOCX_OUTPUT),
        "docx_path": str(saved_output),
        "docx_size_bytes": saved_output.stat().st_size,
        "zip_integrity_error": bad,
        "package_entry_count": len(names),
        "embedded_media_count": len([name for name in names if name.startswith("word/media/")]),
        "heading_1_count": len(heading_1),
        "section_heading_count": len(section_headings),
        "visual_render_verified": False,
        "visual_render_limitation": "LibreOffice unavailable and unattended Word COM rendering previously hung",
    }
    (DOCS / "final_report_docx_audit.json").write_text(json.dumps(audit, indent=2) + "\n", encoding="utf-8")
    # Markdown ``##`` sections intentionally render as Heading 2 beneath the
    # report-level Heading 1 blocks. Audit the complete section hierarchy
    # rather than incorrectly requiring eight top-level Heading 1 paragraphs.
    if bad is not None or audit["embedded_media_count"] < 11 or audit["section_heading_count"] < 8:
        raise RuntimeError(f"final DOCX structural audit failed: {audit}")
    return audit


def main() -> int:
    args = parse_args()
    failures = final_gate()
    if failures:
        print(json.dumps({"ready": False, "gate_failures": failures}, indent=2))
        return 1
    if args.check_only:
        print(json.dumps({"ready": True, "gate_failures": []}, indent=2))
        return 0
    markdown, context = build_markdown()
    MARKDOWN_OUTPUT.write_text(markdown, encoding="utf-8")
    audit = build_docx(markdown, context)
    print(json.dumps({"markdown": str(MARKDOWN_OUTPUT), "docx": audit["docx_path"], "audit": audit}, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
