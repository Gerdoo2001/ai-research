"""Append the second-batch image flow and complete training tables to Word."""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[3]
WORK = Path(__file__).resolve().parent
RUNS = ROOT / "model" / "results" / "runs"
REFERENCE = WORK / "reference_copy.docx"
FIGURES = WORK / "figures"
PREVIEW = WORK / "second_batch_section_preview.docx"
FINAL = (
    ROOT
    / "model"
    / "docs"
    / "Pneumonia_Model_Training_Details_Black_Second_Batch_Complete_20260804.docx"
)

MAIN_RUNS = [
    "baseline_densenet121_s42",
    "baseline_densenet169_s42",
    "ablation_densenet169_clahe_s42",
    "ablation_densenet169_histogram_equalization_s42",
    "ablation_densenet169_pseudo_color_s42",
    "baseline_densenet201_s42",
    "ablation_densenet201_clahe_s42",
    "ablation_densenet201_histogram_equalization_s42",
    "ablation_densenet201_pseudo_color_s42",
    "ablation_densenet201_histogram_equalization_then_clahe_s42",
    "ablation_densenet201_clahe_then_histogram_equalization_s42",
    "baseline_resnet50_s42",
    "ablation_resnet50_clahe_s42",
    "ablation_resnet50_histogram_equalization_s42",
    "ablation_resnet50_pseudo_color_s42",
    "baseline_resnet101_s42",
    "baseline_resnet152_s42",
    "baseline_yolov5_cls_s42",
    "baseline_yolov8_cls_s42",
    "experimental_densenet201_sas_mff_clahe_s42",
    "experimental_densenet201_sas_mff_histogram_equalization_s42",
]

TUNING_RUNS = [
    f"tuning_top{rank}_trial_{trial:03d}"
    for rank in (1, 2, 3)
    for trial in range(4)
]
SEED_RUNS = ["final_selected_s42", "final_selected_s123", "final_selected_s2026"]

ARCHITECTURE_LABELS = {
    "densenet121": "DenseNet-121",
    "densenet169": "DenseNet-169",
    "densenet201": "DenseNet-201",
    "densenet201_sas_mff": "DenseNet-201 + SAS/MFF",
    "resnet50": "ResNet-50",
    "resnet101": "ResNet-101",
    "resnet152": "ResNet-152",
    "yolov5_cls": "YOLOv5-cls",
    "yolov8_cls": "YOLOv8-cls",
}

PREPROCESSING_LABELS = {
    "baseline": "Base/original",
    "clahe": "CLAHE",
    "histogram_equalization": "Global HE",
    "pseudo_color": "VIRIDIS",
    "histogram_equalization_then_clahe": "GE -> CLAHE",
    "clahe_then_histogram_equalization": "CLAHE -> GE",
}


def set_run_font(run, size: float | None = None, bold: bool | None = None) -> None:
    run.font.name = "Arial"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Arial")
    run.font.color.rgb = RGBColor(0, 0, 0)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def configure_styles(doc: Document) -> None:
    specs = {
        "Normal": (10.5, False, 0, 6, 1.1),
        "Title": (28, True, 0, 8, 1.0),
        "Subtitle": (14, False, 0, 8, 1.0),
        "Heading 1": (16, True, 16, 8, 1.0),
        "Heading 2": (13, True, 12, 6, 1.0),
        "Heading 3": (11.5, True, 8, 4, 1.0),
        "Caption": (8.5, False, 0, 0, 1.0),
    }
    for name, (size, bold, before, after, spacing) in specs.items():
        style = doc.styles[name]
        style.font.name = "Arial"
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Arial")
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = spacing
    doc.styles["Subtitle"].font.italic = True


def configure_page(doc: Document) -> None:
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)


def add_body(doc: Document, text: str = "", *, bold_lead: str | None = None):
    p = doc.add_paragraph(style="Normal")
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        set_run_font(lead, bold=True)
        rest = p.add_run(text[len(bold_lead) :])
        set_run_font(rest)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_bullet(doc: Document, text: str, lead: str | None = None):
    p = doc.add_paragraph(style="List Bullet")
    if lead and text.startswith(lead):
        a = p.add_run(lead)
        set_run_font(a, bold=True)
        b = p.add_run(text[len(lead) :])
        set_run_font(b)
    else:
        r = p.add_run(text)
        set_run_font(r)
    p.paragraph_format.space_after = Pt(3)
    return p


def add_heading(doc: Document, text: str, level: int):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    for run in p.runs:
        set_run_font(run)
    return p


def add_figure(doc: Document, filename: str, width: float, caption: str, alt_text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    inline = run.add_picture(str(FIGURES / filename), width=Inches(width))
    inline._inline.docPr.set("descr", alt_text)
    inline._inline.docPr.set("title", caption.split(".", 1)[0])
    c = doc.add_paragraph(style="Caption")
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_after = Pt(6)
    cr = c.add_run(caption)
    set_run_font(cr, 8.5)


def read_run(run_id: str) -> dict[str, object]:
    run_dir = RUNS / run_id
    status = json.loads((run_dir / "status.json").read_text(encoding="utf-8"))
    config = json.loads((run_dir / "config.json").read_text(encoding="utf-8"))
    history_lines = (run_dir / "history.csv").read_text(encoding="utf-8").splitlines()
    headers = history_lines[0].split(",")
    rows = [dict(zip(headers, line.split(","))) for line in history_lines[1:] if line.strip()]
    best = next(row for row in rows if int(row["epoch"]) == int(status["best_epoch"]))
    return {"run_id": run_id, "status": status, "config": config, "best": best}


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 75, start: int = 120, bottom: int = 75, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_table_geometry(table, widths: list[int]) -> None:
    total = sum(widths)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "6")
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), "000000")


def format_table(table, widths: list[int], font_size: float = 8.0) -> None:
    set_table_geometry(table, widths)
    set_table_borders(table)
    set_repeat_header(table.rows[0])
    for row_index, row in enumerate(table.rows):
        if row_index == 0:
            for cell in row.cells:
                shade_cell(cell, "D9D9D9")
        for col_index, cell in enumerate(row.cells):
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                paragraph.alignment = (
                    WD_ALIGN_PARAGRAPH.LEFT if col_index < 2 else WD_ALIGN_PARAGRAPH.CENTER
                )
                for run in paragraph.runs:
                    set_run_font(run, font_size, bold=(row_index == 0))


def add_table_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="Caption")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    set_run_font(r, 8.5, bold=True)


def pct(value: str) -> str:
    return f"{100 * float(value):.3f}%"


def add_main_results_table(doc: Document) -> None:
    add_table_caption(doc, "Table 11.1. Complete main architecture and preprocessing training results (validation).")
    headers = ["Model", "Preprocessing", "Accuracy", "Balanced acc.", "Macro-F1", "Best epoch"]
    table = doc.add_table(rows=1, cols=len(headers))
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text
    for run_id in MAIN_RUNS:
        record = read_run(run_id)
        cfg = record["config"]
        best = record["best"]
        row = table.add_row().cells
        values = [
            ARCHITECTURE_LABELS[str(cfg["architecture"])],
            PREPROCESSING_LABELS[str(cfg["preprocessing"])],
            pct(str(best["val_accuracy"])),
            pct(str(best["val_balanced_accuracy"])),
            pct(str(best["val_macro_f1"])),
            str(record["status"]["best_epoch"]),
        ]
        for cell, text in zip(row, values):
            cell.text = text
    format_table(table, [2050, 2100, 1150, 1350, 1350, 1120], font_size=7.8)


def tuning_label(run_id: str) -> str:
    if "top1" in run_id:
        return "DenseNet-201 + GE"
    if "top2" in run_id:
        return "DenseNet-169 + Base"
    return "ResNet-50 + VIRIDIS"


def add_tuning_table(doc: Document) -> None:
    add_table_caption(doc, "Table 11.2. Bounded automated tuning trials (validation).")
    headers = ["Candidate", "Trial", "Optimizer / loss", "Augmentation", "LR / WD", "Macro-F1"]
    table = doc.add_table(rows=1, cols=len(headers))
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text
    for run_id in TUNING_RUNS:
        record = read_run(run_id)
        cfg = record["config"]
        best = record["best"]
        aug = cfg["augmentation"]
        aug_text = "Flip 0.5; rot. 5 deg" if float(aug["horizontal_flip_probability"]) else "Disabled"
        trial = str(int(run_id.rsplit("_", 1)[1]))
        values = [
            tuning_label(run_id),
            trial,
            f"{str(cfg['optimizer']).upper()} / {str(cfg['loss']).replace('_', ' ')}",
            aug_text,
            f"{float(cfg['learning_rate']):.3g} / {float(cfg['weight_decay']):.3g}",
            pct(str(best["val_macro_f1"])),
        ]
        row = table.add_row().cells
        for cell, text in zip(row, values):
            cell.text = text
    format_table(table, [2200, 700, 1800, 1500, 1700, 1220], font_size=7.6)


def add_seed_table(doc: Document) -> None:
    add_table_caption(doc, "Table 11.3. DenseNet-201 + global HE seed-confirmation runs (validation).")
    headers = ["Seed", "Accuracy", "Balanced accuracy", "Macro-F1", "Recall N / V / B", "Best epoch"]
    table = doc.add_table(rows=1, cols=len(headers))
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text
    for run_id in SEED_RUNS:
        record = read_run(run_id)
        cfg = record["config"]
        best = record["best"]
        recalls = " / ".join(
            [pct(str(best["recall_normal"])), pct(str(best["recall_viral_or_pneumonia"])), pct(str(best["recall_bacterial"]))]
        )
        values = [
            str(cfg["seed"]),
            pct(str(best["val_accuracy"])),
            pct(str(best["val_balanced_accuracy"])),
            pct(str(best["val_macro_f1"])),
            recalls,
            str(record["status"]["best_epoch"]),
        ]
        row = table.add_row().cells
        for cell, text in zip(row, values):
            cell.text = text
    format_table(table, [800, 1350, 1700, 1450, 2600, 1220], font_size=7.8)


def append_second_batch(doc: Document, *, begin_with_page_break: bool) -> None:
    if begin_with_page_break:
        doc.add_page_break()
    add_heading(doc, "11. Second-Batch Training: Shared Data Flow and Complete Results", 1)
    add_body(
        doc,
        "This section consolidates all 36 completed full validation-training runs from the second experimental batch. "
        "Smoke tests and the one-epoch pilot are excluded because they verify execution rather than model quality. "
        "The locked test set was not reopened for these additional ablations.",
    )

    add_heading(doc, "11.1 What was shared and what intentionally changed", 2)
    add_body(
        doc,
        "Dataset preparation was shared, but image preprocessing and some tuning controls were intentionally varied. "
        "This distinction is important: preparing the dataset establishes trustworthy examples and splits, whereas "
        "preprocessing changes the pixels shown to a model.",
    )
    add_bullet(
        doc,
        "Shared across all 36 full runs: the same canonical-image audit, exact-duplicate exclusions, patient/group-aware "
        "locked manifests, class order (Normal, Viral, Bacterial), 4,076/874/874 train-validation-test counts, "
        "224 x 224 model input, ImageNet initialization, and validation-only configuration selection.",
        "Shared across all 36 full runs:",
    )
    add_bullet(
        doc,
        "Shared across 30 main/control runs: random horizontal flip probability 0.5 and random rotation within +/-5 degrees, "
        "AdamW, and balanced class weighting.",
        "Shared across 30 main/control runs:",
    )
    add_bullet(
        doc,
        "Tuning exception: six SGD/focal-loss trials deliberately disabled horizontal flip, rotation, and class weighting. "
        "Those trials are not augmentation-matched controls and are identified separately in Table 11.2.",
        "Tuning exception:",
    )
    add_bullet(
        doc,
        "Intentionally varied: model architecture, preprocessing condition, architecture-safe batch size, sampled tuning "
        "hyperparameters, SAS/MFF adaptation, and the seed used in the final robustness check.",
        "Intentionally varied:",
    )

    add_heading(doc, "11.2 One real image through data preparation", 2)
    add_figure(
        doc,
        "figure_11_1_data_preparation.png",
        6.2,
        "Figure 11.1. One manifest-listed Viral Pneumonia image passing through the champion global-HE input path.",
        "Five-stage flow using person1000_virus_1681.jpeg: source, RGB conversion, global histogram equalization, 224-pixel crop, and normalized tensor.",
    )
    for lead, explanation in [
        ("Locked manifest.", "Confirms the approved path, class, patient/group identifier, and assigned split before loading."),
        ("Open and convert to RGB.", "The original grayscale image is decoded and replicated into three channels required by ImageNet-pretrained backbones."),
        ("Apply the selected preprocessing.", "The pictured champion path uses global histogram equalization; other experimental branches are shown in Figure 11.2."),
        ("Resize and center crop.", "Aspect ratio is preserved before a centered 224 x 224 crop creates the fixed spatial input."),
        ("Create the tensor.", "Pixels become float32 values, dimensions change from H x W x C to C x H x W, and ImageNet mean/std normalization is applied."),
    ]:
        add_body(doc, f"{lead} {explanation}", bold_lead=lead)

    add_heading(doc, "11.3 The same image under each preprocessing condition", 2)
    add_figure(
        doc,
        "figure_11_2_preprocessing_variants.png",
        6.0,
        "Figure 11.2. Pixel-level comparison of the six preprocessing conditions evaluated on DenseNet/ResNet representatives.",
        "Six views of the same X-ray: base RGB, CLAHE, global histogram equalization, VIRIDIS, global HE then CLAHE, and CLAHE then global HE.",
    )
    for lead, explanation in [
        ("Base/original.", "No contrast enhancement; only RGB conversion, resizing, tensor conversion, and normalization."),
        ("CLAHE.", "Locally redistributes contrast in 8 x 8 tiles with clip limit 2.0, which can reveal local texture but can amplify local noise."),
        ("Global HE.", "Redistributes intensity using the whole image and produced the strongest single validation macro-F1."),
        ("VIRIDIS.", "Maps grayscale intensity to pseudo-color. It adds no new clinical information and was exploratory."),
        ("Ordered combinations.", "GE -> CLAHE and CLAHE -> GE are different because the second transform acts on already remapped pixels."),
    ]:
        add_body(doc, f"{lead} {explanation}", bold_lead=lead)

    add_heading(doc, "11.4 Training-only data augmentation", 2)
    add_figure(
        doc,
        "figure_11_3_training_augmentation.png",
        6.2,
        "Figure 11.3. A deterministic illustration of the standard training-only flip and small-rotation sequence.",
        "Four-stage augmentation example using the same global-HE image: input, horizontal flip, five-degree rotation, and final 224-pixel view.",
    )
    add_body(
        doc,
        "Horizontal flip (p = 0.5). During standard training, approximately half of image presentations are mirrored. "
        "The operation was retained because left-right orientation does not change the three study labels.",
        bold_lead="Horizontal flip (p = 0.5).",
    )
    add_body(
        doc,
        "Small rotation (+/-5 degrees). A random angle represents modest acquisition or positioning variation without "
        "creating anatomically implausible large rotations.",
        bold_lead="Small rotation (+/-5 degrees).",
    )
    add_body(
        doc,
        "Training only. Validation and locked-test images receive deterministic preprocessing, resizing, tensor conversion, "
        "and normalization, but no random flip or rotation.",
        bold_lead="Training only.",
    )

    add_heading(doc, "11.5 End-to-end training and validation pipeline", 2)
    add_figure(
        doc,
        "figure_11_4_training_pipeline.png",
        6.2,
        "Figure 11.4. Shared experiment flow from the locked manifest to the best validation checkpoint.",
        "Nine-box training flow: manifest, image transform, tensor input, CNN backbone, logits, loss, validation, macro-F1 selection, and best checkpoint.",
    )
    for lead, explanation in [
        ("Batch construction.", "The loader groups transformed tensors using the largest architecture-safe batch size tested on the 4 GB RTX 2050."),
        ("CNN backbone.", "DenseNet, ResNet, or YOLO classification features are fine-tuned from ImageNet weights; the experimental SAS/MFF module was added only to its named DenseNet-201 runs."),
        ("Three logits.", "The network outputs one unnormalized score for Normal, Viral, and Bacterial; probabilities are derived for evaluation."),
        ("Optimization.", "The standard comparison uses class-weighted cross-entropy and AdamW; loss gradients update the model after each training batch."),
        ("Validation and checkpointing.", "Macro-F1 is primary, with balanced accuracy and per-class recalls as safeguards. The checkpoint from the best validation epoch is retained."),
    ]:
        add_body(doc, f"{lead} {explanation}", bold_lead=lead)

    doc.add_page_break()
    add_heading(doc, "11.6 Complete second-batch training results", 2)
    add_body(
        doc,
        "Table 11.1 contains the 21 main architecture/preprocessing runs. A dash is unnecessary because each row records "
        "only a configuration that actually completed. Results are percentages from the best validation checkpoint; no "
        "smoke score is included.",
    )
    add_main_results_table(doc)

    doc.add_page_break()
    add_heading(doc, "11.7 Bounded tuning and seed confirmation", 2)
    add_body(
        doc,
        "The twelve bounded tuning trials tested four configurations for each of the top three parent models. Trial 0 is the "
        "validated parent control. The original control won all three searches.",
    )
    add_tuning_table(doc)
    add_seed_table(doc)

    add_heading(doc, "11.8 Overall selection", 2)
    add_body(
        doc,
        "DenseNet-201 + global histogram equalization remains the selected validation configuration: 85.355% accuracy, "
        "84.281% balanced accuracy, and 84.471% macro-F1 at epoch 6 for seed 42. CLAHE -> GE reached a slightly higher "
        "85.584% accuracy, but its 84.396% macro-F1, 83.839% balanced accuracy, and 63.229% Viral recall did not improve the "
        "predeclared balanced-selection objective.",
    )
    add_body(
        doc,
        "The three-seed confirmation mean for DenseNet-201 + global HE was 83.895% validation macro-F1. This shows why a "
        "single seed should be interpreted as one controlled realization rather than a guarantee of identical retraining results.",
    )


def build_preview() -> None:
    doc = Document()
    configure_styles(doc)
    configure_page(doc)
    append_second_batch(doc, begin_with_page_break=False)
    doc.save(PREVIEW)


def build_final() -> None:
    if not REFERENCE.exists():
        raise FileNotFoundError(REFERENCE)
    doc = Document(REFERENCE)
    configure_styles(doc)
    configure_page(doc)
    append_second_batch(doc, begin_with_page_break=True)
    doc.save(FINAL)


def main() -> None:
    required = MAIN_RUNS + TUNING_RUNS + SEED_RUNS
    missing = [run_id for run_id in required if not (RUNS / run_id / "status.json").exists()]
    if missing:
        raise RuntimeError(f"Missing recorded runs: {missing}")
    build_preview()
    build_final()
    print(PREVIEW, PREVIEW.stat().st_size)
    print(FINAL, FINAL.stat().st_size)


if __name__ == "__main__":
    main()
