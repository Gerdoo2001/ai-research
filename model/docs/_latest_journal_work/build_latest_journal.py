"""Build the final black-only journal-style pneumonia experiment report."""

from __future__ import annotations

import csv
import hashlib
import io
import json
import shutil
import sys
import tempfile
import zipfile
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

import numpy as np
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[3]
MODEL = ROOT / "model"
RESULTS = MODEL / "results"
RUNS = RESULTS / "runs"
STAGED = RESULTS / "staged_preprocessing_experiment"
DOCS = MODEL / "docs"
WORK = Path(__file__).resolve().parent
FIGURES = WORK / "figures"
SOURCE = DOCS / "Pneumonia_Model_Training_Details_Black_Second_Batch_Complete_20260804.docx"
OUTPUT = DOCS / "Pneumonia_Three_Class_Classification_Full_Journal_20260809.docx"
AUDIT = OUTPUT.with_suffix(".audit.json")

SECOND_BATCH_WORK = DOCS / "_second_batch_word_work"
sys.path.insert(0, str(SECOND_BATCH_WORK))
import build_second_batch_word as base  # noqa: E402


TOP1_RUN = "staged_densenet201_pseudo_color_then_histogram_equalization_then_clahe_s42"
TOP2_RUN = "staged_densenet169_pseudo_color_then_histogram_equalization_then_clahe_s42"
RESNET50_WINNER_RUN = "staged_resnet50_clahe_then_pseudo_color_then_histogram_equalization_s42"

MODEL_LABELS = {
    "densenet121": "DenseNet-121",
    "densenet169": "DenseNet-169",
    "densenet201": "DenseNet-201",
    "densenet201_sas_mff": "DenseNet-201 + SAS/MFF",
    "resnet50": "ResNet-50",
    "resnet101": "ResNet-101",
    "resnet152": "ResNet-152",
    "yolov5_cls": "YOLOv5 classification",
    "yolov8_cls": "YOLOv8 classification",
}

OP_LABELS = {
    "baseline": "Base",
    "clahe": "CLAHE",
    "histogram_equalization": "GE",
    "pseudo_color": "VIRIDIS",
}


def read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def run_payload(run_id: str) -> dict[str, Any]:
    run_dir = RUNS / run_id
    return {
        "run_id": run_id,
        "config": read_json(run_dir / "config.json"),
        "status": read_json(run_dir / "status.json"),
        "metrics": read_json(run_dir / "metrics.json"),
    }


def pipeline_label(value: str) -> str:
    if value in OP_LABELS:
        return OP_LABELS[value]
    return " -> ".join(OP_LABELS.get(part, part.upper()) for part in value.split("_then_"))


def pct(value: Any, digits: int = 3) -> str:
    return f"{100.0 * float(value):.{digits}f}%"


def signed_pp(value: float) -> str:
    return f"{value * 100:+.3f} pp"


def minutes(value: Any) -> str:
    return f"{float(value) / 60.0:.1f} min"


def mib(value: Any) -> str:
    return f"{float(value) / (1024.0 ** 2):.1f} MiB"


def clear_paragraph(paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def rewrite_paragraph(paragraph, text: str, *, bold: bool = False, italic: bool = False) -> None:
    clear_paragraph(paragraph)
    run = paragraph.add_run(text)
    base.set_run_font(run, bold=bold)
    run.italic = italic


def set_cell(cell, text: str) -> None:
    cell.text = str(text)
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            base.set_run_font(run)


def add_table(
    doc: Document,
    caption: str,
    headers: list[str],
    rows: Iterable[Iterable[Any]],
    widths: list[int],
    *,
    font_size: float = 8.0,
) -> None:
    base.add_table_caption(doc, caption)
    table = doc.add_table(rows=1, cols=len(headers))
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = str(text)
    for values in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, values):
            cell.text = str(text)
    base.format_table(table, widths, font_size=font_size)


def add_figure(doc: Document, path: Path, caption: str, alt_text: str, width: float = 6.15) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    inline = run.add_picture(str(path), width=Inches(width))
    inline._inline.docPr.set("descr", alt_text)
    inline._inline.docPr.set("title", caption.split(".", 1)[0])
    cap = doc.add_paragraph(style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(7)
    cap_run = cap.add_run(caption)
    base.set_run_font(cap_run, 8.5)


def ensure_reference_style(doc: Document) -> None:
    if "Journal Reference" in [style.name for style in doc.styles]:
        return
    style = doc.styles.add_style("Journal Reference", WD_STYLE_TYPE.PARAGRAPH)
    style.font.name = "Arial"
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Arial")
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Arial")
    style.font.size = Pt(9)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.left_indent = Inches(0.28)
    style.paragraph_format.first_line_indent = Inches(-0.28)
    style.paragraph_format.space_after = Pt(5)
    style.paragraph_format.line_spacing = 1.05


def configure_journal_chrome(doc: Document) -> None:
    # Explicit user override: retain the established all-black Arial system, but
    # repackage the opening and final sections as a journal-style manuscript.
    base.configure_styles(doc)
    base.configure_page(doc)
    for section in doc.sections:
        section.different_first_page_header_footer = True
        header = section.header.paragraphs[0]
        rewrite_paragraph(header, "Three-Class Pneumonia Classification | Full Experimental Documentation")
        header.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in header.runs:
            base.set_run_font(run, 8.5)
        first_header = section.first_page_header.paragraphs[0]
        rewrite_paragraph(first_header, "")
    ensure_reference_style(doc)


def rewrite_front_matter(doc: Document) -> None:
    paragraphs = doc.paragraphs
    rewrite_paragraph(
        paragraphs[3],
        "Leakage-Controlled Three-Class Pneumonia Classification from Pediatric Chest Radiographs",
        bold=True,
    )
    rewrite_paragraph(
        paragraphs[4],
        "Complete journal-style documentation of data preparation, augmentation, model training, ordered preprocessing, tuning, validation, and final selection",
        italic=True,
    )
    rewrite_paragraph(paragraphs[5], "Normal | Viral Pneumonia | Bacterial Pneumonia")
    paragraphs[5].alignment = WD_ALIGN_PARAGRAPH.CENTER
    rewrite_paragraph(paragraphs[7], "Updated 9 August 2026 | Reconstructed exclusively from audited repository artifacts")
    paragraphs[7].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in paragraphs[7].runs:
        base.set_run_font(run, 9)
    rewrite_paragraph(
        paragraphs[8],
        "Reporting boundary: the staged ordered-preprocessing study used the locked validation set only. The test set was evaluated once for the earlier frozen DenseNet-201 + global-histogram-equalization model and was not reopened after the new validation winner emerged.",
        italic=True,
    )

    rewrite_paragraph(paragraphs[10], "Abstract", bold=True)
    paragraphs[10].style = doc.styles["Heading 1"]
    rewrite_paragraph(
        paragraphs[11],
        "This study developed and audited a three-class pediatric chest-radiograph classifier for Normal, Viral Pneumonia, and Bacterial Pneumonia. From 17,591 source files, the preparation pipeline identified 5,824 unique canonical images after excluding metadata, verified nested copies, source-layout anomalies, and 32 redundant exact images. Deterministic group-aware manifests assigned 4,076 images to training, 874 to validation, and 874 to the sealed test split with zero assigned exact-hash or group leakage.",
    )
    rewrite_paragraph(
        paragraphs[12],
        "Eight ImageNet-initialized classifiers were compared under matched optimization controls. DenseNet-201 and ResNet-50 then received exhaustive evaluation of the baseline plus every ordered subset of CLAHE, global histogram equalization (GE), and VIRIDIS pseudo-coloring (16 conditions per model). The DenseNet-201 winner, VIRIDIS -> GE -> CLAHE, achieved 85.294% validation macro F1, 84.546% balanced accuracy, and 86.613% accuracy. Transfer of that pipeline across all supported architectures placed DenseNet-169 second overall at 84.742% macro F1. ResNet-50's own winner was CLAHE -> VIRIDIS -> GE at 84.663%.",
    )
    rewrite_paragraph(
        paragraphs[13],
        "The leading DenseNet-201 configuration exceeded the second-ranked DenseNet-169 by 0.551 percentage points in macro F1, while DenseNet-169 was smaller and faster. The result also demonstrated an architecture-by-preprocessing-order interaction: the DenseNet-201 winning order produced only 82.587% on ResNet-50, whereas ResNet-50 required a different order to reach its best score.",
    )
    rewrite_paragraph(
        paragraphs[14],
        "The current validation winner has not yet received multi-seed confirmation or locked-test evaluation. Therefore, it is the best completed validation configuration, not a clinically validated diagnostic system. The only locked-test result remains the earlier DenseNet-201 + GE model, which achieved 78.877% macro F1.",
    )
    rewrite_paragraph(
        paragraphs[15],
        "Keywords: chest radiography; pneumonia classification; DenseNet; ResNet; CLAHE; histogram equalization; VIRIDIS; data leakage; macro F1; transfer learning",
        italic=True,
    )
    paragraphs[15].style = doc.styles["Normal"]
    rewrite_paragraph(paragraphs[16], "Highlights", bold=True)
    paragraphs[16].style = doc.styles["Heading 2"]
    highlights = [
        "The experiment ledger contains 119 completed entries: 95 full/analysis entries and 24 smoke or pilot checks; one historical failure remains preserved for auditability.",
        "DenseNet-201 + VIRIDIS -> GE -> CLAHE is the current validation leader at 85.294% macro F1.",
        "DenseNet-169 with the same transferred pipeline is second overall at 84.742%; ResNet-50's architecture-specific winner is third at 84.663%.",
        "The sealed test split was not used to choose any staged preprocessing order and was not reopened for the new winner.",
    ]
    for paragraph, text in zip(paragraphs[17:21], highlights):
        rewrite_paragraph(paragraph, text)
        paragraph.style = doc.styles["List Bullet"]
    rewrite_paragraph(paragraphs[21], "Research Completion Status", bold=True)
    paragraphs[21].style = doc.styles["Heading 2"]


def rewrite_status_tables(doc: Document) -> None:
    table = doc.tables[0]
    rows = [
        ("Report field", "Final value"),
        ("Status", "Ordered-preprocessing and cross-model validation stages complete"),
        ("Current validation winner", TOP1_RUN),
        ("Architecture / preprocessing", "DenseNet-201 / VIRIDIS -> GE -> CLAHE"),
        ("Winner validation macro F1", "0.852936"),
        ("Winner validation balanced accuracy", "0.845464"),
        ("Winner validation accuracy", "0.866133"),
        ("Locked-test boundary", "Current winner not tested; earlier DenseNet-201 + GE test macro F1 = 0.788772"),
        ("Generated", "August 09, 2026 (Asia/Manila)"),
    ]
    while len(table.rows) < len(rows):
        table.add_row()
    for row, values in zip(table.rows, rows):
        for cell, value in zip(row.cells, values):
            set_cell(cell, value)
    base.format_table(table, [2400, 6720], font_size=8.5)

    status = doc.tables[1]
    # Preserve the existing audit history while adding the final staged phases.
    additions = [
        (
            "Ordered preprocessing",
            "Not started",
            "Complete; DenseNet-201 16/16 and ResNet-50 16/16",
            "model/results/staged_preprocessing_experiment/*_all_orders.csv",
        ),
        (
            "Winner transfer",
            "Not started",
            "Complete; selected DenseNet-201 pipeline evaluated on 8/8 models",
            "model/results/staged_preprocessing_experiment/winner_cross_model_table.csv",
        ),
        (
            "Current final validation selection",
            "DenseNet-201 + GE",
            "DenseNet-201 + VIRIDIS -> GE -> CLAHE; validation macro F1 85.294%",
            "model/results/staged_preprocessing_experiment/densenet201_winner.json",
        ),
    ]
    for values in additions:
        cells = status.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value
    # Clarify the prior one-time test row without erasing history.
    for row in status.rows[1:]:
        if row.cells[0].text.strip() == "Locked-test evaluation":
            row.cells[2].text = "Completed once for earlier frozen DenseNet-201 + GE; not reopened for staged winner"
    base.format_table(status, [1800, 1450, 3600, 2270], font_size=7.4)


def update_previous_selection_language(doc: Document) -> None:
    for paragraph in doc.paragraphs:
        if paragraph.text.startswith("DenseNet-201 + global histogram equalization remains the selected"):
            rewrite_paragraph(
                paragraph,
                "At the close of the earlier second-batch study, DenseNet-201 + global histogram equalization was the selected validation configuration at 84.471% macro F1. The later ordered-preprocessing experiment superseded that validation ranking: DenseNet-201 + VIRIDIS -> GE -> CLAHE reached 85.294%. The earlier model remains important because it alone received the predeclared one-time locked-test evaluation.",
            )
        elif paragraph.text.startswith("The three-seed confirmation mean for DenseNet-201 + global HE"):
            rewrite_paragraph(
                paragraph,
                "The three-seed confirmation mean of 83.895% applies only to DenseNet-201 + global HE. It must not be transferred to the newer VIRIDIS -> GE -> CLAHE winner, which currently has a single completed seed-42 result.",
            )


def completed_ledger_summary() -> dict[str, int]:
    rows = read_csv(RESULTS / "experiment_ledger.csv")
    completed = [row for row in rows if row.get("status") == "completed"]
    smoke = [
        row
        for row in completed
        if row.get("run_id", "").startswith(("smoke", "pilot"))
    ]
    return {
        "total": len(rows),
        "completed": len(completed),
        "full": len(completed) - len(smoke),
        "smoke": len(smoke),
        "failed": len([row for row in rows if row.get("status") == "failed"]),
    }


def distinct_validation_ranking() -> list[dict[str, Any]]:
    candidates: dict[tuple[Any, ...], dict[str, Any]] = {}
    for row in read_csv(RESULTS / "experiment_ledger.csv"):
        run_id = row.get("run_id", "")
        if row.get("status") != "completed" or run_id.startswith(("smoke", "pilot")):
            continue
        metrics_path = RUNS / run_id / "metrics.json"
        config_path = RUNS / run_id / "config.json"
        if not metrics_path.exists() or not config_path.exists():
            continue
        metrics = read_json(metrics_path)
        config = read_json(config_path)
        if metrics.get("split") != "validation" or "macro_f1" not in metrics:
            continue
        record = {
            "run_id": run_id,
            "architecture": str(config["architecture"]),
            "preprocessing": str(config["preprocessing"]),
            "seed": int(config["seed"]),
            "metrics": metrics,
            "config": config,
            "status": read_json(RUNS / run_id / "status.json"),
        }
        # Some preregistered tuning controls intentionally reproduced an existing
        # parent checkpoint under a shorter epoch ceiling.  Deduplicate aliases
        # that have the same architecture, pixel pipeline, seed, and complete
        # validation outcome so the journal ranking does not show an
        # indistinguishable model twice.
        alias_key = (
            record["architecture"],
            record["preprocessing"],
            record["seed"],
            round(float(metrics["macro_f1"]), 12),
            round(float(metrics.get("balanced_accuracy", 0.0)), 12),
            round(float(metrics.get("accuracy", 0.0)), 12),
        )
        current = candidates.get(alias_key)
        if current is None or float(metrics["macro_f1"]) > float(current["metrics"]["macro_f1"]):
            candidates[alias_key] = record
    return sorted(candidates.values(), key=lambda row: float(row["metrics"]["macro_f1"]), reverse=True)


def make_figures(ranking: list[dict[str, Any]]) -> None:
    FIGURES.mkdir(parents=True, exist_ok=True)
    regular_path = Path(r"C:\Windows\Fonts\arial.ttf")
    bold_path = Path(r"C:\Windows\Fonts\arialbd.ttf")

    def font(size: int, *, bold: bool = False):
        path = bold_path if bold else regular_path
        return ImageFont.truetype(str(path), size) if path.exists() else ImageFont.load_default()

    def centered(draw: ImageDraw.ImageDraw, xy: tuple[int, int], text: str, selected_font, fill="black"):
        box = draw.textbbox((0, 0), text, font=selected_font)
        draw.text((xy[0] - (box[2] - box[0]) / 2, xy[1] - (box[3] - box[1]) / 2), text, font=selected_font, fill=fill)

    # Figure 12.1: monochrome horizontal ranking bars.
    canvas = Image.new("RGB", (2200, 1420), "white")
    draw = ImageDraw.Draw(canvas)
    centered(draw, (1100, 65), "TOP DISTINCT COMPLETED VALIDATION CONFIGURATIONS", font(42, bold=True))
    top = ranking[:8]
    values = [100.0 * float(row["metrics"]["macro_f1"]) for row in top]
    lower = min(values) - 0.8
    upper = max(values) + 0.45
    plot_left, plot_right = 820, 2040
    plot_top, row_gap, bar_height = 150, 148, 54
    for tick in np.arange(np.floor(lower * 2) / 2, upper + 0.01, 0.5):
        x = int(plot_left + (tick - lower) / (upper - lower) * (plot_right - plot_left))
        draw.line((x, plot_top, x, plot_top + row_gap * len(top)), fill="#D0D0D0", width=2)
        centered(draw, (x, plot_top + row_gap * len(top) + 34), f"{tick:.1f}", font(24))
    for idx, (row, value) in enumerate(zip(top, values)):
        y = plot_top + idx * row_gap + 42
        model = MODEL_LABELS.get(row["architecture"], row["architecture"])
        pipe = pipeline_label(row["preprocessing"])
        draw.text((35, y - 29), model, font=font(27, bold=True), fill="black")
        draw.text((35, y + 8), pipe, font=font(24), fill="#333333")
        x_end = int(plot_left + (value - lower) / (upper - lower) * (plot_right - plot_left))
        fill = "#111111" if idx == 0 else "#777777"
        draw.rectangle((plot_left, y - bar_height // 2, x_end, y + bar_height // 2), fill=fill, outline="black", width=2)
        draw.text((x_end + 15, y - 16), f"{value:.3f}%", font=font(25, bold=True), fill="black")
    centered(draw, (1430, 1385), "Validation macro F1 (%)", font(27, bold=True))
    canvas.save(FIGURES / "figure_12_1_top_validation_ranking.png")

    # Figure 12.2: two grayscale confusion matrices.
    canvas = Image.new("RGB", (2200, 1050), "white")
    draw = ImageDraw.Draw(canvas)
    centered(draw, (1100, 58), "TOP-TWO VALIDATION CONFUSION MATRICES", font(42, bold=True))
    top1 = run_payload(TOP1_RUN)["metrics"]
    top2 = run_payload(TOP2_RUN)["metrics"]
    for panel, (metrics, title) in enumerate(
        (
            (top1, "DenseNet-201 | VIRIDIS -> GE -> CLAHE"),
            (top2, "DenseNet-169 | VIRIDIS -> GE -> CLAHE"),
        )
    ):
        matrix = np.asarray(metrics["confusion_matrix"], dtype=int)
        panel_x = 90 + panel * 1080
        grid_x, grid_y, cell = panel_x + 260, 245, 205
        centered(draw, (panel_x + 535, 145), title, font(28, bold=True))
        for row in range(3):
            for col in range(3):
                value = int(matrix[row, col])
                intensity = int(245 - 205 * value / max(int(matrix.max()), 1))
                fill = (intensity, intensity, intensity)
                x0, y0 = grid_x + col * cell, grid_y + row * cell
                draw.rectangle((x0, y0, x0 + cell, y0 + cell), fill=fill, outline="black", width=3)
                centered(draw, (x0 + cell // 2, y0 + cell // 2), str(value), font(36, bold=True), "white" if intensity < 115 else "black")
        labels = ["Normal", "Viral", "Bacterial"]
        for idx, label in enumerate(labels):
            centered(draw, (grid_x + idx * cell + cell // 2, grid_y + 3 * cell + 43), label, font(24))
            box = draw.textbbox((0, 0), label, font=font(24))
            draw.text((grid_x - 25 - (box[2] - box[0]), grid_y + idx * cell + cell // 2 - 14), label, font=font(24), fill="black")
        centered(draw, (grid_x + 1.5 * cell, grid_y + 3 * cell + 95), "Predicted label", font(25, bold=True))
        centered(draw, (panel_x + 75, grid_y + 1.5 * cell), "True label", font(25, bold=True))
    canvas.save(FIGURES / "figure_12_2_top_two_confusion_matrices.png")

    # Figure 12.3: two monochrome epoch-history panels.
    canvas = Image.new("RGB", (2200, 980), "white")
    draw = ImageDraw.Draw(canvas)
    centered(draw, (1100, 55), "TOP-TWO TRAINING HISTORIES", font(42, bold=True))
    histories = {
        "DenseNet-201": read_csv(RUNS / TOP1_RUN / "history.csv"),
        "DenseNet-169": read_csv(RUNS / TOP2_RUN / "history.csv"),
    }

    def draw_panel(panel_x: int, title: str, key: str, ylabel: str, scale: tuple[float, float]):
        left, top_y, width, height = panel_x + 130, 190, 760, 600
        draw.rectangle((left, top_y, left + width, top_y + height), outline="black", width=3)
        centered(draw, (left + width // 2, 135), title, font(29, bold=True))
        y_min, y_max = scale
        max_epoch = max(int(row["epoch"]) for rows in histories.values() for row in rows)
        for i in range(6):
            y = top_y + int(i / 5 * height)
            value = y_max - i / 5 * (y_max - y_min)
            draw.line((left, y, left + width, y), fill="#D0D0D0", width=2)
            draw.text((left - 85, y - 13), f"{value:.1f}", font=font(21), fill="black")
        for epoch in range(1, max_epoch + 1):
            x = left + int((epoch - 1) / max(max_epoch - 1, 1) * width)
            if epoch == 1 or epoch == max_epoch or epoch % 2 == 0:
                centered(draw, (x, top_y + height + 32), str(epoch), font(20))
        for label, color, dashed in (("DenseNet-201", "#111111", False), ("DenseNet-169", "#777777", True)):
            points = []
            for row in histories[label]:
                epoch = int(row["epoch"])
                value = float(row[key]) * (100.0 if key == "val_macro_f1" else 1.0)
                x = left + int((epoch - 1) / max(max_epoch - 1, 1) * width)
                y = top_y + height - int((value - y_min) / (y_max - y_min) * height)
                points.append((x, y))
            for a, b in zip(points, points[1:]):
                if dashed:
                    segments = 8
                    for seg in range(0, segments, 2):
                        x0 = int(a[0] + (b[0] - a[0]) * seg / segments)
                        y0 = int(a[1] + (b[1] - a[1]) * seg / segments)
                        x1 = int(a[0] + (b[0] - a[0]) * (seg + 1) / segments)
                        y1 = int(a[1] + (b[1] - a[1]) * (seg + 1) / segments)
                        draw.line((x0, y0, x1, y1), fill=color, width=5)
                else:
                    draw.line((*a, *b), fill=color, width=5)
            for point in points:
                draw.ellipse((point[0] - 6, point[1] - 6, point[0] + 6, point[1] + 6), fill=color, outline="black")
        centered(draw, (left + width // 2, top_y + height + 80), "Epoch", font(23, bold=True))
        draw.text((panel_x + 5, top_y + height // 2), ylabel, font=font(23, bold=True), fill="black")
        draw.line((left + 80, 850, left + 145, 850), fill="#111111", width=5)
        draw.text((left + 155, 835), "DenseNet-201", font=font(21), fill="black")
        for x in range(left + 410, left + 475, 18):
            draw.line((x, 850, min(x + 10, left + 475), 850), fill="#777777", width=5)
        draw.text((left + 485, 835), "DenseNet-169", font=font(21), fill="black")

    f1_values = [100.0 * float(row["val_macro_f1"]) for rows in histories.values() for row in rows]
    loss_values = [float(row["val_loss"]) for rows in histories.values() for row in rows]
    draw_panel(50, "Validation macro F1 by epoch", "val_macro_f1", "Macro F1 (%)", (max(0.0, min(f1_values) - 3), min(100.0, max(f1_values) + 2)))
    draw_panel(1130, "Validation loss by epoch", "val_loss", "Cross-entropy loss", (max(0.0, min(loss_values) - 0.08), max(loss_values) + 0.08))
    canvas.save(FIGURES / "figure_12_3_top_two_training_history.png")


def all_orders_rows(path: Path) -> list[list[str]]:
    rows: list[list[str]] = []
    for row in read_csv(path):
        metrics = run_payload(row["Run ID"])["metrics"]
        rows.append(
            [
                row["Order"],
                row["Validation Macro F1"],
                row["Validation Balanced Accuracy"],
                row["Validation Accuracy"],
                str(metrics["best_epoch"]),
            ]
        )
    return rows


def append_latest_sections(doc: Document, ranking: list[dict[str, Any]]) -> None:
    ledger = completed_ledger_summary()
    top1 = run_payload(TOP1_RUN)
    top2 = run_payload(TOP2_RUN)
    r50 = run_payload(RESNET50_WINNER_RUN)
    m1, m2, mr = top1["metrics"], top2["metrics"], r50["metrics"]
    c1, c2, cr = top1["config"], top2["config"], r50["config"]

    doc.add_page_break()
    base.add_heading(doc, "12. Final Ordered-Preprocessing Study and Updated Validation Selection", 1)
    base.add_body(
        doc,
        f"The final repository ledger contains {ledger['completed']} completed entries: {ledger['full']} full or analysis entries and {ledger['smoke']} smoke/pilot checks. One historical failed row is preserved with its resolution trail. The ordered-preprocessing controller completed all 16 DenseNet-201 conditions, all 16 ResNet-50 conditions, and transfer of the DenseNet-201 winner to all eight supported architectures. Every selection in this section is validation-only.",
    )

    base.add_heading(doc, "12.1 Experimental design and why order is a real variable", 2)
    base.add_body(
        doc,
        "Each of CLAHE (C), global histogram equalization (G), and VIRIDIS (V) appears at most once. The resulting design contains the baseline, three one-operation pipelines, six ordered two-operation pipelines, and six ordered three-operation pipelines: 16 conditions per exhaustively tested architecture. All target configurations were derived from the architecture's seed-42 baseline so optimizer, learning rate, weight decay, augmentation, epoch budget, loss, scheduler, manifests, and class definitions remained matched.",
    )
    add_table(
        doc,
        "Table 12.1. Controlled variables and intentionally varied factors in the ordered-preprocessing study.",
        ["Component", "Controlled or varied", "Implementation"],
        [
            ("Dataset and split", "Controlled", "Locked group-aware manifests: 4,076 train; 874 validation; 874 sealed test"),
            ("Random seed", "Controlled", "42 for every ordered-preprocessing comparison"),
            ("Training augmentation", "Controlled", "Horizontal flip p=0.5; rotation uniformly sampled within +/-5 degrees"),
            ("Optimization", "Controlled", "AdamW; LR 0.0001; weight decay 0.0001; balanced class-weighted cross-entropy"),
            ("Training budget", "Controlled", "Maximum 12 epochs; early-stopping patience 4; ReduceLROnPlateau factor 0.5/patience 2"),
            ("Architecture", "Varied by phase", "DenseNet-201 exhaustive; ResNet-50 exhaustive; winner transferred to eight architectures"),
            ("Enhancement sequence", "Varied", "All ordered subsets of CLAHE, GE, and VIRIDIS"),
            ("Selection", "Controlled", "Maximum validation macro F1; balanced accuracy and accuracy as tie-breakers"),
        ],
        [2100, 1800, 5220],
        font_size=7.8,
    )
    base.add_body(
        doc,
        "Order matters because every operation consumes the output of the prior operation. In this implementation, CLAHE and GE convert their input to grayscale and then replicate it to RGB. VIRIDIS also begins by converting its input to grayscale before applying the color map. Therefore, a pipeline that ends in GE or CLAHE is finally grayscale-replicated even if VIRIDIS appeared earlier; VIRIDIS can still change the intensity mapping that later grayscale operations receive. The observed order effects are algorithmic interactions, not evidence that artificial color adds new clinical information.",
    )

    doc.add_page_break()
    base.add_heading(doc, "12.2 Completed six-condition screening across all eight architectures", 2)
    six = read_csv(STAGED / "current_six_column_table.csv")
    six_headers = list(six[0])
    add_table(
        doc,
        "Table 12.2. Validation macro F1 (%) for the completed six-condition screening matrix.",
        six_headers,
        [[row[column] for column in six_headers] for row in six],
        [1650, 1050, 1050, 1050, 1100, 1610, 1610],
        font_size=7.2,
    )
    base.add_body(
        doc,
        "The six-condition screen shows that preprocessing is architecture-dependent. GE -> CLAHE improved ResNet-50 from 83.692% to 84.579% macro F1, but ResNet-101 and ResNet-152 remained strongest at baseline. YOLOv5 also remained strongest at baseline, while YOLOv8's best screened condition was VIRIDIS.",
    )

    base.add_heading(doc, "12.3 Complete DenseNet-201 and ResNet-50 order tables", 2)
    add_table(
        doc,
        "Table 12.3. DenseNet-201: all 16 completed preprocessing orders (validation).",
        ["Order", "Macro F1", "Balanced acc.", "Accuracy", "Best epoch"],
        all_orders_rows(STAGED / "densenet201_all_orders.csv"),
        [3000, 1600, 1700, 1600, 1220],
        font_size=7.8,
    )
    doc.add_page_break()
    add_table(
        doc,
        "Table 12.4. ResNet-50: all 16 completed preprocessing orders (validation).",
        ["Order", "Macro F1", "Balanced acc.", "Accuracy", "Best epoch"],
        all_orders_rows(STAGED / "resnet50_all_orders.csv"),
        [3000, 1600, 1700, 1600, 1220],
        font_size=7.8,
    )
    base.add_body(
        doc,
        f"DenseNet-201 selected {pipeline_label(c1['preprocessing'])} at {pct(m1['macro_f1'])}. ResNet-50 selected {pipeline_label(cr['preprocessing'])} at {pct(mr['macro_f1'])}. DenseNet-201 leads the best ResNet-50 result by {signed_pp(float(m1['macro_f1']) - float(mr['macro_f1']))}.",
    )

    base.add_heading(doc, "12.4 Transfer of the DenseNet-201 winning pipeline", 2)
    transfer = read_csv(STAGED / "winner_cross_model_table.csv")
    add_table(
        doc,
        "Table 12.5. Cross-model transfer of VIRIDIS -> GE -> CLAHE (validation).",
        ["Model", "Pipeline", "Macro F1", "Balanced acc.", "Accuracy", "Status"],
        [
            [
                row["Model"],
                row["Winning Pipeline"],
                row["Validation Macro F1"],
                row["Validation Balanced Accuracy"],
                row["Validation Accuracy"],
                row["Status"],
            ]
            for row in transfer
        ],
        [1650, 2250, 1250, 1450, 1250, 1270],
        font_size=7.5,
    )
    base.add_body(
        doc,
        "The winning DenseNet-201 order did not transfer uniformly. It ranked second overall on DenseNet-169 (84.742%) but reduced ResNet-50 to 82.587%, below both its baseline and its architecture-specific winner. This is why one model's best preprocessing cannot be assumed to be optimal for another architecture.",
    )

    doc.add_page_break()
    base.add_heading(doc, "12.5 Overall ranking of distinct completed validation configurations", 2)
    ranking_rows = []
    for rank, row in enumerate(ranking[:12], start=1):
        metrics = row["metrics"]
        ranking_rows.append(
            [
                rank,
                MODEL_LABELS.get(row["architecture"], row["architecture"]),
                pipeline_label(row["preprocessing"]),
                pct(metrics["macro_f1"]),
                pct(metrics["balanced_accuracy"]),
                pct(metrics["accuracy"]),
                row["status"].get("best_epoch", metrics.get("best_epoch", "")),
            ]
        )
    add_table(
        doc,
        "Table 12.6. Top 12 distinct completed seed-42 validation configurations after configuration-alias deduplication.",
        ["Rank", "Model", "Preprocessing", "Macro F1", "Balanced", "Accuracy", "Epoch"],
        ranking_rows,
        [620, 1650, 2700, 1130, 1130, 1050, 840],
        font_size=7.2,
    )
    add_figure(
        doc,
        FIGURES / "figure_12_1_top_validation_ranking.png",
        "Figure 12.1. Highest distinct completed validation configurations ranked by macro F1.",
        "Horizontal grayscale bar chart showing the top eight distinct validation configurations. DenseNet-201 with VIRIDIS then GE then CLAHE is highest, followed by DenseNet-169 with the same pipeline.",
    )

    base.add_heading(doc, "12.6 Top-two overall comparison", 2)
    base.add_body(
        doc,
        f"The top two distinct completed configurations use the same preprocessing order but different DenseNet depths. DenseNet-201 leads DenseNet-169 by {signed_pp(float(m1['macro_f1']) - float(m2['macro_f1']))} in macro F1 and {signed_pp(float(m1['accuracy']) - float(m2['accuracy']))} in accuracy. Balanced accuracy differs by only {signed_pp(float(m1['balanced_accuracy']) - float(m2['balanced_accuracy']))}, indicating that their average class recalls are nearly identical.",
    )
    comparison_rows = [
        ("Validation macro F1", pct(m1["macro_f1"]), pct(m2["macro_f1"]), signed_pp(float(m1["macro_f1"]) - float(m2["macro_f1"]))),
        ("Balanced accuracy", pct(m1["balanced_accuracy"]), pct(m2["balanced_accuracy"]), signed_pp(float(m1["balanced_accuracy"]) - float(m2["balanced_accuracy"]))),
        ("Accuracy", pct(m1["accuracy"]), pct(m2["accuracy"]), signed_pp(float(m1["accuracy"]) - float(m2["accuracy"]))),
        ("Macro precision", pct(m1["macro_precision"]), pct(m2["macro_precision"]), signed_pp(float(m1["macro_precision"]) - float(m2["macro_precision"]))),
        ("Macro ROC AUC", pct(m1["roc_auc_ovr_macro"]), pct(m2["roc_auc_ovr_macro"]), signed_pp(float(m1["roc_auc_ovr_macro"]) - float(m2["roc_auc_ovr_macro"]))),
        ("Macro PR AUC", pct(m1["pr_auc_macro"]), pct(m2["pr_auc_macro"]), signed_pp(float(m1["pr_auc_macro"]) - float(m2["pr_auc_macro"]))),
        ("Expected calibration error", f"{float(m1['expected_calibration_error']):.4f}", f"{float(m2['expected_calibration_error']):.4f}", f"{float(m1['expected_calibration_error']) - float(m2['expected_calibration_error']):+.4f}"),
        ("Brier score", f"{float(m1['brier_score']):.4f}", f"{float(m2['brier_score']):.4f}", f"{float(m1['brier_score']) - float(m2['brier_score']):+.4f}"),
        ("Parameters", f"{int(m1['parameter_count']):,}", f"{int(m2['parameter_count']):,}", f"{int(m1['parameter_count']) - int(m2['parameter_count']):+,}"),
        ("Model file size", mib(m1["model_file_size_bytes"]), mib(m2["model_file_size_bytes"]), f"{(float(m1['model_file_size_bytes']) - float(m2['model_file_size_bytes'])) / (1024 ** 2):+.1f} MiB"),
        ("Batch-1 model latency", f"{float(m1['inference']['latency_ms_per_image']):.2f} ms", f"{float(m2['inference']['latency_ms_per_image']):.2f} ms", f"{float(m1['inference']['latency_ms_per_image']) - float(m2['inference']['latency_ms_per_image']):+.2f} ms"),
        ("Training duration", minutes(m1["training_duration_seconds"]), minutes(m2["training_duration_seconds"]), f"{(float(m1['training_duration_seconds']) - float(m2['training_duration_seconds'])) / 60:+.1f} min"),
        ("Best epoch", str(m1["best_epoch"]), str(m2["best_epoch"]), "same" if m1["best_epoch"] == m2["best_epoch"] else "different"),
    ]
    add_table(
        doc,
        "Table 12.7. Detailed comparison of the top two overall validation configurations.",
        ["Metric", "DenseNet-201", "DenseNet-169", "D201 - D169"],
        comparison_rows,
        [3100, 1900, 1900, 2220],
        font_size=7.8,
    )
    per_class_rows = []
    for class_name in ("Normal", "Viral", "Bacterial"):
        for label, metrics in (("DenseNet-201", m1), ("DenseNet-169", m2)):
            values = metrics["per_class"][class_name]
            per_class_rows.append(
                [
                    class_name,
                    label,
                    pct(values["precision"]),
                    pct(values["recall"]),
                    pct(values["f1"]),
                    pct(values["roc_auc"]),
                    pct(values["pr_auc"]),
                ]
            )
    add_table(
        doc,
        "Table 12.8. Per-class validation metrics for the top two configurations.",
        ["Class", "Model", "Precision", "Recall", "F1", "ROC AUC", "PR AUC"],
        per_class_rows,
        [1450, 1600, 1250, 1150, 1100, 1250, 1320],
        font_size=7.6,
    )
    add_figure(
        doc,
        FIGURES / "figure_12_2_top_two_confusion_matrices.png",
        "Figure 12.2. Validation confusion matrices for the top two overall configurations.",
        "Two grayscale confusion matrices compare DenseNet-201 and DenseNet-169 using the VIRIDIS then GE then CLAHE preprocessing order across Normal, Viral, and Bacterial labels.",
    )
    add_figure(
        doc,
        FIGURES / "figure_12_3_top_two_training_history.png",
        "Figure 12.3. Validation macro F1 and validation loss across epochs for the top two configurations.",
        "Two grayscale line charts show epoch-level validation macro F1 and validation loss for DenseNet-201 and DenseNet-169.",
    )

    base.add_heading(doc, "12.7 Configuration comparison of the leading models", 2)
    config_rows = [
        ("Architecture", "DenseNet-201", "DenseNet-169", "ResNet-50"),
        ("Preprocessing", pipeline_label(c1["preprocessing"]), pipeline_label(c2["preprocessing"]), pipeline_label(cr["preprocessing"])),
        ("ImageNet initialization", c1["weights"], c2["weights"], cr["weights"]),
        ("Input", f"3 x {c1['input_size']} x {c1['input_size']}", f"3 x {c2['input_size']} x {c2['input_size']}", f"3 x {cr['input_size']} x {cr['input_size']}"),
        ("Batch size", c1["batch_size"], c2["batch_size"], cr["batch_size"]),
        ("Optimizer", c1["optimizer"], c2["optimizer"], cr["optimizer"]),
        ("Initial learning rate", c1["learning_rate"], c2["learning_rate"], cr["learning_rate"]),
        ("Weight decay", c1["weight_decay"], c2["weight_decay"], cr["weight_decay"]),
        ("Loss", "balanced weighted cross-entropy", "balanced weighted cross-entropy", "balanced weighted cross-entropy"),
        ("Maximum epochs", c1["epochs"], c2["epochs"], cr["epochs"]),
        ("Early stopping patience", c1["early_stopping_patience"], c2["early_stopping_patience"], cr["early_stopping_patience"]),
        ("Scheduler", "ReduceLROnPlateau", "ReduceLROnPlateau", "ReduceLROnPlateau"),
        ("Augmentation", "flip p=0.5; rotation +/-5 degrees", "same", "same"),
        ("Mixed precision", c1["amp"], c2["amp"], cr["amp"]),
        ("Seed", c1["seed"], c2["seed"], cr["seed"]),
        ("Validation macro F1", pct(m1["macro_f1"]), pct(m2["macro_f1"]), pct(mr["macro_f1"])),
    ]
    add_table(
        doc,
        "Table 12.9. Configuration and result comparison: overall top two plus the ResNet-50 exhaustive winner.",
        ["Field", "DenseNet-201", "DenseNet-169", "ResNet-50"],
        config_rows,
        [2450, 2230, 2210, 2230],
        font_size=7.5,
    )

    base.add_heading(doc, "12.8 Relationship to tuning, SAS/MFF, and seed robustness", 2)
    sas_rows = []
    for run_id in (
        "experimental_densenet201_sas_mff_clahe_s42",
        "experimental_densenet201_sas_mff_histogram_equalization_s42",
    ):
        record = run_payload(run_id)
        sas_rows.append(
            [
                MODEL_LABELS[record["config"]["architecture"]],
                pipeline_label(record["config"]["preprocessing"]),
                pct(record["metrics"]["macro_f1"]),
                pct(record["metrics"]["balanced_accuracy"]),
                pct(record["metrics"]["accuracy"]),
                record["metrics"]["best_epoch"],
            ]
        )
    add_table(
        doc,
        "Table 12.10. Completed experimental DenseNet-201 + SAS/MFF comparisons (validation).",
        ["Architecture", "Preprocessing", "Macro F1", "Balanced", "Accuracy", "Epoch"],
        sas_rows,
        [2450, 1800, 1250, 1250, 1250, 1120],
        font_size=7.7,
    )
    base.add_body(
        doc,
        "The bounded Optuna studies had already shown that the original AdamW, balanced cross-entropy, conservative augmentation, learning rate 0.0001, and weight decay 0.0001 controls outperformed the tested SGD/focal alternatives for each parent candidate. The SAS/MFF additions also did not exceed standard DenseNet-201. These findings justify retaining the controlled optimization recipe while varying preprocessing order. However, the three-seed robustness experiment applies only to the earlier DenseNet-201 + GE configuration; the new winner still requires independent seeds.",
    )

    base.add_heading(doc, "12.9 Locked-test result and current selection boundary", 2)
    final_selection = read_json(RESULTS / "final_selection.json")
    test = final_selection["test_metrics"]
    add_table(
        doc,
        "Table 12.11. The only locked-test evaluation, retained for context but not reassigned to the new winner.",
        ["Evaluated model", "Test macro F1", "Balanced acc.", "Accuracy", "Viral F1", "Status"],
        [
            [
                "DenseNet-201 + GE (earlier frozen model)",
                pct(test["macro_f1"]),
                pct(test["balanced_accuracy"]),
                pct(test["accuracy"]),
                pct(test["per_class"]["Viral"]["f1"]),
                "Single sealed evaluation",
            ]
        ],
        [2780, 1250, 1350, 1150, 1120, 1470],
        font_size=7.7,
    )
    base.add_body(
        doc,
        "The 78.877% locked-test macro F1 belongs to DenseNet-201 + GE, not to DenseNet-201 + VIRIDIS -> GE -> CLAHE. Reporting the old test score as the new winner's performance would be incorrect. A future test of the new winner would constitute an additional test-set access and should be governed by a newly documented protocol or, preferably, an independent external dataset.",
    )

    doc.add_page_break()
    base.add_heading(doc, "13. Discussion", 1)
    base.add_heading(doc, "13.1 Why DenseNet-201 is the best current validation model", 2)
    base.add_body(
        doc,
        "DenseNet-201 + VIRIDIS -> GE -> CLAHE is the best current validation configuration because it has the highest completed validation macro F1 after distinct-configuration deduplication. Macro F1 gives equal weight to Normal, Viral, and Bacterial F1 scores, making it more appropriate than accuracy alone for the imbalanced class distribution. Its 86.613% accuracy is also the highest among the two leading configurations, but the choice is based primarily on 85.294% macro F1 rather than raw accuracy.",
    )
    base.add_body(
        doc,
        "The model is not uniformly superior on every measure. DenseNet-169 is approximately 31% smaller by parameter count, has a smaller model file, and its measured batch-1 model latency is lower. DenseNet-169 also produced a slightly higher Viral F1, while DenseNet-201 produced higher Normal and Bacterial F1. Thus DenseNet-201 is preferred for balanced predictive performance; DenseNet-169 is the stronger efficiency alternative.",
    )

    base.add_heading(doc, "13.2 Why preprocessing order changed the result", 2)
    base.add_body(
        doc,
        "Histogram operations are nonlinear and non-commutative: applying A then B does not generally equal applying B then A. CLAHE redistributes local contrast under a clip limit, GE redistributes the whole-image histogram, and VIRIDIS maps intensity to artificial RGB before later operations may convert that output back to luminance. Consequently, each order presents a different pixel distribution to the pretrained network even though it uses the same source radiograph.",
    )
    base.add_body(
        doc,
        "The strongest evidence for architecture interaction is ResNet-50. DenseNet-201's winner scored 82.587% on ResNet-50, whereas CLAHE -> VIRIDIS -> GE reached 84.663% on the same ResNet-50 architecture under matched controls. The data therefore support testing preprocessing jointly with architecture instead of treating preprocessing as a universal front end.",
    )

    base.add_heading(doc, "13.3 Interpretation of accuracy, macro F1, and balanced accuracy", 2)
    base.add_body(
        doc,
        "Accuracy counts all correct predictions and is influenced most by the largest class. Macro F1 computes F1 separately for each class and averages the three values equally, so weak Viral performance cannot be hidden by strong Bacterial performance. Balanced accuracy averages the three recalls and answers a complementary question: how much of each class was found. The selection rule therefore prioritizes macro F1, uses balanced accuracy and accuracy as tie-breakers, and reports per-class precision and recall to expose the clinical error pattern.",
    )

    base.add_heading(doc, "13.4 Limitations", 2)
    for text in [
        "The dataset is a single public pediatric chest-X-ray source; external hospitals, devices, age groups, and prevalence patterns were not tested.",
        "The new DenseNet-201 and DenseNet-169 leaders currently have only seed-42 evidence. Their small 0.551-point separation may change under additional seeds.",
        "The current winner was selected after the earlier one-time locked-test evaluation. It has validation evidence only and must not inherit the earlier test result.",
        "The ordered preprocessing study tested pixel transformations but did not add lung segmentation, ROI masks, PCA, FiLM conditioning, or a new SAS/MFF architecture to every backbone.",
        "VIRIDIS is exploratory and adds no clinical information. When it precedes GE/CLAHE, the final image becomes grayscale-replicated; the gain is an empirical intensity-remapping effect, not proof that color is clinically meaningful.",
        "Grad-CAM, calibration, and end-to-end efficiency measurements were produced for the earlier frozen model and have not yet been regenerated for the new validation winner.",
        "No statistical confidence interval or paired significance test was computed for the 0.551-point top-two difference.",
        "This work is an engineering research benchmark and not a medical device, diagnosis, treatment recommendation, or prospective clinical validation.",
    ]:
        base.add_bullet(doc, text)

    base.add_heading(doc, "13.5 Recommended next validation sequence", 2)
    for text in [
        "Retrain DenseNet-201 + VIRIDIS -> GE -> CLAHE and DenseNet-169 + the same pipeline at seeds 123 and 2026.",
        "Compare mean macro F1, standard deviation, per-class recall, calibration, and paired predictions across seeds.",
        "Freeze one final checkpoint-selection rule without consulting the sealed test labels.",
        "Prefer an external dataset for final confirmation; if the existing locked test is reopened, document that this is a second access and no longer the original one-shot test protocol.",
        "Regenerate Grad-CAM, calibration, and true end-to-end latency for the final frozen winner.",
    ]:
        base.add_bullet(doc, text)

    base.add_heading(doc, "14. Conclusion", 1)
    base.add_body(
        doc,
        "The repository now provides a complete, leakage-controlled comparison of eight supported classifiers, exhaustive ordered preprocessing for DenseNet-201 and ResNet-50, cross-model transfer of the DenseNet-201 winner, bounded hyperparameter trials, a prior three-seed robustness study, and a preserved one-time locked-test evaluation. The strongest current validation result is DenseNet-201 + VIRIDIS -> GE -> CLAHE at 85.294% macro F1. DenseNet-169 with the same pipeline is second overall at 84.742%, and ResNet-50's best architecture-specific order reaches 84.663%.",
    )
    base.add_body(
        doc,
        "DenseNet-201 is therefore the best overall validation model, while DenseNet-169 offers the more efficient top-two alternative. This conclusion is deliberately limited to the audited validation evidence. Multi-seed confirmation and independent test or external validation are still required before the new winner can be presented as a robust final model.",
    )

    doc.add_page_break()
    base.add_heading(doc, "15. Verified References", 1)
    evidence = read_csv(RESULTS / "evidence" / "research_evidence_matrix.csv")
    base.add_body(
        doc,
        f"The repository evidence matrix contains {len(evidence)} verified primary or official sources: {Counter(row['Evidence Applicability'] for row in evidence)['Direct']} Direct, {Counter(row['Evidence Applicability'] for row in evidence)['Partial']} Partial, {Counter(row['Evidence Applicability'] for row in evidence)['Background']} Background, and {Counter(row['Evidence Applicability'] for row in evidence)['Excluded']} Excluded. The following author-year bibliography preserves the verified title, venue, DOI, and official link recorded in that matrix.",
    )
    for row in sorted(evidence, key=lambda item: (item.get("Authors", "").lower(), item.get("Year", ""), item.get("Article Title", "").lower())):
        authors = row.get("Authors", "").strip() or "Author information verified in source"
        year = row.get("Year", "").strip() or "n.d."
        title = row.get("Article Title", "").strip()
        venue = row.get("Venue", "").strip()
        doi = row.get("DOI", "").strip()
        link = row.get("Official Link", "").strip()
        parts = [f"{authors} ({year}). {title}."]
        if venue:
            parts.append(venue + ".")
        if doi and doi.lower() not in {"not reported", "n/a"}:
            parts.append(f"DOI: {doi}.")
        if link:
            parts.append(link)
        paragraph = doc.add_paragraph(" ".join(parts), style="Journal Reference")
        for run in paragraph.runs:
            base.set_run_font(run, 9)

    base.add_heading(doc, "Appendix A. Reproducibility and Artifact Index", 1)
    artifact_rows = [
        ("Locked split summary", "model/data_manifests/split_summary.json"),
        ("Dataset source snapshot", "model/data_manifests/dataset_source_snapshot.json"),
        ("Complete experiment ledger", "model/results/experiment_ledger.csv"),
        ("Six-condition comparison", "model/results/staged_preprocessing_experiment/current_six_column_table.csv"),
        ("DenseNet-201 all orders", "model/results/staged_preprocessing_experiment/densenet201_all_orders.csv"),
        ("ResNet-50 all orders", "model/results/staged_preprocessing_experiment/resnet50_all_orders.csv"),
        ("Cross-model winner transfer", "model/results/staged_preprocessing_experiment/winner_cross_model_table.csv"),
        ("DenseNet-201 winner record", "model/results/staged_preprocessing_experiment/densenet201_winner.json"),
        ("ResNet-50 winner record", "model/results/staged_preprocessing_experiment/resnet50_winner.json"),
        ("Tuning trials", "model/results/tuning_trials.csv"),
        ("Seed robustness", "model/results/seed_robustness.csv"),
        ("Earlier frozen selection/test", "model/results/final_selection.json"),
        ("Verified RRL matrix", "model/results/evidence/research_evidence_matrix.csv"),
        ("Training journal", "model/docs/TRAINING_JOURNAL.md"),
    ]
    add_table(
        doc,
        "Table A.1. Primary reproducibility artifacts used in this manuscript.",
        ["Artifact", "Repository-relative path"],
        artifact_rows,
        [2900, 6220],
        font_size=8.0,
    )
    base.add_body(
        doc,
        "Reproduction environment: Windows, Python 3.12.6, PyTorch 2.5.1+cu121, Torchvision 0.20.1+cu121, Ultralytics 8.3.168, CUDA 12.1, and an NVIDIA GeForce RTX 2050 with 4 GB VRAM. The manifest version is three_class_seed42_3b815cf769c8. Configuration hashes, status files, best checkpoints, per-epoch histories, predictions, confusion matrices, and metric JSON files are retained inside each run directory.",
    )


def validate_sources() -> None:
    required = [
        SOURCE,
        STAGED / "current_six_column_table.csv",
        STAGED / "densenet201_all_orders.csv",
        STAGED / "resnet50_all_orders.csv",
        STAGED / "winner_cross_model_table.csv",
        RESULTS / "experiment_ledger.csv",
        RESULTS / "evidence" / "research_evidence_matrix.csv",
        RESULTS / "final_selection.json",
    ]
    missing = [str(path) for path in required if not path.exists()]
    if missing:
        raise FileNotFoundError(f"Missing required source artifacts: {missing}")
    for path in (STAGED / "densenet201_all_orders.csv", STAGED / "resnet50_all_orders.csv"):
        rows = read_csv(path)
        if len(rows) != 16 or any(row["Status"] != "Completed" for row in rows):
            raise RuntimeError(f"Expected 16 completed rows in {path}")
    progress = read_json(STAGED / "progress.json")
    if progress.get("state") != "completed" or progress.get("locked_test_accessed") is not False:
        raise RuntimeError(f"Staged experiment is not in the expected completed validation-only state: {progress}")


def audit_output(path: Path, ranking: list[dict[str, Any]]) -> dict[str, Any]:
    with zipfile.ZipFile(path) as package:
        corrupt = package.testzip()
        package_files = package.namelist()
    doc = Document(path)
    colors = set()
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            if run.font.color and run.font.color.rgb:
                colors.add(str(run.font.color.rgb))
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        if run.font.color and run.font.color.rgb:
                            colors.add(str(run.font.color.rgb))
    top_ids = [ranking[0]["run_id"], ranking[1]["run_id"]]
    return {
        "path": str(path),
        "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
        "bytes": path.stat().st_size,
        "zip_test": "ok" if corrupt is None else corrupt,
        "package_parts": len(package_files),
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "inline_images": len(doc.inline_shapes),
        "sections": len(doc.sections),
        "visible_text_colors": sorted(colors),
        "all_text_black": all(color == "000000" for color in colors),
        "densenet201_order_rows": len(read_csv(STAGED / "densenet201_all_orders.csv")),
        "resnet50_order_rows": len(read_csv(STAGED / "resnet50_all_orders.csv")),
        "winner_transfer_rows": len(read_csv(STAGED / "winner_cross_model_table.csv")),
        "reference_rows": len(read_csv(RESULTS / "evidence" / "research_evidence_matrix.csv")),
        "top_two_run_ids": top_ids,
        "expected_top_two_run_ids": [TOP1_RUN, TOP2_RUN],
        "top_two_match_expected": top_ids == [TOP1_RUN, TOP2_RUN],
        "locked_test_claim": "Only the earlier DenseNet-201 + GE frozen model has a locked-test result; the current validation winner has not been tested.",
    }


def enforce_black_office_theme(path: Path) -> None:
    """Prevent Word theme overrides from rendering nominally black accents in blue."""
    with zipfile.ZipFile(path, "r") as source, tempfile.NamedTemporaryFile(
        suffix=".docx", delete=False, dir=path.parent
    ) as handle:
        temporary_path = Path(handle.name)
    try:
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(
            temporary_path, "w", compression=zipfile.ZIP_DEFLATED
        ) as target:
            for item in source.infolist():
                data = source.read(item.filename)
                if item.filename == "word/theme/theme1.xml":
                    text = data.decode("utf-8")
                    for original in (
                        "4F81BD",  # accent 1 / legacy Office blue
                        "C0504D",
                        "9BBB59",
                        "8064A2",
                        "4BACC6",
                        "F79646",
                        "0000FF",  # hyperlink
                        "800080",  # followed hyperlink
                    ):
                        text = text.replace(f'val="{original}"', 'val="000000"')
                    data = text.encode("utf-8")
                target.writestr(item, data)
        temporary_path.replace(path)
    finally:
        if temporary_path.exists():
            temporary_path.unlink()


def add_missing_image_alt_text(doc: Document) -> None:
    descriptions = {
        1: "Grayscale comparison of the supported chest X-ray preprocessing conditions, including base, CLAHE, global histogram equalization, and VIRIDIS.",
        2: "Locked-test confusion matrix for the previously frozen DenseNet-201 plus global histogram equalization model.",
        3: "Grad-CAM example for a correctly classified Normal chest X-ray.",
        4: "Grad-CAM example for a Normal chest X-ray incorrectly classified as Viral Pneumonia.",
        5: "Grad-CAM example for a correctly classified Viral Pneumonia chest X-ray.",
        6: "Grad-CAM example for a Viral Pneumonia chest X-ray incorrectly classified as Bacterial Pneumonia.",
        7: "Grad-CAM example for a correctly classified Bacterial Pneumonia chest X-ray.",
        8: "Grad-CAM example for a Bacterial Pneumonia chest X-ray incorrectly classified as Viral Pneumonia.",
        9: "Locked-test reliability diagrams before and after validation-fitted temperature scaling.",
    }
    for index, shape in enumerate(doc.inline_shapes, start=1):
        doc_properties = shape._inline.docPr
        if not doc_properties.get("descr"):
            doc_properties.set("descr", descriptions.get(index, f"Report figure {index}."))
        if not doc_properties.get("title"):
            doc_properties.set("title", f"Figure {index}")


def grayscale_legacy_confusion_matrix(doc: Document) -> None:
    """Convert the old blue locked-test confusion matrix to grayscale."""
    if len(doc.inline_shapes) < 2:
        return
    blip = doc.inline_shapes[1]._inline.xpath(".//a:blip")[0]
    relationship_id = blip.get(qn("r:embed"))
    image_part = doc.part.related_parts[relationship_id]
    with Image.open(io.BytesIO(image_part.blob)) as source:
        grayscale = source.convert("L").convert("RGB")
        output = io.BytesIO()
        grayscale.save(output, format="PNG", optimize=True)
    image_part._blob = output.getvalue()


def compact_phash_review_table(doc: Document) -> None:
    """Keep the 16-row manual pHash review together without a one-row spill page."""
    for table in doc.tables:
        header = [cell.text.strip() for cell in table.rows[0].cells]
        if header[:2] != ["Pair", "Assigned splits"]:
            continue
        for row in table.rows:
            for cell in row.cells:
                base.set_cell_margins(cell, top=20, start=80, bottom=20, end=80)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.paragraph_format.line_spacing = 1.0
                    for run in paragraph.runs:
                        base.set_run_font(run, 7.0)
        break


def main() -> None:
    validate_sources()
    ranking = distinct_validation_ranking()
    if [ranking[0]["run_id"], ranking[1]["run_id"]] != [TOP1_RUN, TOP2_RUN]:
        raise RuntimeError(
            "Unexpected top-two validation ranking: "
            + str([ranking[0]["run_id"], ranking[1]["run_id"]])
        )
    make_figures(ranking)
    working = WORK / "working_journal.docx"
    shutil.copy2(SOURCE, working)
    doc = Document(working)
    configure_journal_chrome(doc)
    rewrite_front_matter(doc)
    rewrite_status_tables(doc)
    update_previous_selection_language(doc)
    append_latest_sections(doc, ranking)
    doc.core_properties.title = "Leakage-Controlled Three-Class Pneumonia Classification from Pediatric Chest Radiographs"
    doc.core_properties.subject = "Complete journal-style experimental documentation"
    doc.core_properties.keywords = "pneumonia, chest X-ray, DenseNet, ResNet, preprocessing, macro F1"
    doc.core_properties.author = "Audited project repository"
    doc.core_properties.last_modified_by = "Audited project repository"
    doc.core_properties.modified = datetime.now()
    compact_phash_review_table(doc)
    grayscale_legacy_confusion_matrix(doc)
    add_missing_image_alt_text(doc)
    doc.save(OUTPUT)
    enforce_black_office_theme(OUTPUT)
    audit = audit_output(OUTPUT, ranking)
    AUDIT.write_text(json.dumps(audit, indent=2), encoding="utf-8")
    print(json.dumps(audit, indent=2))


if __name__ == "__main__":
    main()
